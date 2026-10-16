"""
Enterprise Presentation Status Overlay Manager
==============================================
Headless status badge stamper, stripper, and inspector for PowerPoint (.pptx).
Enforces zero drop shadows (purges OpenXML `<a:effectLst>` and sets `effectRef idx="0"`),
sharp rectangular geometry (`MSO_SHAPE.RECTANGLE`), and standardized badge identifiers.
"""

from __future__ import annotations

import logging
import os
import uuid
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Union

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from src.core.delivery.exceptions import DeliveryError
from src.core.delivery.models import OverlayReport, OverlayStatus, SlideBadgeInfo

logger = logging.getLogger(__name__)

# Standard internal name for status overlay badges
STATUS_BADGE_NAME = "__STATUS_OVERLAY_BADGE__"

# Color mappings for standard delivery statuses
_STATUS_PALETTE: Dict[str, Dict[str, RGBColor]] = {
    OverlayStatus.DRAFT.value: {
        "fill": RGBColor(217, 119, 6),       # Amber (#D97706)
        "text": RGBColor(255, 255, 255),
    },
    OverlayStatus.IN_REVIEW.value: {
        "fill": RGBColor(2, 132, 199),       # Sky Blue (#0284C7)
        "text": RGBColor(255, 255, 255),
    },
    OverlayStatus.APPROVED.value: {
        "fill": RGBColor(22, 163, 74),       # Emerald Green (#16A34A)
        "text": RGBColor(255, 255, 255),
    },
    OverlayStatus.FINAL.value: {
        "fill": RGBColor(15, 23, 42),        # Slate Navy (#0F172A)
        "text": RGBColor(255, 255, 255),
    },
    OverlayStatus.CONFIDENTIAL.value: {
        "fill": RGBColor(220, 38, 38),       # Crimson Red (#DC2626)
        "text": RGBColor(255, 255, 255),
    },
}

_DEFAULT_PALETTE = {
    "fill": RGBColor(71, 85, 105),           # Slate Gray (#475569)
    "text": RGBColor(255, 255, 255),
}


def _verify_pptx_file(path: Path) -> None:
    """Validates that path exists, is non-empty, and represents a readable PKZip archive."""
    if not path.exists() or not path.is_file():
        raise FileNotFoundError(f"Presentation file '{path}' does not exist.")
    if path.stat().st_size == 0:
        raise DeliveryError(f"Presentation file '{path.name}' is empty (0 bytes).")
    if not zipfile.is_zipfile(path):
        raise DeliveryError(f"Presentation file '{path.name}' is not a valid zip/OpenXML archive.")


def _strip_shape_shadows(shape: Any) -> None:
    """
    OpenXML Sanitizer that strips theme drop shadows and style effects.
    
    Guarantees flat consulting aesthetic with zero drop shadow inheritance:
    1. Sets shape.shadow.inherit = False (if attribute exists).
    2. Overrides <p:style>/<a:effectRef idx="0"> so theme drop shadows aren't inherited.
    3. Cleans any existing effectLst or outerShdw in spPr and maintains valid OpenXML schema ordering.
    """
    try:
        if hasattr(shape, "shadow"):
            shape.shadow.inherit = False
    except Exception:
        pass

    try:
        sp = getattr(shape, "_element", getattr(shape, "element", shape))
        if sp is None:
            return

        # 1. Reset theme style effectRef to idx="0"
        style = sp.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style")
        if style is not None:
            effect_ref = style.find("{http://schemas.openxmlformats.org/drawingml/2006/main}effectRef")
            if effect_ref is not None:
                effect_ref.set("idx", "0")

        # 2. Schema-compliant effectLst handling in spPr
        sp_pr = getattr(sp, "spPr", None)
        if sp_pr is not None:
            # Remove any standalone outerShdw tags directly inside spPr
            for child in list(sp_pr):
                if child.tag.endswith("outerShdw"):
                    sp_pr.remove(child)

            # Locate existing effectLst if present
            effect_lst = None
            for child in sp_pr:
                if child.tag.endswith("effectLst"):
                    effect_lst = child
                    break

            if effect_lst is not None:
                # Clear all children from effectLst to leave it empty without breaking XML sequence
                effect_lst.clear()
            else:
                # Insert empty effectLst in schema-valid sequence order:
                # Must precede effectDag, scene3d, sp3d, extLst
                new_effect_lst = OxmlElement("a:effectLst")
                inserted = False
                for idx, child in enumerate(list(sp_pr)):
                    tag = child.tag
                    if any(tag.endswith(t) for t in ("effectDag", "scene3d", "sp3d", "extLst")):
                        sp_pr.insert(idx, new_effect_lst)
                        inserted = True
                        break
                if not inserted:
                    sp_pr.append(new_effect_lst)
    except Exception:
        pass


class OverlayManager:
    """
    Headless status overlay badge controller for PowerPoint decks.
    """

    @staticmethod
    def _save_atomic(prs: Presentation, dest: Path) -> None:
        """Saves presentation using an atomic write-and-rename pattern to prevent file corruption."""
        dest.parent.mkdir(parents=True, exist_ok=True)
        temp_dest = dest.with_name(f".{dest.name}.tmp_{uuid.uuid4().hex[:8]}")
        try:
            prs.save(str(temp_dest))
            temp_dest.replace(dest)
        except Exception as exc:
            if temp_dest.exists():
                temp_dest.unlink(missing_ok=True)
            raise DeliveryError(f"Failed saving presentation to '{dest.name}': {exc}") from exc

    def stamp(
        self,
        pptx_path: Union[str, Path],
        status: Union[str, OverlayStatus],
        slides: Optional[List[int]] = None,
        output_path: Optional[Union[str, Path]] = None,
    ) -> OverlayReport:
        """
        Stamps a sharp status badge onto specified slides with zero drop shadows.
        
        Args:
          pptx_path: Input presentation file path.
          status: Deliverable status (e.g. 'DRAFT', 'IN_REVIEW', 'APPROVED', 'FINAL').
          slides: 1-indexed slide numbers to stamp. If None, stamps all slides.
          output_path: Destination path. Defaults to overwriting pptx_path.
          
        Returns:
          OverlayReport detailing the stamped slides.
        """
        src = Path(pptx_path).resolve()
        _verify_pptx_file(src)

        status_key = status.value if isinstance(status, OverlayStatus) else str(status).strip().upper()

        palette = _STATUS_PALETTE.get(status_key, _DEFAULT_PALETTE)
        fill_color = palette["fill"]
        text_color = palette["text"]

        try:
            prs = Presentation(str(src))
        except Exception as exc:
            raise DeliveryError(f"Failed opening presentation '{src.name}': {exc}") from exc

        total_slides = len(prs.slides)
        target_indices: Set[int]
        if slides is None:
            target_indices = set(range(total_slides))
        else:
            target_indices = set()
            for s in slides:
                try:
                    s_int = int(s)
                    if 1 <= s_int <= total_slides:
                        target_indices.add(s_int - 1)
                except (TypeError, ValueError):
                    continue

        badge_w = Inches(1.75)
        badge_h = Inches(0.35)
        badge_left = prs.slide_width - badge_w - Inches(0.40)
        badge_top = Inches(0.35)

        stamped_slides: List[int] = []

        for idx, slide in enumerate(prs.slides):
            if idx not in target_indices:
                continue

            # 1. Remove any pre-existing badge on this slide
            self._remove_badges_from_slide(slide)

            # 2. Add sharp rectangle container card (Zero Top Stripes on Rounded Cards invariant)
            shape = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                badge_left,
                badge_top,
                badge_w,
                badge_h,
            )
            shape.name = STATUS_BADGE_NAME

            # 3. Apply solid fill and border styling
            shape.fill.solid()
            shape.fill.fore_color.rgb = fill_color
            shape.line.color.rgb = fill_color
            shape.line.width = Pt(0.75)

            # 4. Enforce Zero Drop Shadows
            _strip_shape_shadows(shape)

            # 5. Populate status typography
            tf = shape.text_frame
            tf.word_wrap = False
            tf.margin_left = Inches(0.06)
            tf.margin_right = Inches(0.06)
            tf.margin_top = Inches(0.04)
            tf.margin_bottom = Inches(0.04)

            p = tf.paragraphs[0]
            p.text = status_key
            p.alignment = PP_ALIGN.CENTER
            p.font.name = "Arial"
            p.font.size = Pt(9.5)
            p.font.bold = True
            p.font.color.rgb = text_color

            stamped_slides.append(idx + 1)

        dest = Path(output_path).resolve() if output_path else src
        self._save_atomic(prs, dest)
        logger.info("Stamped %d slide(s) with '%s' overlay into '%s'", len(stamped_slides), status_key, dest.name)

        return self.inspect(dest)

    def strip(
        self,
        pptx_path: Union[str, Path],
        slides: Optional[List[int]] = None,
        output_path: Optional[Union[str, Path]] = None,
    ) -> OverlayReport:
        """
        Removes status overlay badges from specified slides.
        
        Args:
          pptx_path: Presentation file path.
          slides: 1-indexed slide numbers to strip. If None, strips all slides.
          output_path: Destination path. Defaults to overwriting pptx_path.
        """
        src = Path(pptx_path).resolve()
        _verify_pptx_file(src)

        try:
            prs = Presentation(str(src))
        except Exception as exc:
            raise DeliveryError(f"Failed opening presentation '{src.name}': {exc}") from exc

        total_slides = len(prs.slides)
        target_indices: Set[int]
        if slides is None:
            target_indices = set(range(total_slides))
        else:
            target_indices = set()
            for s in slides:
                try:
                    s_int = int(s)
                    if 1 <= s_int <= total_slides:
                        target_indices.add(s_int - 1)
                except (TypeError, ValueError):
                    continue

        stripped_slides: List[int] = []

        for idx, slide in enumerate(prs.slides):
            if idx not in target_indices:
                continue
            removed_count = self._remove_badges_from_slide(slide)
            if removed_count > 0:
                stripped_slides.append(idx + 1)

        dest = Path(output_path).resolve() if output_path else src
        self._save_atomic(prs, dest)
        logger.info("Stripped overlay badges from %d slide(s) in '%s'", len(stripped_slides), dest.name)

        report = self.inspect(dest)
        report.stripped_slides = stripped_slides
        return report

    def inspect(self, pptx_path: Union[str, Path]) -> OverlayReport:
        """
        Audits a presentation and returns slide-by-slide overlay badge telemetry.
        """
        src = Path(pptx_path).resolve()
        _verify_pptx_file(src)

        try:
            prs = Presentation(str(src))
        except Exception as exc:
            raise DeliveryError(f"Failed opening presentation '{src.name}': {exc}") from exc

        badge_infos: List[SlideBadgeInfo] = []
        stamped_slide_nums: List[int] = []
        detected_statuses: List[str] = []

        for idx, slide in enumerate(prs.slides):
            slide_num = idx + 1
            badge_found = False
            for shape in slide.shapes:
                name = getattr(shape, "name", "")
                if name == STATUS_BADGE_NAME or name.startswith(f"{STATUS_BADGE_NAME}_"):
                    badge_found = True
                    text = ""
                    try:
                        if getattr(shape, "has_text_frame", False):
                            text = shape.text_frame.text.strip()
                    except Exception:
                        text = ""
                    shape_id = getattr(shape, "shape_id", None)
                    badge_infos.append(
                        SlideBadgeInfo(
                            slide_number=slide_num,
                            has_badge=True,
                            status=text,
                            text=text,
                            shape_id=shape_id,
                        )
                    )
                    stamped_slide_nums.append(slide_num)
                    if text:
                        detected_statuses.append(text)
                    break

            if not badge_found:
                badge_infos.append(
                    SlideBadgeInfo(slide_number=slide_num, has_badge=False)
                )

        deck_status: Optional[OverlayStatus] = None
        if detected_statuses:
            primary_status_str = detected_statuses[0].upper()
            if primary_status_str in OverlayStatus._value2member_map_:
                deck_status = OverlayStatus(primary_status_str)

        return OverlayReport(
            pptx_path=str(src),
            status=deck_status,
            stamped_slides=stamped_slide_nums,
            stripped_slides=[],
            slide_badges=badge_infos,
            total_slides=len(prs.slides),
        )

    def _remove_badges_from_slide(self, slide: Any) -> int:
        """Removes all status overlay badge shapes from a single slide."""
        removed = 0
        for shape in list(slide.shapes):
            name = getattr(shape, "name", "")
            if name == STATUS_BADGE_NAME or name.startswith(f"{STATUS_BADGE_NAME}_"):
                try:
                    sp = getattr(shape, "_element", None)
                    if sp is not None:
                        parent = sp.getparent()
                        if parent is not None:
                            parent.remove(sp)
                            removed += 1
                except Exception as exc:
                    logger.warning("Failed removing badge shape from slide: %s", exc)
        return removed
