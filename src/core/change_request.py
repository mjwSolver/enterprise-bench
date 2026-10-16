"""
Change Request (CR) Processing Subsystem
=========================================
Automates the intake, impact analysis modeling, document stamping, and ledger
registration for enterprise Change Requests (CRs).
"""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Optional
import docx
import openpyxl
from pydantic import BaseModel, Field

from src.core.slug_registry import (
    DEFAULT_CLIENT_SHORT_NAME,
    EngagementContext,
    substitute_slugs_in_document,
    substitute_slugs_in_workbook,
)
from src.core.workspace import WorkspaceRouter


def _substitute_docx(doc: Any, context: EngagementContext) -> None:
    """Substitutes legacy slugs and placeholders across paragraphs, tables, headers, and footers."""
    substitute_slugs_in_document(doc, context)


def _substitute_xlsx(wb: Any, context: EngagementContext) -> None:
    """Substitutes legacy slugs across all sheets and cells in an openpyxl workbook."""
    substitute_slugs_in_workbook(wb, context)


def _set_cell_text(table: Any, row_idx: int, col_idx: int, text: str) -> None:
    """Safely set text on a table cell if row and column indices are within bounds."""
    if 0 <= row_idx < len(table.rows):
        row = table.rows[row_idx]
        if 0 <= col_idx < len(row.cells):
            row.cells[col_idx].text = text


class CRSubmission(BaseModel):
    """Payload representing an incoming enterprise Change Request."""
    cr_id: Optional[str] = Field(None, description="Identifier (e.g. '7' or 'CR-07'). Auto-increments if omitted.")
    title: str = Field(..., description="Concise title of the change request")
    category: str = Field("Scope", description="Scope, Schedule, Requirement, Design, Technical, Bug Fix")
    priority: str = Field("High", description="Low, Medium, High, Critical")
    requester: str = Field(..., description="Name and title of requester")
    department: str = Field("Consolidated Accounting", description="Department / Role")
    submission_date: str = Field(default_factory=lambda: datetime.now().strftime("%d %B %Y"))
    description: str = Field(..., description="Detailed functional and business requirement")
    reason: str = Field("Business user operational requirement", description="Justification / Business driver")
    scope_impact: str = Field(..., description="Technical & architectural components affected")
    schedule_days: int = Field(15, description="Estimated calendar days required")
    buffer_days: int = Field(2, description="Buffer days for regression testing")
    cost_estimate: str = Field("To be quoted via commercial addendum", description="Cost / Billing estimate")
    mandays_breakdown: Dict[str, int] = Field(
        default_factory=lambda: {
            "analysis_pm": 4,
            "development": 10,
            "testing": 6,
            "deployment": 2,
        },
        description="Effort breakdown by role/phase",
    )
    risk_impact: str = Field("Low - isolated feature addition, minimal regression risk", description="Risk assessment")
    quality_impact: str = Field("Positive - enhances data accuracy and user reporting agility", description="Quality impact")


class CRProcessingResult(BaseModel):
    """Outcome of processing a Change Request submission."""
    cr_id: str
    project_id: str
    cr_form_path: str
    ledger_updated: bool
    mandays_sheet_path: Optional[str] = None
    bast_draft_path: Optional[str] = None
    total_mandays: int
    summary: str


class ChangeRequestProcessor:
    """Orchestrates document generation, ledger logging, and scoping for Change Requests."""

    def __init__(
        self,
        project_id: str = "TTI_Snowflake_Analytics",
        context: Optional[EngagementContext] = None,
    ):
        self.context = context or EngagementContext.default_ngl()
        self.project_id = project_id
        self.project = WorkspaceRouter.get_project(project_id)

    def _get_next_cr_id(self) -> str:
        """Inspect the current ledger to determine the next sequential CR ID."""
        ledger_path = self.project.resolve_clean_file("4.4_Change_Log_Ledger_Template.xlsx")
        if not ledger_path or not ledger_path.exists():
            return "7"

        wb = openpyxl.load_workbook(str(ledger_path), data_only=True, read_only=True)
        try:
            ws = wb["Change Log"] if "Change Log" in wb.sheetnames else wb.active
            max_id = 0
            for row in ws.iter_rows(min_row=3, max_col=1, values_only=True):
                first_cell = row[0] if row else None
                if first_cell is not None:
                    try:
                        val = int(str(first_cell).strip().replace("CR-", "").replace("CR", ""))
                        if val > max_id:
                            max_id = val
                    except ValueError:
                        pass
        finally:
            wb.close()
        return str(max_id + 1 if max_id > 0 else 7)

    def file_change_request(
        self,
        submission: CRSubmission,
        output_dir: Optional[Path | str] = None,
    ) -> CRProcessingResult:
        """
        Process a complete Change Request:
        1. Formats and stamps Change_Request_Form.docx
        2. Appends registration row to Change_Log_Ledger.xlsx
        3. Generates CR_Scoping_and_Mandays.xlsx
        4. Prepares draft BAST_Change_Request.docx
        """
        raw_id = (submission.cr_id or self._get_next_cr_id()).strip()
        cr_id = re.sub(r"^CR[-_\s]*", "", raw_id, flags=re.IGNORECASE) or raw_id
        if output_dir:
            out_p = Path(output_dir)
        else:
            client_prefix = self.context.client_short_name if self.context else DEFAULT_CLIENT_SHORT_NAME
            out_p = Path("output") / f"{client_prefix}_Snowflake_Analytics" / "change_requests" / f"CR_{cr_id.zfill(2)}"
        out_p.mkdir(parents=True, exist_ok=True)

        total_mandays = sum(submission.mandays_breakdown.values())

        # 1. Generate Change_Request_Form.docx
        form_template = self.project.resolve_clean_file("4.4_Change_Request_Form_Template.docx")
        if not form_template:
            raise FileNotFoundError("4.4_Change_Request_Form_Template.docx not found in clean workspace.")

        doc = docx.Document(str(form_template))

        # Table 0: Metadata
        if len(doc.tables) > 0:
            t0 = doc.tables[0]
            _set_cell_text(t0, 0, 2, cr_id)
            _set_cell_text(t0, 1, 2, submission.submission_date)
            _set_cell_text(t0, 2, 2, submission.requester)
            _set_cell_text(t0, 3, 2, submission.department)

        # Table 1: Change Details
        if len(doc.tables) > 1:
            t1 = doc.tables[1]
            _set_cell_text(t1, 0, 2, submission.title)
            _set_cell_text(t1, 1, 2, submission.category)
            _set_cell_text(t1, 2, 2, submission.priority)
            _set_cell_text(t1, 3, 2, submission.description)
            _set_cell_text(t1, 4, 2, submission.reason)

        # Table 2: Impact Analysis
        if len(doc.tables) > 2:
            t2 = doc.tables[2]
            _set_cell_text(t2, 0, 2, submission.scope_impact)
            _set_cell_text(t2, 1, 2, f"Estimasi waktu pelaksanaan: {submission.schedule_days} hari + {submission.buffer_days} hari buffer.")
            _set_cell_text(t2, 2, 2, submission.cost_estimate)
            _set_cell_text(t2, 3, 2, f"Total estimasi kebutuhan sumber daya: {total_mandays} mandays.")
            _set_cell_text(t2, 4, 2, submission.risk_impact)
            _set_cell_text(t2, 5, 2, submission.quality_impact)

        # Table 3: Effort Breakdown
        if len(doc.tables) > 3:
            t3 = doc.tables[3]
            _set_cell_text(t3, 0, 2, f"{submission.mandays_breakdown.get('analysis_pm', 0)} mandays")
            _set_cell_text(t3, 1, 2, f"{submission.mandays_breakdown.get('development', 0)} mandays")
            _set_cell_text(t3, 2, 2, f"{submission.mandays_breakdown.get('testing', 0)} mandays")
            _set_cell_text(t3, 3, 2, f"{submission.mandays_breakdown.get('deployment', 0)} mandays")
            _set_cell_text(t3, 4, 2, f"{total_mandays} mandays")

        doc_out_path = out_p / f"Change_Request_Form_CR_{cr_id.zfill(2)}.docx"
        _substitute_docx(doc, self.context)
        doc.save(str(doc_out_path))

        # Enforce zero-distortion aspect ratio preservation on embedded images
        from src.core.image_aspect import ImageAspectEngine
        ImageAspectEngine().fix_docx(doc_out_path)

        # 2. Append to Change_Log_Ledger.xlsx
        ledger_template = self.project.resolve_clean_file("4.4_Change_Log_Ledger_Template.xlsx")
        ledger_out_path = out_p / f"Change_Log_Ledger_Updated_CR_{cr_id.zfill(2)}.xlsx"
        ledger_updated = False

        if ledger_template and ledger_template.exists():
            wb = openpyxl.load_workbook(str(ledger_template))
            try:
                ws = wb["Change Log"] if "Change Log" in wb.sheetnames else wb.active
                new_row = [
                    int(cr_id) if cr_id.isdigit() else cr_id,
                    datetime.now().strftime("%Y-%m-%d"),
                    submission.title,
                    submission.description[:250] + ("..." if len(submission.description) > 250 else ""),
                    submission.category,
                    "Business Request",
                    submission.requester,
                    submission.reason,
                    submission.scope_impact[:120],
                    f"{submission.schedule_days} hari",
                    submission.cost_estimate,
                    f"{total_mandays} mandays",
                    "Impact Analysis Completed",
                    "Under Review",
                ]
                ws.append(new_row)
                _substitute_xlsx(wb, self.context)
                wb.save(str(ledger_out_path))
                ledger_updated = True
            finally:
                wb.close()

        # 3. Generate CR_Scoping_and_Mandays.xlsx
        scoping_template = self.project.resolve_clean_file("CR_Scoping_and_Mandays_Template.xlsx")
        scoping_out_path = out_p / f"CR_Scoping_and_Mandays_CR_{cr_id.zfill(2)}.xlsx"
        scoping_created = False
        if scoping_template and scoping_template.exists():
            wb_sc = openpyxl.load_workbook(str(scoping_template))
            try:
                ws_gen = wb_sc["General"] if "General" in wb_sc.sheetnames else wb_sc.active
                ws_gen.cell(row=1, column=2, value=f"CR-{cr_id}: {submission.title}")
                ws_gen.cell(row=2, column=2, value=f"Total Mandays: {total_mandays}")
                ws_gen.cell(row=3, column=2, value=submission.description[:150])

                # Create specific detail sheet for this CR
                sheet_title = re.sub(r'[:\\/?*\[\]]', '_', f"CR{cr_id} Breakdown")[:31]
                ws_cr = wb_sc.create_sheet(title=sheet_title)
                ws_cr.append(["Activity / Workstream", "Role / Team", "Estimated Mandays"])
                ws_cr.append(["Analysis, Design & Governance", "Project Manager / BA", submission.mandays_breakdown.get("analysis_pm", 0)])
                ws_cr.append(["Snowflake Engineering & UI Dev", "Data & Streamlit Devs", submission.mandays_breakdown.get("development", 0)])
                ws_cr.append(["Integration & Regression Testing", "QA Engineer", submission.mandays_breakdown.get("testing", 0)])
                ws_cr.append(["Deployment & Operational Handover", "DevOps Engineer", submission.mandays_breakdown.get("deployment", 0)])
                ws_cr.append(["Total Committed Effort", "All Teams", total_mandays])

                _substitute_xlsx(wb_sc, self.context)
                wb_sc.save(str(scoping_out_path))
                scoping_created = True
            finally:
                wb_sc.close()

        # 4. Generate draft BAST_Change_Request.docx
        bast_template = self.project.resolve_clean_file("BAST_Change_Request_Template.docx")
        bast_out_path = out_p / f"BAST_Change_Request_Draft_CR_{cr_id.zfill(2)}.docx"
        bast_created = False
        if bast_template and bast_template.exists():
            doc_bast = docx.Document(str(bast_template))
            # Insert CR title in paragraph 2
            if len(doc_bast.paragraphs) > 1:
                p = doc_bast.paragraphs[1]
                p.text = f"Pekerjaan : Jasa Implementasi Snowflake - Change Request #{cr_id} ({submission.title})"
            _substitute_docx(doc_bast, self.context)
            doc_bast.save(str(bast_out_path))
            bast_created = True

        summary = (
            f"Successfully filed Change Request #{cr_id}: '{submission.title}'\n"
            f"- Form: {doc_out_path}\n"
            f"- Ledger: {ledger_out_path if ledger_updated else 'N/A'}\n"
            f"- Scoping: {scoping_out_path if scoping_created else 'N/A'}\n"
            f"- BAST Draft: {bast_out_path if bast_created else 'N/A'}\n"
            f"- Total Effort: {total_mandays} mandays across {submission.schedule_days} calendar days."
        )

        return CRProcessingResult(
            cr_id=cr_id,
            project_id=self.project_id,
            cr_form_path=str(doc_out_path),
            ledger_updated=ledger_updated,
            mandays_sheet_path=str(scoping_out_path) if scoping_created else None,
            bast_draft_path=str(bast_out_path) if bast_created else None,
            total_mandays=total_mandays,
            summary=summary,
        )
