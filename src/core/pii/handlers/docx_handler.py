"""
Word Document (.docx) Format Handler
====================================
Extracts and sanitizes text and metadata from Microsoft Word files, inspecting
body paragraphs, nested tables, headers, footers, and OOXML core properties.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Dict, List, Set
from src.core.pii.handlers.base import BaseFormatHandler
from src.core.pii.models import SanitizeResult, TextNode

try:
    import docx
    from docx.document import Document as _Document
except ImportError:
    docx = None  # type: ignore


class DocxHandler(BaseFormatHandler):
    """Handler for Microsoft Word .docx files."""

    @classmethod
    def supported_extensions(cls) -> Set[str]:
        return {".docx"}

    def _ensure_lib(self):
        if docx is None:
            raise ImportError("python-docx is required for Word document handling. Install with 'uv sync --extra docx'")

    def extract_metadata(self, path: Path) -> Dict[str, str]:
        self._ensure_lib()
        doc = docx.Document(str(path))
        props = {}
        cp = doc.core_properties
        for attr in ["author", "last_modified_by", "comments", "title", "subject", "keywords"]:
            val = getattr(cp, attr, None)
            if val and str(val).strip():
                props[attr] = str(val).strip()
        return props

    def extract_text_nodes(self, path: Path) -> List[TextNode]:
        self._ensure_lib()
        doc = docx.Document(str(path))
        nodes: List[TextNode] = []

        # 1. Body Paragraphs
        for idx, p in enumerate(doc.paragraphs):
            text = p.text.strip()
            if text:
                nodes.append(
                    TextNode(
                        location=f"Body > Paragraph {idx + 1}",
                        text=text,
                        node_type="paragraph",
                        context={"para_index": idx},
                    )
                )

        # 2. Body Tables (handles nested tables recursively)
        def _extract_table(tbl, prefix: str):
            for r_idx, row in enumerate(tbl.rows):
                for c_idx, cell in enumerate(row.cells):
                    loc = f"{prefix} > Row {r_idx + 1} Col {c_idx + 1}"
                    for p_idx, p in enumerate(cell.paragraphs):
                        t = p.text.strip()
                        if t:
                            nodes.append(
                                TextNode(
                                    location=f"{loc} > Para {p_idx + 1}",
                                    text=t,
                                    node_type="table_cell",
                                    context={"row": r_idx, "col": c_idx},
                                )
                            )
                    for n_idx, nested_tbl in enumerate(cell.tables):
                        _extract_table(nested_tbl, f"{loc} > NestedTable {n_idx + 1}")

        for t_idx, tbl in enumerate(doc.tables):
            _extract_table(tbl, f"Body > Table {t_idx + 1}")

        # 3. Headers and Footers
        for s_idx, sec in enumerate(doc.sections):
            for p_idx, p in enumerate(sec.header.paragraphs):
                t = p.text.strip()
                if t:
                    nodes.append(
                        TextNode(
                            location=f"Section {s_idx + 1} > Header > Para {p_idx + 1}",
                            text=t,
                            node_type="header",
                            context={"section": s_idx},
                        )
                    )
            for t_idx, tbl in enumerate(sec.header.tables):
                _extract_table(tbl, f"Section {s_idx + 1} > Header > Table {t_idx + 1}")

            for p_idx, p in enumerate(sec.footer.paragraphs):
                t = p.text.strip()
                if t:
                    nodes.append(
                        TextNode(
                            location=f"Section {s_idx + 1} > Footer > Para {p_idx + 1}",
                            text=t,
                            node_type="footer",
                            context={"section": s_idx},
                        )
                    )
            for t_idx, tbl in enumerate(sec.footer.tables):
                _extract_table(tbl, f"Section {s_idx + 1} > Footer > Table {t_idx + 1}")

        # 4. Metadata as TextNode
        meta = self.extract_metadata(path)
        for k, v in meta.items():
            nodes.append(
                TextNode(
                    location=f"Metadata > {k}",
                    text=v,
                    node_type="core_property",
                    context={"property": k},
                )
            )

        return nodes

    def apply_replacements(
        self,
        input_path: Path,
        output_path: Path,
        replacements: Dict[str, str],
        strip_metadata: bool = True,
        metadata_overrides: Dict[str, str] | None = None,
    ) -> SanitizeResult:
        self._ensure_lib()
        doc = docx.Document(str(input_path))
        repl_count = 0
        details: List[str] = []

        # Sort replacement targets by length descending so longer phrases match before substrings
        sorted_replacements = sorted(replacements.items(), key=lambda x: len(x[0]), reverse=True)

        def _clean_paragraph(p, location: str) -> int:
            nonlocal repl_count
            local_applied = 0
            if not p.text:
                return 0

            # First pass: try run-by-run replacement to preserve character formatting
            for run in p.runs:
                if not run.text:
                    continue
                orig = run.text
                new_val = orig
                for target, replacement in sorted_replacements:
                    if target in new_val:
                        new_val = re.sub(re.escape(target), replacement, new_val, flags=re.IGNORECASE)
                if new_val != orig:
                    run.text = new_val
                    local_applied += 1

            # Second pass: check if target spans across run boundaries
            p_text = p.text
            for target, replacement in sorted_replacements:
                if re.search(re.escape(target), p_text, flags=re.IGNORECASE):
                    # Replace in entire paragraph text
                    p.text = re.sub(re.escape(target), replacement, p_text, flags=re.IGNORECASE)
                    local_applied += 1
                    p_text = p.text

            if local_applied:
                details.append(f"Replaced {local_applied} entity(ies) in {location}")
            return local_applied

        # 1. Clean paragraphs
        for idx, p in enumerate(doc.paragraphs):
            repl_count += _clean_paragraph(p, f"Body > Paragraph {idx + 1}")

        # 2. Clean tables
        def _clean_table(tbl, prefix: str):
            nonlocal repl_count
            for r_idx, row in enumerate(tbl.rows):
                for c_idx, cell in enumerate(row.cells):
                    loc = f"{prefix} > Row {r_idx + 1} Col {c_idx + 1}"
                    for p_idx, p in enumerate(cell.paragraphs):
                        repl_count += _clean_paragraph(p, f"{loc} > Para {p_idx + 1}")
                    for n_idx, nested in enumerate(cell.tables):
                        _clean_table(nested, f"{loc} > NestedTable {n_idx + 1}")

        for t_idx, tbl in enumerate(doc.tables):
            _clean_table(tbl, f"Body > Table {t_idx + 1}")

        # 3. Clean headers and footers
        for s_idx, sec in enumerate(doc.sections):
            for p_idx, p in enumerate(sec.header.paragraphs):
                repl_count += _clean_paragraph(p, f"Section {s_idx + 1} > Header > Para {p_idx + 1}")
            for tbl in sec.header.tables:
                _clean_table(tbl, f"Section {s_idx + 1} > Header > Table")

            for p_idx, p in enumerate(sec.footer.paragraphs):
                repl_count += _clean_paragraph(p, f"Section {s_idx + 1} > Footer > Para {p_idx + 1}")
            for tbl in sec.footer.tables:
                _clean_table(tbl, f"Section {s_idx + 1} > Footer > Table")

        # 4. Scrub or replace core properties
        if strip_metadata:
            overrides = metadata_overrides or {
                "author": "Enterprise Contributor",
                "last_modified_by": "Enterprise Contributor",
                "comments": "",
                "title": "",
                "subject": "",
            }
            cp = doc.core_properties
            for prop_name, default_val in overrides.items():
                if hasattr(cp, prop_name):
                    setattr(cp, prop_name, default_val)
            details.append("Scrubbed document core properties (author, last_modified_by, comments)")

        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))

        return SanitizeResult(
            source_path=str(input_path),
            output_path=str(output_path),
            replacements_applied=repl_count,
            metadata_scrubbed=strip_metadata,
            details=details,
        )
