"""
PowerPoint (.pptx) Format Handler
=================================
Extracts and sanitizes text and metadata from Microsoft PowerPoint decks,
inspecting shapes, text boxes, tables, presenter notes, grouped shapes, and OOXML core properties.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Dict, List, Set
from src.core.pii.handlers.base import BaseFormatHandler
from src.core.pii.models import SanitizeResult, TextNode

try:
    import pptx
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError:
    pptx = None  # type: ignore


class PptxHandler(BaseFormatHandler):
    """Handler for Microsoft PowerPoint .pptx presentations."""

    @classmethod
    def supported_extensions(cls) -> Set[str]:
        return {".pptx"}

    def _ensure_lib(self):
        if pptx is None:
            raise ImportError("python-pptx is required for PowerPoint handling. Install with 'uv sync --extra ppt'")

    def extract_metadata(self, path: Path) -> Dict[str, str]:
        self._ensure_lib()
        prs = pptx.Presentation(str(path))
        props = {}
        cp = prs.core_properties
        for attr in ["author", "last_modified_by", "comments", "title", "subject", "keywords"]:
            val = getattr(cp, attr, None)
            if val and str(val).strip():
                props[attr] = str(val).strip()
        return props

    def extract_text_nodes(self, path: Path) -> List[TextNode]:
        self._ensure_lib()
        prs = pptx.Presentation(str(path))
        nodes: List[TextNode] = []

        def _extract_shape_text(shape, prefix: str):
            # Check standard text frame
            if shape.has_text_frame:
                for p_idx, p in enumerate(shape.text_frame.paragraphs):
                    t = p.text.strip()
                    if t:
                        nodes.append(
                            TextNode(
                                location=f"{prefix} > Para {p_idx + 1}",
                                text=t,
                                node_type="slide_shape",
                                context={"shape_name": shape.name},
                            )
                        )

            # Check table
            if shape.has_table:
                tbl = shape.table
                for r_idx, row in enumerate(tbl.rows):
                    for c_idx, cell in enumerate(row.cells):
                        cell_loc = f"{prefix} > Table Row {r_idx + 1} Col {c_idx + 1}"
                        for p_idx, p in enumerate(cell.text_frame.paragraphs):
                            t = p.text.strip()
                            if t:
                                nodes.append(
                                    TextNode(
                                        location=f"{cell_loc} > Para {p_idx + 1}",
                                        text=t,
                                        node_type="table_cell",
                                        context={"shape_name": shape.name, "row": r_idx, "col": c_idx},
                                    )
                                )

            # Check grouped shapes recursively
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                for sub_idx, sub_shape in enumerate(shape.shapes):
                    _extract_shape_text(sub_shape, f"{prefix} > GroupItem {sub_idx + 1} ({sub_shape.name})")

        for s_idx, slide in enumerate(prs.slides):
            slide_prefix = f"Slide {s_idx + 1}"

            # Extract slide shapes
            for shp_idx, shape in enumerate(slide.shapes):
                _extract_shape_text(shape, f"{slide_prefix} > Shape {shp_idx + 1} ({shape.name})")

            # Extract presenter notes
            if slide.has_notes_slide:
                notes_tf = slide.notes_slide.notes_text_frame
                for p_idx, p in enumerate(notes_tf.paragraphs):
                    t = p.text.strip()
                    if t:
                        nodes.append(
                            TextNode(
                                location=f"{slide_prefix} > Notes > Para {p_idx + 1}",
                                text=t,
                                node_type="notes",
                                context={"slide_index": s_idx},
                            )
                        )

        # Document core properties
        meta = self.extract_metadata(path)
        for k, v in meta.items():
            nodes.append(
                TextNode(
                    location=f"Metadata > {k}",
                    text=v,
                    node_type="core_property",
                    context={"property": k},
                )
            )

        return nodes

    def apply_replacements(
        self,
        input_path: Path,
        output_path: Path,
        replacements: Dict[str, str],
        strip_metadata: bool = True,
        metadata_overrides: Dict[str, str] | None = None,
    ) -> SanitizeResult:
        self._ensure_lib()
        prs = pptx.Presentation(str(input_path))
        repl_count = 0
        details: List[str] = []

        sorted_replacements = sorted(replacements.items(), key=lambda x: len(x[0]), reverse=True)

        def _clean_paragraph(p, location: str) -> int:
            nonlocal repl_count
            local_applied = 0
            if not p.text:
                return 0

            # Run-level replacement
            for run in p.runs:
                if not run.text:
                    continue
                orig = run.text
                new_val = orig
                for target, replacement in sorted_replacements:
                    if target.lower() in new_val.lower():
                        new_val = re.sub(re.escape(target), replacement, new_val, flags=re.IGNORECASE)
                if new_val != orig:
                    run.text = new_val
                    local_applied += 1

            # Paragraph-level fallback if entity was fragmented across runs
            p_text = p.text
            for target, replacement in sorted_replacements:
                if re.search(re.escape(target), p_text, flags=re.IGNORECASE):
                    p.text = re.sub(re.escape(target), replacement, p_text, flags=re.IGNORECASE)
                    local_applied += 1
                    p_text = p.text

            if local_applied:
                details.append(f"Replaced {local_applied} entity(ies) in {location}")
            return local_applied

        def _clean_shape(shape, prefix: str):
            nonlocal repl_count
            if shape.has_text_frame:
                for p_idx, p in enumerate(shape.text_frame.paragraphs):
                    repl_count += _clean_paragraph(p, f"{prefix} > Para {p_idx + 1}")

            if shape.has_table:
                for r_idx, row in enumerate(shape.table.rows):
                    for c_idx, cell in enumerate(row.cells):
                        for p_idx, p in enumerate(cell.text_frame.paragraphs):
                            repl_count += _clean_paragraph(p, f"{prefix} > Table R{r_idx + 1}C{c_idx + 1} > Para {p_idx + 1}")

            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                for sub_idx, sub_shape in enumerate(shape.shapes):
                    _clean_shape(sub_shape, f"{prefix} > GroupItem {sub_idx + 1}")

        # Clean all slides
        for s_idx, slide in enumerate(prs.slides):
            slide_prefix = f"Slide {s_idx + 1}"
            for shp_idx, shape in enumerate(slide.shapes):
                _clean_shape(shape, f"{slide_prefix} > Shape {shp_idx + 1}")

            if slide.has_notes_slide:
                for p_idx, p in enumerate(slide.notes_slide.notes_text_frame.paragraphs):
                    repl_count += _clean_paragraph(p, f"{slide_prefix} > Notes > Para {p_idx + 1}")

        # Clean metadata
        if strip_metadata:
            overrides = metadata_overrides or {
                "author": "Enterprise Contributor",
                "last_modified_by": "Enterprise Contributor",
                "comments": "",
                "title": "",
                "subject": "",
            }
            cp = prs.core_properties
            for prop_name, default_val in overrides.items():
                if hasattr(cp, prop_name):
                    setattr(cp, prop_name, default_val)
            details.append("Scrubbed presentation core properties (author, last_modified_by, comments)")

        output_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(output_path))

        return SanitizeResult(
            source_path=str(input_path),
            output_path=str(output_path),
            replacements_applied=repl_count,
            metadata_scrubbed=strip_metadata,
            details=details,
        )
