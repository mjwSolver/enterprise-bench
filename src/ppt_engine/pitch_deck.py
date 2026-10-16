"""
Presales Modernization Pitch Deck Builder Subsystem
===================================================
Automates and modernizes the 36-slide Presales Pitch Deck
(`Modernize_Data_Platform_Pitch_Deck_Template.pptx`) for enterprise cloud modernizations,
conforming to consulting frameworks, 60-30-10 palette distribution, and geometric rules:

  Slide 01: Master Cover (Cinematic Hero Cover / Clean Typographic Columns)
  Slide 02: Chapter Divider (Our Company Profile)
  Slide 03: Metrodata Corporate Structure (Hierarchical Equity Tree)
  Slide 04: Metrodata 8 Pillars Offering (4x2 Capability Card Matrix)
  Slide 05: Chapter Divider (Our Data & AI Portfolio)
  Slide 06: Product Portfolio (Normalized Tech Logo Grid)
  Slide 07: Modern Data Stack & Orchestration (Pipeline Ingestion Grid)
  Slide 08: Competencies & Certifications (Tiered Badge Matrix)
  Slide 09: Chapter Divider (Our Customer References)
  Slide 10: Enterprise Client References (Multi-Sector Credential Grid)
  Slide 11: Data Science & AI Portfolio (Capability Card Grid)
  Slide 12: Chapter Divider (Data Analytics Journey)
  Slide 13: Analytics Journey – Start with KPI (Split-Hero Inquiry Cards)
  Slide 14: The Iceberg Concept (Visible Interface vs Subsurface Foundation)
  Slide 15: Chapter Divider (Product Introduction & Objective)
  Slide 16: Snowflake: The AI Data Cloud (3-Tier Decoupled Architecture)
  Slide 17: Snowflake Platform Architecture (Official Architecture Blueprint)
  Slide 18: Business Applications & End-to-End Pipeline (Official Solution Flow)
  Slide 19: Chapter Divider (Existing Challenge & Proposed Solution)
  Slide 20: Existing Challenges (SAP ERP Window Mockup + Observation Cards)
  Slide 21: Gap Analysis & Solution (As-Is vs To-Be Comparative Matrix)
  Slide 22: Proposed Solution (4-Stage Value Chain Chevrons)
  Slide 23: Proposed Target Architecture (Architecture Blueprint Container)
  Slide 24: Analytics & Serving Architecture (Streamlit Window Mockup)
  Slide 25: Chapter Divider (Scope of Work)
  Slide 26: Project Delivery Lifecycle (Initiation to Handover Stepped Cards)
  Slide 27: Scope of Work Workstreams (4-Column WBS Structure)
  Slide 28: Assumptions – Sales & Gross Profit (Streamlit Browser Container)
  Slide 29: Assumptions – Opex & Margin (Streamlit Financial Container)
  Slide 30: Assumptions Summary (Strategic Scope Boundaries)
  Slide 31: Chapter Divider (Project Timeline & Organization)
  Slide 32: Project Timeline Estimate (Delivery Gantt Archetype)
  Slide 33: Project Organization Structure (Hierarchical RACI Org Tree)
  Slide 34: Chapter Divider (Change Request Procedure)
  Slide 35: Change Request Procedure (Governance Pipeline & Decision Gates)
  Slide 36: Executive Closing & Contact (Corporate Emblem & Contacts)

Strictly enforces:
  - Zero overlapping top lines on rounded cards (sharp MSO_SHAPE.RECTANGLE).
  - Unified title and subtitle frames (space_before = Pt(10)).
  - Clean typographic metadata on cover (zero boxed containers, no cover footer).
  - Synchronized bottom pagination ('02 / 36' to '36 / 36').
"""

from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, Union

import yaml
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from src.core.slug_registry import EngagementContext, substitute_slugs_in_presentation
from src.ppt_engine.consulting_archetypes import (
    AgendaItem,
    BadgeMatrixSection,
    CardGridItem,
    ConsultingDeckBuilder,
    EquityEntityNode,
    EquityTreeData,
    FeatureMatrixData,
    GanttTask,
    GanttTimelineData,
    GanttWorkstream,
    HorizonColumnData,
    ProcessChevronStep,
    ScorecardMetric,
    ScorecardQuadrantData,
    SplitHeroBlock,
    StrategyPillarData,
    TableColumnDef,
    TechLogoItem,
    add_card,
    add_card_with_top_stripe,
    add_slide_footer,
    add_slide_header,
    add_slide_with_background,
    build_agenda_slide,
    build_badge_matrix_slide,
    build_bcg_3_horizon_slide,
    build_browser_mockup_slide,
    build_card_grid_slide,
    build_chapter_divider_slide,
    build_chevron_process_slide,
    build_cover_slide,
    build_equity_corporate_tree_slide,
    build_feature_matrix_slide,
    build_gap_analysis_slide,
    build_hero_cover_slide,
    build_iceberg_concept_slide,
    build_mckinsey_cascade_slide,
    build_split_hero_slide,
    build_table_slide,
    build_tech_logo_grid_slide,
    build_thank_you_slide,
    build_timeline_gantt_slide,
    create_presentation,
)
from src.ppt_engine.reference_slides import (
    build_change_request_procedure_slide,
    build_governance_org_structure_slide,
    build_snowflake_data_pipeline_slide,
    build_snowflake_platform_architecture_slide,
)
from src.ppt_engine.theme_engine import Theme, get_theme, hex_to_rgb

logger = logging.getLogger(__name__)

# Precompiled regex for footer slide pagination ("02 / 36")
PAGE_REGEX = re.compile(r"^\d{2}\s*/\s*\d{2}$")


class PitchDeckBuilder:
    """Specialized presentation generator for the 36-slide Presales Modernization Pitch Deck."""

    def __init__(
        self,
        theme: Union[str, Theme] = "metrodata",
        config_dict: Optional[Dict[str, Any]] = None,
    ) -> None:
        if isinstance(theme, str):
            self.theme = get_theme(theme)
        else:
            self.theme = theme

        self.prs: Presentation = create_presentation(self.theme)
        self.config: Dict[str, Any] = config_dict or {}
        self.metadata: Dict[str, Any] = self.config.get("metadata") or {}
        self.client_name: str = self.metadata.get("client_name", "[CLIENT_COMPANY_NAME]")
        self.vendor_name: str = self.metadata.get("vendor_name", "PT Metrodata Electronics Tbk")
        self.product: str = self.metadata.get("product", "Snowflake AI Data Cloud")
        self.date_str: str = self.metadata.get("date_str", "September 2026")
        self.notice: str = self.metadata.get("confidentiality", "Enterprise Strategy Group  |  Confidential & Proprietary")

    @classmethod
    def from_yaml(
        cls,
        yaml_path: Union[str, Path],
        theme_override: Optional[str] = None,
    ) -> "PitchDeckBuilder":
        """Instantiates PitchDeckBuilder from a declarative YAML deck configuration."""
        p = Path(yaml_path)
        with open(p, "r", encoding="utf-8") as f:
            data = yaml.safe_load(f)

        theme_name = theme_override or (data.get("theme") if isinstance(data, dict) else "metrodata") or "metrodata"
        builder = cls(theme=theme_name, config_dict=data)
        return builder

    def build_all(self, slide_limit: Optional[int] = None) -> None:
        """Iterates through all declared slides in config and compiles them into the presentation."""
        slides_cfg = self.config.get("slides") or []
        if slide_limit is not None and slide_limit > 0:
            slides_cfg = slides_cfg[:slide_limit]

        for s_idx, s_cfg in enumerate(slides_cfg, start=1):
            archetype = s_cfg.get("archetype", "")
            self._dispatch_slide(archetype, s_cfg, s_idx, len(slides_cfg))

    def _dispatch_slide(
        self,
        archetype: str,
        cfg: Dict[str, Any],
        idx: int,
        total: int,
    ) -> None:
        """Dispatches configuration block to the matching archetype function."""
        arch = archetype.lower().strip()

        if arch in ("cover", "hero_cover", "cinematic_cover"):
            hero_img = cfg.get("hero_image_path")
            if arch in ("hero_cover", "cinematic_cover") or hero_img:
                build_hero_cover_slide(
                    self.prs,
                    self.theme,
                    title=cfg.get("title", "Data Application for Financial Analytics with Snowflake"),
                    subtitle=cfg.get("subtitle", "Modernizing Financial Data Architecture, High-Speed Ingestion & Executive Analytics"),
                    client=cfg.get("client", self.client_name),
                    vendor=cfg.get("vendor", self.vendor_name),
                    product=cfg.get("product", self.product),
                    date_str=cfg.get("date_str", self.date_str),
                    tracker=cfg.get("tracker", "ENTERPRISE PRESALES MODERNIZATION"),
                    hero_image_path=hero_img,
                    hero_height=float(cfg.get("hero_height", 3.65)),
                    scrim_alpha=float(cfg.get("scrim_alpha", 0.28)),
                )
            else:
                build_cover_slide(
                    self.prs,
                    self.theme,
                    title=cfg.get("title", "Data Application for Financial Analytics with Snowflake"),
                    subtitle=cfg.get("subtitle", "Modernizing Financial Data Architecture, High-Speed Ingestion & Executive Analytics"),
                    client=cfg.get("client", self.client_name),
                    vendor=cfg.get("vendor", self.vendor_name),
                    product=cfg.get("product", self.product),
                    date_str=cfg.get("date_str", self.date_str),
                    tracker=cfg.get("tracker", "ENTERPRISE PRESALES MODERNIZATION"),
                )

        elif arch in ("agenda", "table_of_contents", "toc"):
            raw_items = cfg.get("items") or []
            items = []
            for itm in raw_items:
                if isinstance(itm, dict):
                    items.append(AgendaItem(
                        num=str(itm.get("num", f"{len(items)+1:02d}")),
                        title=itm.get("title", ""),
                        description=itm.get("description"),
                        badge=itm.get("badge"),
                        accent_key=itm.get("accent_key", "accent"),
                    ))
            build_agenda_slide(
                self.prs,
                self.theme,
                tracker=cfg.get("tracker", "PROJECT ALIGNMENT | EXECUTIVE AGENDA"),
                action_title=cfg.get("title", "Comprehensive Project Alignment Across Governance, Deliverables, and Timelines"),
                subtitle=cfg.get("subtitle", "Structured discussion sequence establishing strategic consensus and stage-gate readiness."),
                items=items if items else None,
                columns=int(cfg.get("columns", 2)),
                current_idx=idx,
                total_slides=total,
                notice=self.notice,
            )

        elif arch in ("chapter_divider", "section_divider"):
            build_chapter_divider_slide(
                self.prs,
                self.theme,
                tracker=cfg.get("tracker", "SECTION"),
                title=cfg.get("headline", cfg.get("title", "Chapter Title")),
                subtitle=cfg.get("subtitle", "Executive overview and transition context."),
                division_tag=cfg.get("division_tagline", "PT Mitra Integrasi Informatika"),
                hero_image_path=cfg.get("hero_image_path"),
                current_idx=idx,
                total_slides=total,
            )

        elif arch in ("corporate_equity_tree", "equity_tree"):
            tree_raw = cfg.get("tree_data")
            tree_data = None
            if tree_raw and isinstance(tree_raw, dict):
                p_nodes_raw = tree_raw.get("tier1_nodes") or tree_raw.get("public_shareholders") or []
                s_nodes_raw = tree_raw.get("tier2_nodes") or tree_raw.get("subsidiaries") or []
                p_nodes = [
                    EquityEntityNode(
                        name=n.get("name", ""),
                        code=n.get("code", ""),
                        percentage=str(n.get("percentage", "")),
                        business_domain=n.get("business_domain") or n.get("domain", ""),
                        tier=int(n.get("tier", 1)),
                        parent_code=n.get("parent_code"),
                        accent_color=n.get("accent_color"),
                    )
                    for n in p_nodes_raw
                    if isinstance(n, dict)
                ]
                s_nodes = [
                    EquityEntityNode(
                        name=n.get("name", ""),
                        code=n.get("code", ""),
                        percentage=str(n.get("percentage", "")),
                        business_domain=n.get("business_domain") or n.get("domain", ""),
                        tier=int(n.get("tier", 2)),
                        parent_code=n.get("parent_code"),
                        accent_color=n.get("accent_color"),
                    )
                    for n in s_nodes_raw
                    if isinstance(n, dict)
                ]
                tree_data = EquityTreeData(
                    parent_company=tree_raw.get("parent_company") or tree_raw.get("holding_name", "PT METRODATA ELECTRONICS TBK"),
                    parent_ticker=tree_raw.get("parent_ticker") or tree_raw.get("holding_ticker", "MTDL"),
                    parent_subtitle=tree_raw.get("parent_subtitle") or tree_raw.get("holding_subtitle", "Leading ICT Solutions & Distribution Ecosystem in Indonesia"),
                    tier1_nodes=p_nodes,
                    tier2_nodes=s_nodes,
                    footnote=tree_raw.get("footnote", "* Consolidated group financial statements and ownership matrix."),
                )
            build_equity_corporate_tree_slide(
                self.prs,
                self.theme,
                tracker=cfg.get("tracker", "CORPORATE STRUCTURE | SUBSIDIARIES & AFFILIATES"),
                action_title=cfg.get("title", "Comprehensive ICT Ecosystem Connects Distribution, Solutions & Cloud Subsidiaries"),
                subtitle=cfg.get("subtitle", "Strategic corporate equity hierarchy across distribution channels, managed services, and enterprise consulting."),
                tree_data=tree_data,
                current_idx=idx,
                total_slides=total,
                notice=self.notice,
            )

        elif arch in ("card_grid", "capability_cards", "pillars_grid"):
            raw_items = cfg.get("items") or []
            items = []
            for itm in raw_items:
                if isinstance(itm, dict):
                    items.append(CardGridItem(
                        title=itm.get("title", ""),
                        subtitle=itm.get("subtitle"),
                        icon=itm.get("icon"),
                        bullets=itm.get("bullets") or [],
                        accent_key=itm.get("accent_key", "accent"),
                        badge=itm.get("badge"),
                    ))
            build_card_grid_slide(
                self.prs,
                self.theme,
                tracker=cfg.get("tracker", "STRATEGIC CAPABILITIES | CORE SERVICE OFFERINGS"),
                action_title=cfg.get("title", "Eight Strategic Pillars Deliver End-to-End Enterprise Modernization"),
                subtitle=cfg.get("subtitle", "Integrated solution pillars supporting digital transformation from foundation cloud infrastructure to applied intelligence."),
                items=items if items else None,
                columns=int(cfg.get("columns", 4)),
                current_idx=idx,
                total_slides=total,
                notice=self.notice,
            )

        elif arch in ("tech_logo_grid", "product_portfolio", "tech_grid"):
            raw_items = cfg.get("items") or []
            items = []
            for itm in raw_items:
                if isinstance(itm, dict):
                    items.append(TechLogoItem(
                        name=itm.get("name", ""),
                        category=itm.get("category", ""),
                        tier_badge=itm.get("tier_badge", "CORE"),
                        logo=itm.get("logo"),
                        accent_key=itm.get("accent_key", "accent"),
                        description=itm.get("description"),
                    ))
            build_tech_logo_grid_slide(
                self.prs,
                self.theme,
                tracker=cfg.get("tracker", "PRODUCT PORTFOLIO | ENTERPRISE DATA ECOSYSTEM"),
                action_title=cfg.get("title", "Best-in-Class Technology Partnerships Ensure Seamless Solution Interoperability"),
                subtitle=cfg.get("subtitle", "Curated partner portfolio spanning database engines, event streaming, data transformation, and visualization."),
                items=items if items else None,
                columns=int(cfg.get("columns", 4)),
                current_idx=idx,
                total_slides=total,
                notice=self.notice,
            )

        elif arch in ("badge_matrix", "certifications", "partner_matrix"):
            raw_sections = cfg.get("sections") or []
            sections = []
            for s in raw_sections:
                if isinstance(s, dict):
                    sections.append(BadgeMatrixSection(
                        section_title=s.get("section_title", ""),
                        accent_key=s.get("accent_key", "accent"),
                        badge_count_label=s.get("badge_count_label"),
                        items=s.get("items") or [],
                    ))
            build_badge_matrix_slide(
                self.prs,
                self.theme,
                tracker=cfg.get("tracker", "COMPETENCIES & CERTIFICATIONS | ENTERPRISE ASSURANCE"),
                action_title=cfg.get("title", "Certified Engineering Teams Guarantee Flawless Implementation and High Availability"),
                subtitle=cfg.get("subtitle", "Validated competencies across data lakehouse, cloud hyperscalers, and regulatory governance standards."),
                sections=sections if sections else None,
                current_idx=idx,
                total_slides=total,
                notice=self.notice,
            )

        elif arch in ("iceberg_concept", "iceberg"):
            build_iceberg_concept_slide(
                self.prs,
                self.theme,
                tracker=cfg.get("tracker", "ANALYTICS FOUNDATION | THE ICEBERG CONCEPT"),
                action_title=cfg.get("title", "Modern Data Platforms Eliminate the Subsurface Technical Debt That Breaks Dashboards"),
                subtitle=cfg.get("subtitle", "While executives view surface reports, 85% of analytic resilience relies on robust foundation data engineering."),
                visible_items=cfg.get("visible_items"),
                subsurface_items=cfg.get("subsurface_items"),
                takeaways=cfg.get("takeaways"),
                current_idx=idx,
                total_slides=total,
                notice=self.notice,
            )

        elif arch in ("split_hero", "kpi_framework"):
            blocks = []
            for b in (cfg.get("blocks") or []):
                if isinstance(b, dict):
                    blocks.append(SplitHeroBlock(
                        title=b.get("title", ""),
                        description=b.get("description", b.get("body", "")),
                        metric_badge=b.get("metric_badge", b.get("badge")),
                        category_tag=b.get("category_tag", b.get("tag")),
                    ))
            build_split_hero_slide(
                self.prs,
                self.theme,
                tracker=cfg.get("tracker", "ANALYTICS JOURNEY | KPI DISCOVERY FRAMEWORK"),
                action_title=cfg.get("title", "High-Performance Analytics Begins with Clarifying Strategic Business Objectives"),
                subtitle=cfg.get("subtitle", "Aligning technical data schemas with executive decision metrics ensures measurable return on investment."),
                hero_stat=cfg.get("hero_stat", cfg.get("hero_tag", "4 STAGES")),
                hero_stat_label=cfg.get("hero_stat_label", cfg.get("hero_title", "KPI Framework")),
                hero_stat_subtext=cfg.get("hero_stat_subtext", cfg.get("hero_body")),
                blocks=blocks if blocks else None,
                current_idx=idx,
                total_slides=total,
                notice=self.notice,
            )

        elif arch in ("snowflake_platform_architecture", "platform_architecture"):
            build_snowflake_platform_architecture_slide(
                self.prs,
                self.theme,
                tracker=cfg.get("tracker", "TECHNOLOGY PLATFORM BLUEPRINT | SNOWFLAKE DATA CLOUD"),
                action_title=cfg.get("title", "Fully-Managed Unified Platform Eliminates Data Silos Across Multi-Cloud Environments"),
                subtitle=cfg.get("subtitle", "Official Snowflake platform blueprint showing decoupled storage, elastic compute, and Horizon governance."),
                current_idx=idx,
                total_slides=total,
                notice=self.notice,
            )

        elif arch in ("snowflake_data_pipeline", "solution_architecture", "data_pipeline"):
            build_snowflake_data_pipeline_slide(
                self.prs,
                self.theme,
                tracker=cfg.get("tracker", "SOLUTION ARCHITECTURE | END-TO-END DATA PIPELINE"),
                action_title=cfg.get("title", "Modern Data Pipeline Streamlines Continuous Ingestion into Governed Analytics Consumption"),
                subtitle=cfg.get("subtitle", "Official end-to-end data flow: transactional ERP/IoT sources into Snowflake core, virtual warehouses, and delivery channels."),
                current_idx=idx,
                total_slides=total,
                notice=self.notice,
            )

        elif arch in ("bcg_3_horizon", "3_tier_architecture", "horizons"):
            horizons = []
            for h in (cfg.get("horizons") or []):
                if isinstance(h, dict):
                    h_tag = h.get("horizon_tag") or h.get("tag") or h.get("name") or "HORIZON"
                    h_title = h.get("title") or ""
                    h_focus = h.get("strategic_focus") or h.get("focus") or h.get("subtitle") or ""
                    h_metric = h.get("metric_highlight") or h.get("metric") or ""
                    h_label = h.get("metric_label") or h.get("subtext") or ""
                    h_inits = h.get("initiatives") or h.get("bullets") or []
                    horizons.append(HorizonColumnData(
                        horizon_tag=h_tag,
                        title=h_title,
                        strategic_focus=h_focus,
                        metric_highlight=h_metric,
                        metric_label=h_label,
                        initiatives=h_inits,
                        accent_key=h.get("accent_key", "accent"),
                        status_tag=h.get("status_tag") or h.get("status", "IN EXECUTION"),
                        status_bg_key=h.get("status_bg_key", "badge_blue_fill"),
                        status_text_key=h.get("status_text_key", "badge_blue_text"),
                        icon=h.get("icon"),
                        icon_path=h.get("icon_path"),
                    ))
            build_bcg_3_horizon_slide(
                self.prs,
                self.theme,
                tracker=cfg.get("tracker", "SNOWFLAKE ARCHITECTURE | THE AI DATA CLOUD"),
                action_title=cfg.get("title", "Decoupled Architecture Delivers Infinite Compute Elasticity with Zero Resource Contention"),
                subtitle=cfg.get("subtitle", "Three independent layers: centralized cloud storage, multi-cluster elastic compute, and unified cloud services."),
                horizons=horizons if horizons else None,
                current_idx=idx,
                total_slides=total,
                notice=self.notice,
            )

        elif arch in ("browser_mockup", "window_mockup", "app_mockup", "app_walkthrough", "module_walkthrough"):
            build_browser_mockup_slide(
                self.prs,
                self.theme,
                tracker=cfg.get("tracker", "SYSTEM INTERFACE & RUNTIME VALIDATION"),
                action_title=cfg.get("title", "Modernized Analytics Platform Eliminates ERP Latency and Siloed Data Exports"),
                subtitle=cfg.get("subtitle", "High-fidelity interface preview running directly on Snowflake Cortex AI and Streamlit serving tiers."),
                screenshot_path=cfg.get("screenshot_path") or cfg.get("screenshot") or cfg.get("image_path"),
                url=cfg.get("url", "https://finance.snowflakecomputing.com/streamlit/app"),
                telemetry_badges=cfg.get("telemetry_badges"),
                theme_mode=cfg.get("theme_mode", "light"),
                takeaways_title=cfg.get("takeaways_title", "KEY OBSERVATIONS & DRIVERS"),
                takeaways=cfg.get("takeaways"),
                current_idx=idx,
                total_slides=total,
                notice=self.notice,
            )

        elif arch in ("gap_analysis", "as_is_to_be"):
            build_gap_analysis_slide(
                self.prs,
                self.theme,
                tracker=cfg.get("tracker", "GAP ANALYSIS | AS-IS VS TO-BE"),
                action_title=cfg.get("title", "Strategic Modernization Bridges Current Legacy Friction into Agile Intelligence"),
                subtitle=cfg.get("subtitle", "Detailed comparative analysis across latency, data quality, scalability, and operational governance."),
                current_idx=idx,
                total_slides=total,
                notice=self.notice,
            )

        elif arch in ("chevron_process", "process_flow", "value_chain"):
            steps = []
            for s in (cfg.get("steps") or []):
                if isinstance(s, dict):
                    acc_key = s.get("accent_key")
                    acc_col = self.theme.get_rgb(acc_key) if acc_key else None
                    steps.append(ProcessChevronStep(
                        phase_number=s.get("phase_number", "01"),
                        title=s.get("title", ""),
                        duration_badge=s.get("duration_badge", s.get("subtitle")),
                        deliverables=s.get("deliverables") or [],
                        status=s.get("status"),
                        accent_color=acc_col,
                    ))
            build_chevron_process_slide(
                self.prs,
                self.theme,
                tracker=cfg.get("tracker", "DELIVERY METHODOLOGY | EXECUTION ROADMAP"),
                action_title=cfg.get("title", "Structured Four-Stage Delivery Framework Drives Predictable Implementation"),
                subtitle=cfg.get("subtitle", "Sequential phase gates: Initiation, Cloud Development, Integration Testing, and Executive Handover."),
                steps=steps if steps else None,
                current_idx=idx,
                total_slides=total,
                notice=self.notice,
            )

        elif arch in ("timeline_gantt", "delivery_gantt", "gantt"):
            g_raw = cfg.get("gantt_data")
            g_data = None
            if g_raw and isinstance(g_raw, dict):
                ws_list = []
                for ws in (g_raw.get("workstreams") or []):
                    if isinstance(ws, dict):
                        t_list = []
                        for t in (ws.get("tasks") or []):
                            if isinstance(t, dict):
                                t_dict = dict(t)
                                if "start_period" not in t_dict and "start_week" in t_dict:
                                    t_dict["start_period"] = float(t_dict.pop("start_week"))
                                if "end_period" not in t_dict:
                                    if "end_week" in t_dict:
                                        t_dict["end_period"] = float(t_dict.pop("end_week"))
                                    elif "duration_weeks" in t_dict:
                                        dur = float(t_dict.pop("duration_weeks"))
                                        t_dict["end_period"] = float(t_dict.get("start_period", 1.0) + max(1.0, dur) - 1.0)
                                    else:
                                        t_dict["end_period"] = float(t_dict.get("start_period", 1.0))
                                t_dict.pop("duration_weeks", None)
                                t_dict.pop("end_week", None)
                                t_list.append(GanttTask(
                                    name=t_dict.get("name", ""),
                                    start_period=float(t_dict.get("start_period", 1.0)),
                                    end_period=float(t_dict.get("end_period", 1.0)),
                                    progress_pct=float(t_dict.get("progress_pct", 100.0)),
                                    status=t_dict.get("status", "COMPLETED"),
                                    owner=t_dict.get("owner"),
                                    is_milestone=bool(t_dict.get("is_milestone", False)),
                                ))
                            elif isinstance(t, GanttTask):
                                t_list.append(t)
                        cat = ws.get("category") or ws.get("name", "Workstream")
                        accent_c = ws.get("accent_color") or ws.get("accent_key", "primary")
                        badge_c = ws.get("badge", "WORKSTREAM")
                        ws_list.append(GanttWorkstream(
                            category=cat,
                            badge=badge_c,
                            accent_color=accent_c,
                            tasks=t_list,
                        ))
                    elif isinstance(ws, GanttWorkstream):
                        ws_list.append(ws)
                total_p = int(g_raw.get("total_periods") or g_raw.get("total_weeks", 12))
                cur_marker = float(g_raw["current_period_marker"]) if "current_period_marker" in g_raw and g_raw["current_period_marker"] is not None else 6.0
                g_data = GanttTimelineData(
                    time_scale_label=g_raw.get("time_scale_label", "WEEKS"),
                    total_periods=total_p,
                    period_labels=g_raw.get("period_labels"),
                    current_period_marker=cur_marker,
                    workstreams=ws_list,
                    footnote=g_raw.get("footnote", "*) Timeline reflects target PMO sprint velocity and scheduled governance gates."),
                )
            build_timeline_gantt_slide(
                self.prs,
                self.theme,
                tracker=cfg.get("tracker", "PROJECT TIMELINE & DELIVERY GANTT"),
                action_title=cfg.get("title", "16-Week Modernization Roadmap Structured Across Clear Milestone Gates"),
                subtitle=cfg.get("subtitle", "Multi-phase implementation schedule coordinating cloud staging, dbt transformations, and UAT sign-off."),
                gantt_data=g_data,
                current_idx=idx,
                total_slides=total,
                notice=self.notice,
            )

        elif arch in ("governance_org_structure", "org_chart", "raci"):
            build_governance_org_structure_slide(
                self.prs,
                self.theme,
                tracker=cfg.get("tracker", "PROJECT GOVERNANCE | ORGANIZATIONAL STRUCTURE"),
                action_title=cfg.get("title", "Joint Dual-Pillar Governance Matrix Establishes Clear Escalation & Delivery Ownership"),
                subtitle=cfg.get("subtitle", "Hierarchical project structure connects executive steering committees, PMO leads, and specialized execution pods."),
                client_name=cfg.get("client_name", self.client_name),
                vendor_name=cfg.get("vendor_name", self.vendor_name),
                current_idx=idx,
                total_slides=total,
                notice=self.notice,
                tier1_client_text=cfg.get("tier1_client_text"),
                tier1_vendor_text=cfg.get("tier1_vendor_text"),
                client_pm_title=cfg.get("client_pm_title"),
                vendor_pm_title=cfg.get("vendor_pm_title"),
                client_pm_bullets=cfg.get("client_pm_bullets"),
                vendor_pm_bullets=cfg.get("vendor_pm_bullets"),
                pods=cfg.get("pods"),
            )

        elif arch in ("change_request_procedure", "cr_procedure"):
            build_change_request_procedure_slide(
                self.prs,
                self.theme,
                tracker=cfg.get("tracker", "PROJECT GOVERNANCE | SCOPE & CHANGE CONTROL"),
                action_title=cfg.get("title", "Structured Change Request Procedure Governs Scope Adjustments Through Defined Decision Gates"),
                subtitle=cfg.get("subtitle", "Sequential 4-stage governance pipeline with bi-level escalation thresholds and auditable legal sign-offs."),
                current_idx=idx,
                total_slides=total,
                notice=self.notice,
            )

        elif arch in ("table", "matrix_table", "milestones_table", "prerequisites_table", "communication_table", "risk_table", "table_matrix"):
            raw_cols = cfg.get("columns") or []
            columns = []
            for c in raw_cols:
                if isinstance(c, dict):
                    columns.append(TableColumnDef(
                        header=c.get("header", ""),
                        width=float(c.get("width", 2.0)),
                        align=str(c.get("align", "LEFT")).upper(),
                        is_status=bool(c.get("is_status", False)),
                        is_index=bool(c.get("is_index", False)),
                    ))
            rows = cfg.get("rows") or []
            build_table_slide(
                self.prs,
                self.theme,
                tracker=cfg.get("tracker", "PROJECT GOVERNANCE | REGISTER"),
                action_title=cfg.get("title", "Structured Governance Register Establishes Accountabilities and Timelines"),
                subtitle=cfg.get("subtitle", "Detailed tracking matrix with designated owners, baseline targets, and delivery statuses."),
                columns=columns if columns else None,
                rows=rows if rows else None,
                footnote=cfg.get("footnote"),
                current_idx=idx,
                total_slides=total,
                notice=self.notice,
            )

        elif arch in ("thank_you", "closing", "thank_you_closing"):
            build_thank_you_slide(
                self.prs,
                self.theme,
                tracker=cfg.get("tracker", "ENGAGEMENT WRAP-UP | THANK YOU"),
                title=cfg.get("title", "Thank You"),
                subtitle=cfg.get("subtitle", "Partnering to accelerate enterprise data modernization with Snowflake AI Data Cloud."),
                contacts=cfg.get("contacts"),
                company=cfg.get("company", f"{self.vendor_name}  |  PT Mitra Integrasi Informatika"),
                office=cfg.get("office", "APL Tower 37th Floor, Jl. Letjen S. Parman Kav. 28, Jakarta Barat 11470"),
                current_idx=idx,
                total_slides=total,
                notice=self.notice,
            )

        elif arch in ("feature_matrix", "harvey_balls"):
            m_raw = cfg.get("matrix_data")
            m_data = None
            if m_raw and isinstance(m_raw, dict):
                f_rows = []
                for r in (m_raw.get("rows") or []):
                    if isinstance(r, dict):
                        from src.ppt_engine.consulting_archetypes import FeatureScorecardRow
                        f_rows.append(FeatureScorecardRow(
                            category=r.get("category", ""),
                            criteria=r.get("criteria", ""),
                            scores=r.get("scores") or {},
                            business_impact=r.get("business_impact", "HIGH"),
                        ))
                    elif hasattr(r, "category"):
                        f_rows.append(r)
                m_data = FeatureMatrixData(
                    columns=m_raw.get("columns") or [],
                    rows=f_rows,
                    recommendation_badge=m_raw.get("recommendation_badge", "RECOMMENDED ARCHITECTURE"),
                    recommended_column=m_raw.get("recommended_column", "Snowflake Data Cloud"),
                    footnote=m_raw.get("footnote", "Harvey Balls: ● Full Native Support | ◐ Partial / Requires Extension | ○ Complex / Unsupported"),
                )
            build_feature_matrix_slide(
                self.prs,
                self.theme,
                tracker=cfg.get("tracker", "ARCHITECTURE EVALUATION | PLATFORM BENCHMARK"),
                action_title=cfg.get("title", "Snowflake AI Data Cloud Demonstrates Superior Capabilities Across Core Criteria"),
                subtitle=cfg.get("subtitle", "Multi-dimensional comparative assessment evaluating compute elasticity, data governance, and native AI integration."),
                matrix_data=m_data,
                current_idx=idx,
                total_slides=total,
                notice=self.notice,
            )

        else:
            logger.warning(f"Unknown archetype '{archetype}' at slide {idx}; falling back to chapter divider.")
            build_chapter_divider_slide(
                self.prs,
                self.theme,
                tracker=f"SLIDE {idx:02d}",
                title=cfg.get("title", f"Slide {idx}"),
                subtitle=cfg.get("subtitle", "Executive slide content."),
                current_idx=idx,
                total_slides=total,
            )

    def update_pagination(self) -> None:
        """
        Post-processes all slides to synchronize bottom footer pagination ('02 / 36').
        Preserves existing font styling, size, boldness, and color.
        Cover slide (Slide 1) strictly has no footer or pagination.
        """
        total_slides = len(self.prs.slides)
        for s_idx, slide in enumerate(self.prs.slides, start=1):
            if s_idx == 1:
                continue
            for shape in slide.shapes:
                if not getattr(shape, "has_text_frame", False) or not shape.has_text_frame:
                    continue
                try:
                    if getattr(shape, "top", None) is not None and shape.top < Inches(6.8):
                        continue
                except Exception:
                    pass
                for paragraph in shape.text_frame.paragraphs:
                    text_clean = paragraph.text.strip()
                    if PAGE_REGEX.match(text_clean):
                        new_page_str = f"{s_idx:02d} / {total_slides:02d}"
                        if paragraph.runs:
                            r0 = paragraph.runs[0]
                            font_name = r0.font.name
                            font_size = r0.font.size
                            font_bold = r0.font.bold
                            font_color = None
                            try:
                                if r0.font.color and getattr(r0.font.color, "type", None) == 1:
                                    font_color = r0.font.color.rgb
                            except (AttributeError, ValueError):
                                font_color = None
                            paragraph.text = new_page_str
                            if paragraph.runs:
                                nr = paragraph.runs[0]
                                nr.font.name = font_name
                                nr.font.size = font_size
                                nr.font.bold = font_bold
                                if font_color:
                                    nr.font.color.rgb = font_color
                        else:
                            paragraph.text = new_page_str

    def save(
        self,
        output_path: Union[str, Path],
        engagement_context: Optional[EngagementContext] = None,
    ) -> Path:
        """Saves presentation deck to target path after synchronizing pagination and substituting slugs."""
        self.update_pagination()
        if engagement_context is None:
            engagement_context = EngagementContext.default_ngl()
        substitute_slugs_in_presentation(self.prs, engagement_context)

        p = Path(output_path)
        p.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(p))
        logger.info(f"Successfully saved pitch deck to {p}")
        return p
