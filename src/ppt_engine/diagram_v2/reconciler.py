"""
Diagram Engine V2 - Bidirectional PPTX Decompiler & YAML Reconciliation Engine
==============================================================================
Diffs native PowerPoint DrawingML shapes against the declarative DiagramSpec AST,
detecting deleted nodes and flows, pruning orphan connectors, extracting mutated
text (titles, subtitles, kickers, container labels, badges, and flow pills),
updating the YAML specification on disk, and triggering automatic 4-way re-compilation.
"""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Set, Tuple, Union

import yaml
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from src.ppt_engine.diagram_v2.models import (
    DiagramSpec,
    ContainerSpec,
    NodeSpec,
    FlowSpec,
    BadgeSpec,
    BulletListSpec,
)
from src.ppt_engine.diagram_v2.compiler import compile_diagram, load_diagram_spec, EngineV2Result

logger = logging.getLogger(__name__)


class IndentedYamlDumper(yaml.SafeDumper):
    """Custom YAML SafeDumper that maintains 2-space indented list markers."""
    def increase_indent(self, flow=False, indentless=False):
        return super().increase_indent(flow, False)


def _represent_list(dumper: Any, data: Any) -> Any:
    """Format short numeric tuples/lists (e.g. coordinates [x, y]) inline in YAML."""
    if data and all(isinstance(x, (int, float)) for x in data) and len(data) <= 4:
        return dumper.represent_sequence("tag:yaml.org,2002:seq", data, flow_style=True)
    return dumper.represent_sequence("tag:yaml.org,2002:seq", data, flow_style=False)


IndentedYamlDumper.add_representer(list, _represent_list)


def dump_diagram_yaml(data: Dict[str, Any], path: Union[str, Path]) -> Path:
    """Writes a diagram specification dictionary to a beautifully indented YAML file."""
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    content = yaml.dump(data, Dumper=IndentedYamlDumper, sort_keys=False, indent=2, allow_unicode=True)
    p.write_text(content, encoding="utf-8")
    return p


@dataclass
class ReconciledChange:
    """Individual structural, textual, or geometric change detected between PPTX shapes and YAML AST."""
    change_type: str  # "pruned_node", "pruned_flow", "pruned_container", "pruned_badge", "pruned_bullet_list", "mutated_text", "mutated_geometry"
    element_type: str  # "node", "flow", "container", "badge", "bullet_list", "header"
    element_id: str
    field_name: Optional[str] = None
    old_value: Any = None
    new_value: Any = None
    description: str = ""


@dataclass
class ReconciliationResult:
    """Summary of reconciliation execution, detected diffs, and recompilation status."""
    spec_path: Path
    pptx_path: Path
    slide_index: int
    changes: List[ReconciledChange] = field(default_factory=list)
    updated_spec: Optional[DiagramSpec] = None
    recompiled: bool = False
    recompile_result: Optional[EngineV2Result] = None

    @property
    def has_changes(self) -> bool:
        return len(self.changes) > 0

    @property
    def pruned_nodes(self) -> List[ReconciledChange]:
        return [c for c in self.changes if c.change_type == "pruned_node"]

    @property
    def pruned_flows(self) -> List[ReconciledChange]:
        return [c for c in self.changes if c.change_type == "pruned_flow"]

    @property
    def mutated_texts(self) -> List[ReconciledChange]:
        return [c for c in self.changes if c.change_type == "mutated_text"]

    @property
    def mutated_geometries(self) -> List[ReconciledChange]:
        return [c for c in self.changes if c.change_type == "mutated_geometry"]


_BULLET_PREFIX_RE = re.compile(r"^[•\-\*]\s*")


def _clean_text(val: Optional[str]) -> Optional[str]:
    """Normalizes vertical tabs (\x0b / <a:br/>) and Windows line endings to \n."""
    if val is None:
        return None
    return val.replace("\x0b", "\n").replace("\r\n", "\n").strip()


class PPTXDiagramDecompiler:
    """
    Decompiles and inspects an existing PowerPoint slide, extracting shapes tagged
    with deterministic OpenXML names ('node:...', 'flow:...', 'container:...', etc.).
    """

    _DETERMINISTIC_PREFIXES = ("node:", "flow:", "container:", "badge:", "header:", "bullet_list:")

    def __init__(self, pptx_path: Union[str, Path], slide_index: int = 1) -> None:
        self.pptx_path = Path(pptx_path)
        if not self.pptx_path.exists():
            raise FileNotFoundError(f"PowerPoint presentation not found: {self.pptx_path}")

        self.prs = Presentation(str(self.pptx_path))
        if slide_index < 1 or slide_index > len(self.prs.slides):
            raise IndexError(f"Slide index {slide_index} out of bounds (1..{len(self.prs.slides)})")

        self.slide_index = slide_index
        self.slide = self.prs.slides[slide_index - 1]
        self.shapes = self._collect_shapes(self.slide.shapes)
        self.shapes_by_name: Dict[str, Any] = {}
        self._shapes_by_element_key: Dict[str, List[Any]] = {}
        self.all_names: Set[str] = set()

        for s in self.shapes:
            name = (getattr(s, "name", "") or "").strip()
            if name:
                self.shapes_by_name[name] = s
                self.all_names.add(name)
                parts = name.split(":", 2)
                if len(parts) >= 2:
                    elem_key = f"{parts[0]}:{parts[1]}"
                    self._shapes_by_element_key.setdefault(elem_key, []).append(s)

    def _collect_shapes(self, shape_tree: Any) -> List[Any]:
        """Recursively flattens all slide shapes including nested group shapes."""
        flat: List[Any] = []
        for s in shape_tree:
            flat.append(s)
            if getattr(s, "shape_type", None) == MSO_SHAPE_TYPE.GROUP or hasattr(s, "shapes"):
                try:
                    flat.extend(self._collect_shapes(s.shapes))
                except Exception:
                    pass
        return flat

    def has_deterministic_names(self) -> bool:
        """Checks if the slide contains shapes tagged with Diagram Engine V2 deterministic names."""
        return any(name.startswith(self._DETERMINISTIC_PREFIXES) for name in self.all_names)

    def get_shape(self, name: str) -> Optional[Any]:
        return self.shapes_by_name.get(name)

    def find_shapes_by_prefix(self, prefix: str) -> List[Any]:
        """
        Finds shapes matching a deterministic prefix.
        When given a 2-part element key like 'node:app', matches exact 'node:app' or
        sub-elements 'node:app:*' without false-matching sibling prefixes like 'node:app_db'.
        """
        if not prefix.endswith(":") and prefix in self._shapes_by_element_key:
            return list(self._shapes_by_element_key[prefix])
        if prefix.endswith(":"):
            return [s for s in self.shapes if (getattr(s, "name", "") or "").startswith(prefix)]
        colon_prefix = f"{prefix}:"
        return [
            s for s in self.shapes
            if (name := (getattr(s, "name", "") or "")) == prefix or name.startswith(colon_prefix)
        ]

    @staticmethod
    def get_text_paragraphs(shape: Any) -> List[str]:
        """Extracts non-empty stripped text paragraphs from a shape's text frame."""
        if not getattr(shape, "has_text_frame", False) or shape.text_frame is None:
            return []
        paras: List[str] = []
        for p in shape.text_frame.paragraphs:
            cleaned = _clean_text(p.text)
            if cleaned:
                paras.append(cleaned)
        return paras


class DiagramReconciler:
    """
    Core engine comparing PowerPoint slide shape state against the declarative YAML AST,
    applying diffs, pruning orphans, and synchronizing 4-way output deliverables.
    """

    def __init__(
        self,
        spec_path: Union[str, Path],
        pptx_path: Optional[Union[str, Path]] = None,
        slide_index: int = 1,
        snap_grid: float = 0.05,
        snap_threshold: float = 0.02,
    ) -> None:
        self.spec_path = Path(spec_path)
        if not self.spec_path.exists():
            raise FileNotFoundError(f"Diagram specification not found: {self.spec_path}")

        raw_text = self.spec_path.read_text(encoding="utf-8")
        loaded = yaml.safe_load(raw_text)
        if not isinstance(loaded, dict):
            raise ValueError(f"Invalid diagram YAML specification in {self.spec_path}: expected top-level mapping.")
        self.raw_data: Dict[str, Any] = loaded
        self.spec = DiagramSpec.model_validate(self.raw_data)

        # Resolve PPTX path
        if pptx_path:
            self.pptx_path = Path(pptx_path)
        else:
            self.pptx_path = self._auto_locate_pptx()

        self.slide_index = slide_index
        self.snap_grid = max(0.001, float(snap_grid))
        self.snap_threshold = max(0.0001, float(snap_threshold))

    def _auto_locate_pptx(self) -> Path:
        """Attempts to locate the compiled PPTX corresponding to the YAML spec."""
        candidates = [
            Path("output/diagrams_v2") / f"{self.spec.id}.pptx",
            Path("output/diagrams_v2") / f"{self.spec_path.stem}.pptx",
            self.spec_path.parent / f"{self.spec_path.stem}.pptx",
            self.spec_path.parent / f"{self.spec.id}.pptx",
        ]
        for c in candidates:
            if c.exists():
                return c
        raise FileNotFoundError(
            f"Could not auto-locate compiled PPTX for {self.spec_path}. "
            f"Please specify --pptx path/to/presentation.pptx"
        )

    def _to_spec_units(self, inches_val: float) -> float:
        """Converts inches to specification units (e.g. px if canvas.unit == 'px')."""
        if self.spec.canvas.unit == "px":
            return round(inches_val * 90.0, 2)
        return inches_val

    def _to_inches(self, spec_val: float) -> float:
        """Converts specification units to inches for geometric comparison."""
        if self.spec.canvas.unit == "px":
            return spec_val / 90.0
        return spec_val

    def _snap_spec_coord(self, inches_val: float) -> float:
        """Snaps an inch coordinate value to snap_grid and returns in spec units."""
        snapped_in = round(round(inches_val / self.snap_grid) * self.snap_grid, 4)
        if abs(snapped_in) < 1e-9:
            snapped_in = 0.0
        return self._to_spec_units(snapped_in)

    def reconcile(
        self,
        dry_run: bool = False,
        no_recompile: bool = False,
        formats: Sequence[str] = ("all",),
        skip_pptx: bool = False,
    ) -> ReconciliationResult:
        """Executes full bidirectional reconciliation between PPTX shapes and YAML AST."""
        decompiler = PPTXDiagramDecompiler(self.pptx_path, slide_index=self.slide_index)

        if not decompiler.has_deterministic_names():
            raise RuntimeError(
                f"The PowerPoint slide at '{self.pptx_path}' (Slide {self.slide_index}) does not contain "
                f"deterministic OpenXML shape names ('node:...', 'flow:...', etc.). "
                f"Please compile the slide with 'bench diagram v2 compile {self.spec_path}' first."
            )

        changes: List[ReconciledChange] = []
        deleted_node_ids: Set[str] = set()
        deleted_flow_ids: Set[str] = set()

        # ----------------------------------------------------------------------
        # 1. Header Reconciliation
        # ----------------------------------------------------------------------
        header_shape = decompiler.get_shape("header:title") or decompiler.get_shape("header:text")
        if header_shape:
            h_paras = decompiler.get_text_paragraphs(header_shape)
            if h_paras:
                new_title = h_paras[0]
                if new_title != self.spec.header.action_title:
                    changes.append(ReconciledChange(
                        change_type="mutated_text",
                        element_type="header",
                        element_id="header",
                        field_name="action_title",
                        old_value=self.spec.header.action_title,
                        new_value=new_title,
                        description=f"Action title mutated from '{self.spec.header.action_title}' to '{new_title}'",
                    ))
                    self.spec.header.action_title = new_title
                    if isinstance(self.raw_data.get("header"), dict):
                        self.raw_data["header"]["action_title"] = new_title

            if len(h_paras) >= 2:
                new_sub = "\n".join(h_paras[1:])
                if new_sub != self.spec.header.subtitle:
                    changes.append(ReconciledChange(
                        change_type="mutated_text",
                        element_type="header",
                        element_id="header",
                        field_name="subtitle",
                        old_value=self.spec.header.subtitle,
                        new_value=new_sub,
                        description=f"Header subtitle mutated from '{self.spec.header.subtitle}' to '{new_sub}'",
                    ))
                    self.spec.header.subtitle = new_sub
                    if isinstance(self.raw_data.get("header"), dict):
                        self.raw_data["header"]["subtitle"] = new_sub
            elif len(h_paras) == 1 and self.spec.header.subtitle:
                changes.append(ReconciledChange(
                    change_type="mutated_text",
                    element_type="header",
                    element_id="header",
                    field_name="subtitle",
                    old_value=self.spec.header.subtitle,
                    new_value="",
                    description="Header subtitle cleared in PPTX",
                ))
                self.spec.header.subtitle = ""
                if isinstance(self.raw_data.get("header"), dict):
                    self.raw_data["header"]["subtitle"] = ""

        # ----------------------------------------------------------------------
        # 2. Container Boundaries Reconciliation
        # ----------------------------------------------------------------------
        surviving_containers: List[ContainerSpec] = []
        deleted_container_ids: Set[str] = set()

        for c in self.spec.containers:
            c_shapes = decompiler.find_shapes_by_prefix(f"container:{c.id}")
            if not c_shapes:
                deleted_container_ids.add(c.id)
                changes.append(ReconciledChange(
                    change_type="pruned_container",
                    element_type="container",
                    element_id=c.id,
                    old_value=c.label,
                    new_value=None,
                    description=f"Container '{c.id}' ('{c.label}') was removed in PPTX",
                ))
                continue

            surviving_containers.append(c)

            # Check geometric position & dimensions
            c_shape = decompiler.get_shape(f"container:{c.id}")
            if c_shape and hasattr(c_shape, "left") and hasattr(c_shape, "top") and hasattr(c_shape, "width") and hasattr(c_shape, "height"):
                cur_x_in = float(c_shape.left) / 914400.0
                cur_y_in = float(c_shape.top) / 914400.0
                cur_w_in = float(c_shape.width) / 914400.0
                cur_h_in = float(c_shape.height) / 914400.0

                for prop, cur_in, orig_val in [
                    ("x", cur_x_in, c.x),
                    ("y", cur_y_in, c.y),
                    ("w", cur_w_in, c.w),
                    ("h", cur_h_in, c.h),
                ]:
                    orig_in = self._to_inches(orig_val)
                    if abs(cur_in - orig_in) >= self.snap_threshold:
                        snapped_val = self._snap_spec_coord(cur_in)
                        if abs(snapped_val - orig_val) > 1e-4:
                            changes.append(ReconciledChange(
                                change_type="mutated_geometry",
                                element_type="container",
                                element_id=c.id,
                                field_name=prop,
                                old_value=orig_val,
                                new_value=snapped_val,
                                description=f"Container '{c.id}' {prop.upper()} adjusted from {orig_val} to {snapped_val}",
                            ))
                            setattr(c, prop, snapped_val)
                            self._update_container_dict(c.id, prop, snapped_val)

            # Check boundary label
            lbl_shape = decompiler.get_shape(f"container:{c.id}:label")
            if lbl_shape:
                lbl_paras = decompiler.get_text_paragraphs(lbl_shape)
                if lbl_paras:
                    new_lbl = lbl_paras[0]
                    # Note: compiler uppercases label for visual hierarchy
                    if new_lbl != c.label and new_lbl != c.label.upper():
                        changes.append(ReconciledChange(
                            change_type="mutated_text",
                            element_type="container",
                            element_id=c.id,
                            field_name="label",
                            old_value=c.label,
                            new_value=new_lbl,
                            description=f"Container '{c.id}' label mutated from '{c.label}' to '{new_lbl}'",
                        ))
                        c.label = new_lbl
                        self._update_container_dict(c.id, "label", new_lbl)

                # Check label vertical position (bottom-edge policy)
                c_shape = decompiler.get_shape(f"container:{c.id}")
                is_bottom = False
                if c_shape and hasattr(c_shape, "top") and hasattr(c_shape, "height") and hasattr(lbl_shape, "top"):
                    is_bottom = lbl_shape.top > (c_shape.top + c_shape.height / 2)
                elif hasattr(lbl_shape, "top"):
                    is_bottom = (lbl_shape.top / 914400.0) > (c.y + c.h / 2)

                if is_bottom:
                    if c.h < 4.0:
                        if c.vertical_align != "bottom":
                            changes.append(ReconciledChange(
                                change_type="mutated_text",
                                element_type="container",
                                element_id=c.id,
                                field_name="vertical_align",
                                old_value=c.vertical_align,
                                new_value="bottom",
                                description=f"Container '{c.id}' bottom label detected for h={c.h:.2f} < 4.0''",
                            ))
                            c.vertical_align = "bottom"
                            self._update_container_dict(c.id, "vertical_align", "bottom")
                    else:
                        # h >= 4.0'': omit vertical_align or keep clean to avoid YAML churn
                        if c.vertical_align is not None:
                            self._update_container_dict(c.id, "vertical_align", None)
                            c.vertical_align = None
                else:
                    if c.h >= 4.0:
                        if c.vertical_align != "top":
                            changes.append(ReconciledChange(
                                change_type="mutated_text",
                                element_type="container",
                                element_id=c.id,
                                field_name="vertical_align",
                                old_value=c.vertical_align,
                                new_value="top",
                                description=f"Container '{c.id}' top label detected for h={c.h:.2f} >= 4.0''",
                            ))
                            c.vertical_align = "top"
                            self._update_container_dict(c.id, "vertical_align", "top")
                    else:
                        # h < 4.0'': omit vertical_align as top is already default
                        if c.vertical_align is not None:
                            self._update_container_dict(c.id, "vertical_align", None)
                            c.vertical_align = None

            # Check tag badge
            tag_shape = decompiler.get_shape(f"container:{c.id}:tag_badge")
            if tag_shape:
                tag_paras = decompiler.get_text_paragraphs(tag_shape)
                if tag_paras:
                    new_tag = tag_paras[0]
                    if new_tag != c.tag_badge and new_tag != (c.tag_badge or "").upper():
                        changes.append(ReconciledChange(
                            change_type="mutated_text",
                            element_type="container",
                            element_id=c.id,
                            field_name="tag_badge",
                            old_value=c.tag_badge,
                            new_value=new_tag,
                            description=f"Container '{c.id}' tag badge mutated from '{c.tag_badge}' to '{new_tag}'",
                        ))
                        c.tag_badge = new_tag
                        self._update_container_dict(c.id, "tag_badge", new_tag)
            elif c.tag_badge is not None:
                changes.append(ReconciledChange(
                    change_type="mutated_text",
                    element_type="container",
                    element_id=c.id,
                    field_name="tag_badge",
                    old_value=c.tag_badge,
                    new_value=None,
                    description=f"Container '{c.id}' tag badge was removed in PPTX",
                ))
                c.tag_badge = None
                self._update_container_dict(c.id, "tag_badge", None)

        if deleted_container_ids:
            self.spec.containers = surviving_containers
            if "containers" in self.raw_data:
                self.raw_data["containers"] = [
                    cd for cd in self.raw_data["containers"] if cd.get("id") not in deleted_container_ids
                ]

        # ----------------------------------------------------------------------
        # 3. Entity Nodes Reconciliation
        # ----------------------------------------------------------------------
        surviving_nodes: List[NodeSpec] = []
        for n in self.spec.nodes:
            n_shapes = decompiler.find_shapes_by_prefix(f"node:{n.id}")
            if not n_shapes:
                deleted_node_ids.add(n.id)
                changes.append(ReconciledChange(
                    change_type="pruned_node",
                    element_type="node",
                    element_id=n.id,
                    old_value=n.title,
                    new_value=None,
                    description=f"Node '{n.id}' ('{n.title}') was removed in PPTX",
                ))
                continue

            if n.container_id and n.container_id in deleted_container_ids:
                n.container_id = None
                self._update_node_dict(n.id, "container_id", None)

            surviving_nodes.append(n)

            # Check geometric position & dimensions
            card_shape = decompiler.get_shape(f"node:{n.id}")
            if card_shape and hasattr(card_shape, "left") and hasattr(card_shape, "top") and hasattr(card_shape, "width") and hasattr(card_shape, "height"):
                cur_x_in = float(card_shape.left) / 914400.0
                cur_y_in = float(card_shape.top) / 914400.0
                cur_w_in = float(card_shape.width) / 914400.0
                cur_h_in = float(card_shape.height) / 914400.0

                for prop, cur_in, orig_val in [
                    ("x", cur_x_in, n.x),
                    ("y", cur_y_in, n.y),
                    ("w", cur_w_in, n.w),
                ]:
                    orig_in = self._to_inches(orig_val)
                    if abs(cur_in - orig_in) >= self.snap_threshold:
                        snapped_val = self._snap_spec_coord(cur_in)
                        if abs(snapped_val - orig_val) > 1e-4:
                            changes.append(ReconciledChange(
                                change_type="mutated_geometry",
                                element_type="node",
                                element_id=n.id,
                                field_name=prop,
                                old_value=orig_val,
                                new_value=snapped_val,
                                description=f"Node '{n.id}' {prop.upper()} adjusted from {orig_val} to {snapped_val}",
                            ))
                            setattr(n, prop, snapped_val)
                            self._update_node_dict(n.id, prop, snapped_val)

                # Node H: check against n.h (or compiled default 1.25 if n.h is None)
                eff_h = n.h if n.h is not None else 1.25
                eff_h_in = self._to_inches(eff_h)
                if abs(cur_h_in - eff_h_in) >= self.snap_threshold:
                    snapped_h = self._snap_spec_coord(cur_h_in)
                    if abs(snapped_h - eff_h) > 1e-4:
                        changes.append(ReconciledChange(
                            change_type="mutated_geometry",
                            element_type="node",
                            element_id=n.id,
                            field_name="h",
                            old_value=n.h,
                            new_value=snapped_h,
                            description=f"Node '{n.id}' H adjusted from {eff_h} to {snapped_h}",
                        ))
                        n.h = snapped_h
                        self._update_node_dict(n.id, "h", snapped_h)

            # Extract text from node text frames or card shape
            text_shapes = [
                s for s in decompiler.find_shapes_by_prefix(f"node:{n.id}:text")
                + decompiler.find_shapes_by_prefix(f"node:{n.id}:title")
                if s.has_text_frame
            ]
            if not text_shapes:
                card_fallback = decompiler.get_shape(f"node:{n.id}")
                if card_fallback and card_fallback.has_text_frame:
                    text_shapes = [card_fallback]

            if text_shapes:
                paras: List[str] = []
                for ts in text_shapes:
                    paras.extend(decompiler.get_text_paragraphs(ts))
                if paras:
                    remaining_paras = list(paras)

                    # 1. Kicker extraction
                    if n.kicker is not None and remaining_paras:
                        new_kicker = remaining_paras.pop(0)
                        clean_new_kicker = _clean_text(new_kicker) or ""
                        clean_old_kicker = _clean_text(n.kicker) or ""
                        if clean_new_kicker != clean_old_kicker and clean_new_kicker != clean_old_kicker.upper():
                            changes.append(ReconciledChange(
                                change_type="mutated_text",
                                element_type="node",
                                element_id=n.id,
                                field_name="kicker",
                                old_value=n.kicker,
                                new_value=new_kicker,
                                description=f"Node '{n.id}' kicker mutated from '{n.kicker}' to '{new_kicker}'",
                            ))
                            n.kicker = new_kicker
                            self._update_node_dict(n.id, "kicker", new_kicker)

                    # 2. Separate bullet items from body paragraphs
                    text_paras: List[str] = []
                    bullet_paras: List[str] = []
                    for idx, p in enumerate(remaining_paras):
                        stripped = p.strip()
                        # Only classify as bullet if not a single-line title or if existing node has bullets/multi-line
                        is_single_title = len(remaining_paras) == 1 and n.bullets is None and (n.title or "").startswith("•")
                        if stripped.startswith("•") and not is_single_title:
                            bullet_paras.append(stripped.lstrip("•").strip())
                        else:
                            text_paras.append(p)

                    # 3. Partition title vs subtitle respecting multi-line titles
                    orig_title_lines = [s for s in n.title.split("\n") if s.strip()] if n.title else []
                    expected_title_lines = max(1, len(orig_title_lines))

                    if len(text_paras) <= expected_title_lines:
                        new_title = "\n".join(text_paras) if text_paras else n.title
                        new_sub = None
                    else:
                        new_title = "\n".join(text_paras[:expected_title_lines])
                        new_sub = "\n".join(text_paras[expected_title_lines:])

                    # 4. Update title
                    if new_title and _clean_text(new_title) != _clean_text(n.title):
                        changes.append(ReconciledChange(
                            change_type="mutated_text",
                            element_type="node",
                            element_id=n.id,
                            field_name="title",
                            old_value=n.title,
                            new_value=new_title,
                            description=f"Node '{n.id}' title mutated from '{n.title}' to '{new_title}'",
                        ))
                        n.title = new_title
                        self._update_node_dict(n.id, "title", new_title)

                    # 5. Update subtitle
                    if new_sub:
                        if _clean_text(new_sub) != _clean_text(n.subtitle or ""):
                            changes.append(ReconciledChange(
                                change_type="mutated_text",
                                element_type="node",
                                element_id=n.id,
                                field_name="subtitle",
                                old_value=n.subtitle,
                                new_value=new_sub,
                                description=f"Node '{n.id}' subtitle mutated from '{n.subtitle}' to '{new_sub}'",
                            ))
                            n.subtitle = new_sub
                            self._update_node_dict(n.id, "subtitle", new_sub)
                    elif n.subtitle is not None:
                        changes.append(ReconciledChange(
                            change_type="mutated_text",
                            element_type="node",
                            element_id=n.id,
                            field_name="subtitle",
                            old_value=n.subtitle,
                            new_value=None,
                            description=f"Node '{n.id}' subtitle cleared in PPTX",
                        ))
                        n.subtitle = None
                        self._update_node_dict(n.id, "subtitle", None)

                    # 6. Update bullets
                    if bullet_paras:
                        old_bullets = n.bullets or []
                        if bullet_paras != old_bullets:
                            changes.append(ReconciledChange(
                                change_type="mutated_text",
                                element_type="node",
                                element_id=n.id,
                                field_name="bullets",
                                old_value=n.bullets,
                                new_value=bullet_paras,
                                description=f"Node '{n.id}' bullets mutated in PPTX",
                            ))
                            n.bullets = bullet_paras
                            self._update_node_dict(n.id, "bullets", bullet_paras)
                    elif n.bullets is not None:
                        changes.append(ReconciledChange(
                            change_type="mutated_text",
                            element_type="node",
                            element_id=n.id,
                            field_name="bullets",
                            old_value=n.bullets,
                            new_value=None,
                            description=f"Node '{n.id}' bullets cleared in PPTX",
                        ))
                        n.bullets = None
                        self._update_node_dict(n.id, "bullets", None)

        if deleted_node_ids:
            self.spec.nodes = surviving_nodes
            if "nodes" in self.raw_data:
                self.raw_data["nodes"] = [
                    nd for nd in self.raw_data["nodes"] if nd.get("id") not in deleted_node_ids
                ]

        # ----------------------------------------------------------------------
        # 4. Routed Flows / Connectors Reconciliation
        # ----------------------------------------------------------------------
        surviving_flows: List[FlowSpec] = []
        deleted_endpoint_ids = deleted_node_ids | deleted_container_ids
        pruned_flow_indices: Set[int] = set()

        for i, f in enumerate(self.spec.flows):
            flow_id = f.id or (f"{f.source}->{f.target}" if f.source and f.target else f"flow_{i}")

            # Check if flow endpoints (node or container) were deleted
            if (f.source and f.source in deleted_endpoint_ids) or (f.target and f.target in deleted_endpoint_ids):
                deleted_flow_ids.add(flow_id)
                pruned_flow_indices.add(i)
                culprit = f.source if f.source in deleted_endpoint_ids else f.target
                changes.append(ReconciledChange(
                    change_type="pruned_flow",
                    element_type="flow",
                    element_id=flow_id,
                    old_value=f"{f.source} -> {f.target}",
                    new_value=None,
                    description=f"Flow '{flow_id}' pruned because connected endpoint '{culprit}' was deleted",
                ))
                continue

            # Check if connector line shape exists
            line_shape = decompiler.get_shape(f"flow:{flow_id}")
            if not line_shape:
                deleted_flow_ids.add(flow_id)
                pruned_flow_indices.add(i)
                changes.append(ReconciledChange(
                    change_type="pruned_flow",
                    element_type="flow",
                    element_id=flow_id,
                    old_value=f"{f.source} -> {f.target}",
                    new_value=None,
                    description=f"Flow connector '{flow_id}' was removed in PPTX",
                ))
                continue

            surviving_flows.append(f)

            # Check label pill shape
            label_shape = decompiler.get_shape(f"flow:{flow_id}:label")
            if label_shape:
                l_paras = decompiler.get_text_paragraphs(label_shape)
                if l_paras:
                    new_label = l_paras[0]
                    if new_label != (f.label or "").strip():
                        changes.append(ReconciledChange(
                            change_type="mutated_text",
                            element_type="flow",
                            element_id=flow_id,
                            field_name="label",
                            old_value=f.label,
                            new_value=new_label,
                            description=f"Flow '{flow_id}' label mutated from '{f.label}' to '{new_label}'",
                        ))
                        f.label = new_label
                        self._update_flow_dict(flow_id, i, "label", new_label)
            elif f.label is not None:
                # Label pill shape was removed
                changes.append(ReconciledChange(
                    change_type="mutated_text",
                    element_type="flow",
                    element_id=flow_id,
                    field_name="label",
                    old_value=f.label,
                    new_value=None,
                    description=f"Flow '{flow_id}' label pill was deleted in PPTX",
                ))
                f.label = None
                self._update_flow_dict(flow_id, i, "label", None)

        if deleted_flow_ids:
            self.spec.flows = surviving_flows
            if "flows" in self.raw_data:
                self.raw_data["flows"] = [
                    fd for i, fd in enumerate(self.raw_data["flows"])
                    if i not in pruned_flow_indices
                    and (fd.get("id") or (f"{fd.get('source')}->{fd.get('target')}" if fd.get("source") and fd.get("target") else f"flow_{i}")) not in deleted_flow_ids
                ]

        # ----------------------------------------------------------------------
        # 5. Badges Reconciliation
        # ----------------------------------------------------------------------
        surviving_badges: List[BadgeSpec] = []
        deleted_badge_ids: Set[str] = set()

        for i, b in enumerate(self.spec.badges):
            badge_id = b.id or f"badge_{i}"
            b_shape = decompiler.get_shape(f"badge:{badge_id}")
            if not b_shape:
                deleted_badge_ids.add(badge_id)
                changes.append(ReconciledChange(
                    change_type="pruned_badge",
                    element_type="badge",
                    element_id=badge_id,
                    old_value=b.text,
                    new_value=None,
                    description=f"Badge '{badge_id}' ('{b.text}') was removed in PPTX",
                ))
                continue

            surviving_badges.append(b)

            # Check geometric position & dimensions
            if hasattr(b_shape, "left") and hasattr(b_shape, "top") and hasattr(b_shape, "width") and hasattr(b_shape, "height"):
                cur_x_in = float(b_shape.left) / 914400.0
                cur_y_in = float(b_shape.top) / 914400.0
                cur_w_in = float(b_shape.width) / 914400.0
                cur_h_in = float(b_shape.height) / 914400.0

                for prop, cur_in, orig_val in [
                    ("x", cur_x_in, b.x),
                    ("y", cur_y_in, b.y),
                    ("w", cur_w_in, b.w),
                    ("h", cur_h_in, b.h),
                ]:
                    orig_in = self._to_inches(orig_val)
                    if abs(cur_in - orig_in) >= self.snap_threshold:
                        snapped_val = self._snap_spec_coord(cur_in)
                        if abs(snapped_val - orig_val) > 1e-4:
                            changes.append(ReconciledChange(
                                change_type="mutated_geometry",
                                element_type="badge",
                                element_id=badge_id,
                                field_name=prop,
                                old_value=orig_val,
                                new_value=snapped_val,
                                description=f"Badge '{badge_id}' {prop.upper()} adjusted from {orig_val} to {snapped_val}",
                            ))
                            setattr(b, prop, snapped_val)
                            self._update_badge_dict(badge_id, i, prop, snapped_val)

            b_paras = decompiler.get_text_paragraphs(b_shape)
            if b_paras:
                new_text = b_paras[0]
                if new_text != b.text:
                    changes.append(ReconciledChange(
                        change_type="mutated_text",
                        element_type="badge",
                        element_id=badge_id,
                        field_name="text",
                        old_value=b.text,
                        new_value=new_text,
                        description=f"Badge '{badge_id}' text mutated from '{b.text}' to '{new_text}'",
                    ))
                    b.text = new_text
                    self._update_badge_dict(badge_id, i, "text", new_text)

        if deleted_badge_ids:
            self.spec.badges = surviving_badges
            if "badges" in self.raw_data:
                self.raw_data["badges"] = [
                    bd for i, bd in enumerate(self.raw_data["badges"])
                    if (bd.get("id") or f"badge_{i}") not in deleted_badge_ids
                ]

        # ----------------------------------------------------------------------
        # 6. Bullet Lists Reconciliation
        # ----------------------------------------------------------------------
        surviving_bl: List[BulletListSpec] = []
        deleted_bl_ids: Set[str] = set()

        for i, bl in enumerate(self.spec.bullet_lists):
            bl_id = bl.id or f"bullet_list_{i}"
            bl_shape = decompiler.get_shape(f"bullet_list:{bl_id}")
            if not bl_shape:
                deleted_bl_ids.add(bl_id)
                changes.append(ReconciledChange(
                    change_type="pruned_bullet_list",
                    element_type="bullet_list",
                    element_id=bl_id,
                    old_value=str(bl.items),
                    new_value=None,
                    description=f"Bullet list '{bl_id}' was removed in PPTX",
                ))
                continue

            surviving_bl.append(bl)

            # Check geometric position & dimensions
            if hasattr(bl_shape, "left") and hasattr(bl_shape, "top") and hasattr(bl_shape, "width") and hasattr(bl_shape, "height"):
                cur_x_in = float(bl_shape.left) / 914400.0
                cur_y_in = float(bl_shape.top) / 914400.0
                cur_w_in = float(bl_shape.width) / 914400.0
                cur_h_in = float(bl_shape.height) / 914400.0

                for prop, cur_in, orig_val in [
                    ("x", cur_x_in, bl.x),
                    ("y", cur_y_in, bl.y),
                    ("w", cur_w_in, bl.w),
                    ("h", cur_h_in, bl.h),
                ]:
                    orig_in = self._to_inches(orig_val)
                    if abs(cur_in - orig_in) >= self.snap_threshold:
                        snapped_val = self._snap_spec_coord(cur_in)
                        if abs(snapped_val - orig_val) > 1e-4:
                            changes.append(ReconciledChange(
                                change_type="mutated_geometry",
                                element_type="bullet_list",
                                element_id=bl_id,
                                field_name=prop,
                                old_value=orig_val,
                                new_value=snapped_val,
                                description=f"Bullet list '{bl_id}' {prop.upper()} adjusted from {orig_val} to {snapped_val}",
                            ))
                            setattr(bl, prop, snapped_val)
                            self._update_bullet_list_dict(bl_id, i, prop, snapped_val)
            raw_paras = decompiler.get_text_paragraphs(bl_shape)
            cleaned_items = [_BULLET_PREFIX_RE.sub("", p).strip() for p in raw_paras if p.strip()]
            if cleaned_items != bl.items:
                changes.append(ReconciledChange(
                    change_type="mutated_text",
                    element_type="bullet_list",
                    element_id=bl_id,
                    field_name="items",
                    old_value=bl.items,
                    new_value=cleaned_items,
                    description=f"Bullet list '{bl_id}' items mutated",
                ))
                bl.items = cleaned_items
                self._update_bullet_list_dict(bl_id, i, "items", cleaned_items)

        if deleted_bl_ids:
            self.spec.bullet_lists = surviving_bl
            if "bullet_lists" in self.raw_data:
                self.raw_data["bullet_lists"] = [
                    bld for i, bld in enumerate(self.raw_data["bullet_lists"])
                    if (bld.get("id") or f"bullet_list_{i}") not in deleted_bl_ids
                ]

        # ----------------------------------------------------------------------
        # 7. Persistence & Automatic 4-Way Re-compilation Gate
        # ----------------------------------------------------------------------
        res = ReconciliationResult(
            spec_path=self.spec_path,
            pptx_path=self.pptx_path,
            slide_index=self.slide_index,
            changes=changes,
            updated_spec=self.spec,
        )

        if changes and not dry_run:
            # Validate updated dictionary structure against Pydantic model
            validated_spec = DiagramSpec.model_validate(self.raw_data)
            self.spec = validated_spec
            res.updated_spec = validated_spec
            dump_diagram_yaml(self.raw_data, self.spec_path)

            if not no_recompile:
                effective_formats = list(formats)
                if skip_pptx:
                    if "all" in effective_formats:
                        effective_formats = ["drawio", "svg", "png"]
                    else:
                        effective_formats = [f for f in effective_formats if f.lower() != "pptx"]
                if effective_formats:
                    recompile_res = compile_diagram(
                        spec_path_or_dict=validated_spec,
                        output_dir_or_file=self.pptx_path,
                        formats=effective_formats,
                    )
                    res.recompiled = True
                    res.recompile_result = recompile_res

        return res

    def _update_node_dict(self, node_id: str, field_name: str, new_val: Any) -> None:
        if "nodes" not in self.raw_data or not isinstance(self.raw_data["nodes"], list):
            return
        for nd in self.raw_data["nodes"]:
            if isinstance(nd, dict) and nd.get("id") == node_id:
                if new_val is None:
                    nd.pop(field_name, None)
                else:
                    nd[field_name] = new_val
                break

    def _update_flow_dict(self, flow_id: str, index: int, field_name: str, new_val: Any) -> None:
        flows_list = self.raw_data.get("flows")
        if not isinstance(flows_list, list):
            return
        if 0 <= index < len(flows_list) and isinstance(flows_list[index], dict):
            fd = flows_list[index]
            if new_val is None:
                fd.pop(field_name, None)
            else:
                fd[field_name] = new_val
            return
        for i, fd in enumerate(flows_list):
            if isinstance(fd, dict):
                fid = fd.get("id") or (f"{fd.get('source')}->{fd.get('target')}" if fd.get("source") and fd.get("target") else f"flow_{i}")
                if fid == flow_id:
                    if new_val is None:
                        fd.pop(field_name, None)
                    else:
                        fd[field_name] = new_val
                    break

    def _update_container_dict(self, container_id: str, field_name: str, new_val: Any) -> None:
        if "containers" not in self.raw_data or not isinstance(self.raw_data["containers"], list):
            return
        for cd in self.raw_data["containers"]:
            if isinstance(cd, dict) and cd.get("id") == container_id:
                if new_val is None:
                    cd.pop(field_name, None)
                else:
                    cd[field_name] = new_val
                break

    def _update_badge_dict(self, badge_id: str, index: int, field_name: str, new_val: Any) -> None:
        badges_list = self.raw_data.get("badges")
        if not isinstance(badges_list, list):
            return
        if 0 <= index < len(badges_list) and isinstance(badges_list[index], dict):
            bd = badges_list[index]
            if new_val is None:
                bd.pop(field_name, None)
            else:
                bd[field_name] = new_val
            return
        for i, bd in enumerate(badges_list):
            if isinstance(bd, dict):
                bid = bd.get("id") or f"badge_{i}"
                if bid == badge_id:
                    if new_val is None:
                        bd.pop(field_name, None)
                    else:
                        bd[field_name] = new_val
                    break

    def _update_bullet_list_dict(self, bl_id: str, index: int, field_name: str, new_val: Any) -> None:
        bl_list = self.raw_data.get("bullet_lists")
        if not isinstance(bl_list, list):
            return
        if 0 <= index < len(bl_list) and isinstance(bl_list[index], dict):
            bld = bl_list[index]
            if new_val is None:
                bld.pop(field_name, None)
            else:
                bld[field_name] = new_val
            return
        for i, bld in enumerate(bl_list):
            if isinstance(bld, dict):
                bid = bld.get("id") or f"bullet_list_{i}"
                if bid == bl_id:
                    if new_val is None:
                        bld.pop(field_name, None)
                    else:
                        bld[field_name] = new_val
                    break


def reconcile_diagram(
    spec_path: Union[str, Path],
    pptx_path: Optional[Union[str, Path]] = None,
    slide_index: int = 1,
    dry_run: bool = False,
    no_recompile: bool = False,
    formats: Sequence[str] = ("all",),
    skip_pptx: bool = False,
    snap_grid: float = 0.05,
    snap_threshold: float = 0.02,
) -> ReconciliationResult:
    """Public facade function to run diagram reconciliation and return the structured result."""
    reconciler = DiagramReconciler(
        spec_path=spec_path,
        pptx_path=pptx_path,
        slide_index=slide_index,
        snap_grid=snap_grid,
        snap_threshold=snap_threshold,
    )
    return reconciler.reconcile(
        dry_run=dry_run,
        no_recompile=no_recompile,
        formats=formats,
        skip_pptx=skip_pptx,
    )
