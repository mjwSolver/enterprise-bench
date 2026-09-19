"""
Snowflake Cloud Sizing & Financial Model Engine
===============================================
Automates infrastructure parameterization, credit modeling, and formula-preserving
injection into Cloud_Sizing_Calculator_Template.xlsx.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, Union
import openpyxl
from pydantic import BaseModel, Field

from src.core.config import OUTPUT_DIR, get_template_path, validate_clean_path


# Credit Consumption Reference Table (Snowflake Standard)
WAREHOUSE_CREDIT_RATES: Dict[str, float] = {
    "XS": 1.0,
    "S": 2.0,
    "M": 4.0,
    "L": 8.0,
    "XL": 16.0,
    "2XL": 32.0,
    "3XL": 64.0,
    "4XL": 128.0,
}

# AWS Jakarta Credit Pricing ($ / credit)
TIER_PRICES_PER_CREDIT: Dict[str, float] = {
    "standard": 2.50,
    "premier": 2.80,
    "enterprise": 3.70,
    "business_critical": 5.00,
}


class WarehouseSpec(BaseModel):
    """Specification for a Snowflake Virtual Warehouse."""
    name: str = Field(..., description="Warehouse purpose or name")
    size: str = Field("XS", description="T-shirt size (XS, S, M, L, XL, 2XL)")
    days_per_month: int = Field(22, ge=0, le=31, description="Active working days per month")
    hours_per_day: float = Field(8.0, ge=0.0, le=24.0, description="Active running hours per day")
    months_per_year: int = Field(12, ge=1, le=12, description="Operating months per year")


class CloudSizingConfig(BaseModel):
    """Full parameterization for Cloud Sizing Financial Estimation."""
    client_name: str = Field("Enterprise Client", description="Client or company name")
    platform: str = Field("Snowflake on AWS (Jakarta)", description="Cloud platform and region")
    warehouses: List[WarehouseSpec] = Field(
        default_factory=lambda: [
            WarehouseSpec(name="Data Ingestion & Pipelines", size="XS", days_per_month=30, hours_per_day=4.0),
            WarehouseSpec(name="dbt Transformations & Staging", size="S", days_per_month=22, hours_per_day=6.0),
            WarehouseSpec(name="BI Reporting & Dashboard", size="XS", days_per_month=22, hours_per_day=10.0),
            WarehouseSpec(name="Ad-hoc Analytics & Development", size="XS", days_per_month=22, hours_per_day=4.0),
        ]
    )
    storage_tb: float = Field(1.0, ge=0.0, description="Provisioned monthly data storage in TB")
    storage_price_per_tb: float = Field(25.0, description="Price per TB per month in USD")
    cortex_ai_annual_usd: float = Field(3660.0, description="Snowflake Cortex AI annual allocation USD")
    training_seats_usd: float = Field(1500.0, description="Mandatory Snowflake training cost USD")
    wht_rate: float = Field(0.10, description="Withholding Tax (WHT) rate (0.10 = 10%)")
    exchange_rate_idr: float = Field(16500.0, description="USD to IDR currency conversion rate")


@dataclass
class TierCostSummary:
    tier_name: str
    price_per_credit: float
    annual_compute_usd: float
    annual_storage_usd: float
    annual_cortex_usd: float
    total_subscription_usd: float
    wht_tax_usd: float
    training_usd: float
    grand_total_usd: float
    grand_total_idr: float


@dataclass
class CloudSizingResult:
    config: CloudSizingConfig
    monthly_credits: float
    annual_credits: float
    annual_storage_usd: float
    tiers: Dict[str, TierCostSummary]
    output_path: Path

    def summary_table(self) -> str:
        """Produce clean ASCII financial summary."""
        lines = [
            f"Snowflake Cloud Sizing Financial Summary — {self.config.client_name}",
            f"Platform: {self.config.platform} | FX Rate: 1 USD = {self.config.exchange_rate_idr:,.0f} IDR",
            f"Compute Credits: {self.monthly_credits:,.1f} / mo | {self.annual_credits:,.1f} / yr",
            f"Storage: {self.config.storage_tb} TB (${self.annual_storage_usd:,.2f}/yr) | AI/Cortex: ${self.config.cortex_ai_annual_usd:,.2f}/yr",
            "-" * 80,
            f"{'Edition':<18} | {'Compute ($)':<12} | {'Total Subs ($)':<14} | {'Grand Total ($)':<15} | {'Grand Total (IDR)':<20}",
            "-" * 80,
        ]
        for key in ["standard", "enterprise", "business_critical"]:
            if key in self.tiers:
                t = self.tiers[key]
                lines.append(
                    f"{t.tier_name:<18} | ${t.annual_compute_usd:>10,.2f} | ${t.total_subscription_usd:>12,.2f} | "
                    f"${t.grand_total_usd:>13,.2f} | Rp {t.grand_total_idr:>17,.0f}"
                )
        lines.append("-" * 80)
        return "\n".join(lines)


def calculate_cloud_sizing(
    config: Optional[CloudSizingConfig] = None,
    template_path: Optional[Union[str, Path]] = None,
    output_path: Optional[Union[str, Path]] = None,
) -> CloudSizingResult:
    """
    Parameterize and calculate Snowflake infrastructure sizing, preserving all template formulas.
    """
    cfg = config or CloudSizingConfig()

    if len(cfg.warehouses) > 8:
        raise ValueError(
            f"Snowflake Cloud Sizing template supports maximum 8 warehouse configurations; "
            f"received {len(cfg.warehouses)}. Please consolidate warehouse tiers."
        )

    p = template_path or get_template_path("Cloud_Sizing_Calculator_Template.xlsx")
    if not p or not Path(p).exists():
        raise FileNotFoundError("Cloud_Sizing_Calculator_Template.xlsx template not found.")

    wb = openpyxl.load_workbook(str(p), data_only=False)
    ws = wb["Calculator (live)"]

    # 1. Inject Header Metadata
    ws["C2"] = cfg.client_name
    ws["C4"] = cfg.platform

    # 2. Inject Warehouse Configurations (Rows 7 to 14)
    # Clear / zero out all 8 available slots first
    monthly_credits = 0.0
    annual_credits = 0.0

    for slot_idx in range(8):
        row = 7 + slot_idx
        if slot_idx < len(cfg.warehouses):
            wh = cfg.warehouses[slot_idx]
            size_upper = wh.size.upper()
            rate = WAREHOUSE_CREDIT_RATES.get(size_upper, 1.0)

            ws[f"C{row}"] = wh.name
            ws[f"D{row}"] = size_upper
            ws[f"F{row}"] = wh.days_per_month
            ws[f"G{row}"] = wh.hours_per_day
            ws[f"H{row}"] = wh.months_per_year

            wh_monthly = rate * wh.days_per_month * wh.hours_per_day
            wh_annual = rate * wh.hours_per_day * (wh.days_per_month * wh.months_per_year)
            monthly_credits += wh_monthly
            annual_credits += wh_annual
        else:
            ws[f"C{row}"] = ""
            ws[f"D{row}"] = "XS"
            ws[f"F{row}"] = 0
            ws[f"G{row}"] = 0
            ws[f"H{row}"] = 0

    # 3. Inject Storage Parameters
    ws["I19"] = cfg.storage_tb
    ws["H18"] = cfg.storage_price_per_tb
    annual_storage_usd = cfg.storage_tb * cfg.storage_price_per_tb * 12

    # 4. Inject AI & Training Parameters
    ws["F32"] = cfg.cortex_ai_annual_usd
    ws["F35"] = cfg.training_seats_usd

    # 5. Inject Taxes & Foreign Exchange
    ws["H39"] = cfg.wht_rate
    ws["K39"] = cfg.exchange_rate_idr

    # 6. Save Workbook
    out = Path(output_path) if output_path else OUTPUT_DIR / f"calculated_Cloud_Sizing_{cfg.client_name.replace(' ', '_')}.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    wb.close()

    # 7. Compute Python-Side Mathematical Verification
    tiers: Dict[str, TierCostSummary] = {}
    tier_labels = {
        "standard": "Standard",
        "premier": "Premier",
        "enterprise": "Enterprise",
        "business_critical": "Business Critical",
    }

    for t_key, label in tier_labels.items():
        price = TIER_PRICES_PER_CREDIT.get(t_key, 3.70)
        compute_usd = annual_credits * price
        total_subs = compute_usd + annual_storage_usd + cfg.cortex_ai_annual_usd
        wht_usd = total_subs * cfg.wht_rate
        grand_total_usd = total_subs + wht_usd + cfg.training_seats_usd
        grand_total_idr = grand_total_usd * cfg.exchange_rate_idr

        tiers[t_key] = TierCostSummary(
            tier_name=label,
            price_per_credit=price,
            annual_compute_usd=compute_usd,
            annual_storage_usd=annual_storage_usd,
            annual_cortex_usd=cfg.cortex_ai_annual_usd,
            total_subscription_usd=total_subs,
            wht_tax_usd=wht_usd,
            training_usd=cfg.training_seats_usd,
            grand_total_usd=grand_total_usd,
            grand_total_idr=grand_total_idr,
        )

    return CloudSizingResult(
        config=cfg,
        monthly_credits=monthly_credits,
        annual_credits=annual_credits,
        annual_storage_usd=annual_storage_usd,
        tiers=tiers,
        output_path=out,
    )
