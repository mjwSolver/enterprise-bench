"""
Timeline Aggregator & S-Curve Synchronizer
==========================================
Binds 4.2_Weekly_Progress_Timeline_Update_Template.xlsx directly to the
SCurveGenerator to extract daily planned/actual progression, compute weekly
cumulative milestones, and embed executive S-curve line charts.
"""

from __future__ import annotations

from datetime import date, datetime, timedelta
import math
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, Union
import openpyxl

from src.core.config import OUTPUT_DIR, get_template_path, validate_clean_path
from src.xlsx_engine.s_curve_generator import (
    SCurveGenerator,
    SCurvePoint,
    VarianceAnalysis,
)


def _format_date(val: Any) -> str:
    """Format excel date value into ISO YYYY-MM-DD."""
    if val is None:
        return ""
    if isinstance(val, (datetime, date)):
        return val.strftime("%Y-%m-%d")
    s = str(val).strip()
    if not s:
        return ""
    head = s[:10]
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(head, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return head if len(s) >= 10 else s


def _short_date_label(date_str: str) -> str:
    """Extract compact MM-DD label from ISO YYYY-MM-DD (or first 5 chars for custom strings)."""
    if len(date_str) >= 10 and date_str[4] in ("-", "/") and date_str[:4].isdigit():
        return date_str[5:10]
    return date_str[:5]


def _coerce_float(val: Any, default: Optional[float] = 0.0) -> Optional[float]:
    """Safely convert numeric Excel cell value to float, rejecting booleans and NaN/Inf."""
    if val is None or isinstance(val, bool):
        return default
    if isinstance(val, (int, float)):
        f_val = float(val)
        return default if (math.isnan(f_val) or math.isinf(f_val)) else f_val
    if isinstance(val, str):
        stripped = val.strip()
        if not stripped:
            return default
        try:
            f_val = float(stripped)
            return default if (math.isnan(f_val) or math.isinf(f_val)) else f_val
        except ValueError:
            return default
    return default


class TimelineAggregator:
    """
    Extracts and aggregates daily schedule tracking into weekly cumulative S-curves.
    """

    def __init__(self, timeline_path: Optional[Union[str, Path]] = None):
        if timeline_path is not None:
            candidate = Path(timeline_path)
            validate_clean_path(candidate)
            if candidate.exists():
                resolved = validate_clean_path(candidate)
            else:
                found = get_template_path(str(timeline_path))
                if not found or not Path(found).exists():
                    raise FileNotFoundError(
                        f"4.2_Weekly_Progress_Timeline_Update_Template.xlsx template not found: {timeline_path}"
                    )
                resolved = validate_clean_path(found)
        else:
            found = get_template_path("4.2_Weekly_Progress_Timeline_Update_Template.xlsx")
            if not found or not Path(found).exists():
                raise FileNotFoundError(
                    "4.2_Weekly_Progress_Timeline_Update_Template.xlsx template not found."
                )
            resolved = validate_clean_path(found)

        self.timeline_path: Path = Path(resolved)
        self._cached_extraction: Optional[Tuple[int, int, VarianceAnalysis]] = None

    def extract_s_curve(self, chunk_days: int = 7) -> VarianceAnalysis:
        """
        Parse daily planned (row 3) and actual (row 4) weights across date columns (col E onwards)
        and aggregate into weekly S-curve points.
        """
        if chunk_days < 1:
            raise ValueError(f"chunk_days must be >= 1, got {chunk_days}")

        try:
            mtime_ns = self.timeline_path.stat().st_mtime_ns
        except OSError:
            mtime_ns = -1

        if (
            self._cached_extraction is not None
            and self._cached_extraction[0] == chunk_days
            and self._cached_extraction[1] == mtime_ns
        ):
            return self._cached_extraction[2]

        wb = openpyxl.load_workbook(str(self.timeline_path), data_only=True)
        try:
            ws = wb.active
            if ws is None:
                raise ValueError(f"No active worksheet found in {self.timeline_path}")

            daily_points: List[Dict[str, Any]] = []

            # Fetch rows 2..5 from column E (5) onwards in a single batch read
            rows = list(
                ws.iter_rows(
                    min_row=2,
                    max_row=5,
                    min_col=5,
                    max_col=ws.max_column,
                    values_only=True,
                )
            )
            if len(rows) >= 4:
                date_row, planned_row, actual_row, baseline_row = rows[0], rows[1], rows[2], rows[3]

                # If any date cells are None (un-cached formula dates like =E2+1), inspect formula row
                formula_date_row = None
                if any(dt is None for dt in date_row):
                    try:
                        wb_form = openpyxl.load_workbook(str(self.timeline_path), data_only=False, read_only=True)
                        ws_form = wb_form.active
                        if ws_form:
                            f_rows = list(ws_form.iter_rows(min_row=2, max_row=2, min_col=5, max_col=ws.max_column, values_only=True))
                            if f_rows:
                                formula_date_row = f_rows[0]
                        wb_form.close()
                    except Exception:
                        formula_date_row = None

                last_valid_date: Optional[date] = None
                last_valid_offset: Optional[int] = None

                for offset, (raw_dt, raw_p, raw_a, raw_b) in enumerate(
                    zip(date_row, planned_row, actual_row, baseline_row), start=5
                ):
                    date_str = _format_date(raw_dt)
                    if not date_str:
                        form_val = (
                            formula_date_row[offset - 5]
                            if formula_date_row and (offset - 5) < len(formula_date_row)
                            else None
                        )
                        if form_val is not None and last_valid_date is not None and last_valid_offset is not None:
                            inferred_date = last_valid_date + timedelta(days=offset - last_valid_offset)
                            date_str = inferred_date.strftime("%Y-%m-%d")

                    if not date_str:
                        continue

                    try:
                        last_valid_date = datetime.strptime(date_str, "%Y-%m-%d").date()
                        last_valid_offset = offset
                    except Exception:
                        pass

                    p_val = _coerce_float(raw_p, default=0.0) or 0.0
                    a_val = _coerce_float(raw_a, default=None)
                    b_val = _coerce_float(raw_b, default=0.0) or 0.0

                    daily_points.append({
                        "date": date_str,
                        "col": offset,
                        "planned": p_val,
                        "actual": a_val,
                        "baseline": b_val,
                    })
        finally:
            wb.close()

        if not daily_points:
            raise ValueError(f"No date progression points found in {self.timeline_path}")

        total_planned_weight = sum(d["planned"] for d in daily_points)
        if total_planned_weight <= 0:
            total_actual_weight = sum(d["actual"] for d in daily_points if d["actual"] is not None)
            total_planned_weight = total_actual_weight if total_actual_weight > 0 else 100.0

        # Group into intervals (default 7 calendar days per week)
        periods: List[SCurvePoint] = []
        cum_p = 0.0
        cum_a = 0.0
        current_period_idx: Optional[int] = None

        for idx in range(0, len(daily_points), chunk_days):
            chunk = daily_points[idx:idx + chunk_days]
            w_num = len(periods) + 1
            w_label = f"W{w_num:02d}"

            chunk_p = sum(d["planned"] for d in chunk)
            cum_p += chunk_p
            p_pct = round(min(1.0, max(0.0, cum_p / total_planned_weight)), 4)

            actual_entries = [d["actual"] for d in chunk if d["actual"] is not None]
            has_actual = len(actual_entries) > 0

            if has_actual:
                current_period_idx = len(periods)
                chunk_a = sum(actual_entries)
                cum_a += chunk_a
                a_pct = round(min(1.0, max(0.0, cum_a / total_planned_weight)), 4)
                var = round(a_pct - p_pct, 4)
                var_pct = round(var * 100, 2)

                # Health evaluation
                if var >= 0.0:
                    health = "ON TRACK"
                elif var >= -0.02:
                    health = "AT RISK"
                elif var >= -0.05:
                    health = "DELAYED"
                else:
                    health = "CRITICAL"
            else:
                a_pct = None
                var = None
                var_pct = None
                health = "PLANNED"

            start_d = chunk[0]["date"]

            periods.append(
                SCurvePoint(
                    period=f"{w_label} ({_short_date_label(start_d)})",
                    planned_pct=p_pct,
                    actual_pct=a_pct,
                    variance=var,
                    variance_pct=var_pct,
                    health=health,
                )
            )

        # Consolidate overall variance metrics at latest active period
        curr_pt = periods[current_period_idx] if current_period_idx is not None else periods[0]
        spi = 1.0
        if curr_pt.planned_pct and curr_pt.planned_pct > 0 and curr_pt.actual_pct is not None:
            spi = round(curr_pt.actual_pct / curr_pt.planned_pct, 3)

        analysis = VarianceAnalysis(
            current_period=curr_pt.period,
            current_period_idx=current_period_idx,
            current_planned_pct=curr_pt.planned_pct,
            current_actual_pct=curr_pt.actual_pct,
            current_variance=curr_pt.variance,
            current_variance_pct=curr_pt.variance_pct,
            overall_health=curr_pt.health,
            spi=spi,
            points=periods,
        )
        self._cached_extraction = (chunk_days, mtime_ns, analysis)
        return analysis

    def synchronize_and_chart(
        self,
        output_path: Optional[Union[str, Path]] = None,
        chunk_days: int = 7,
        analysis: Optional[VarianceAnalysis] = None,
    ) -> Path:
        """
        Extract cumulative curve points, load the original timeline workbook,
        and inject an executive 'S-Curve Analysis' tab with KPI cards, summary table,
        and native openpyxl LineChart.
        """
        resolved_analysis = analysis or self.extract_s_curve(chunk_days=chunk_days)

        out_path = Path(output_path) if output_path else OUTPUT_DIR / "updated_Weekly_Progress_Timeline_SCurve.xlsx"
        out_path.parent.mkdir(parents=True, exist_ok=True)

        generator = SCurveGenerator(template=self.timeline_path)
        out_file = generator.stamp_and_chart(
            curve_data=resolved_analysis.points,
            output_path=out_path,
            sheet_name="S-Curve Analysis",
            chart_title="Data Application for Financial Analytics — Cumulative S-Curve",
            include_kpis=True,
            analysis=resolved_analysis,
        )
        return Path(out_file)


def synchronize_timeline_s_curve(
    timeline_path: Optional[Union[str, Path]] = None,
    output_path: Optional[Union[str, Path]] = None,
) -> Tuple[Path, VarianceAnalysis]:
    """Helper to extract and chart S-curve from project timeline update spreadsheet."""
    agg = TimelineAggregator(timeline_path=timeline_path)
    analysis = agg.extract_s_curve()
    saved_path = agg.synchronize_and_chart(output_path=output_path, analysis=analysis)
    return saved_path, analysis
