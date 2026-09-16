"""
RAID Log & Spreadsheet Ledger Subsystem
=======================================
Pydantic schemas and automated row injection for enterprise Risk Registers,
Issue Logs, Defect Trackers, Stakeholder Registers, and Closeout Checklists
without corrupting formulas, styles, or named ranges.
"""

from __future__ import annotations

from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Union
import openpyxl
from pydantic import BaseModel, Field

from src.core.config import OUTPUT_DIR, get_template_path, validate_clean_path


def _copy_cell_style(src_cell: Any, tgt_cell: Any) -> None:
    """Copy formatting, border, fill, and font from a reference cell."""
    if getattr(src_cell, "has_style", False):
        tgt_cell.font = copy(src_cell.font)
        tgt_cell.border = copy(src_cell.border)
        tgt_cell.fill = copy(src_cell.fill)
        tgt_cell.number_format = src_cell.number_format
        tgt_cell.protection = copy(src_cell.protection)
        tgt_cell.alignment = copy(src_cell.alignment)


# ============================================================================
# 1. Risk Register Models & Ingestion
# ============================================================================

class RiskEntry(BaseModel):
    """Payload representing a project risk item."""
    title: str = Field(..., description="Short title of the risk")
    description: str = Field(..., description="Root cause and event description")
    impact_statement: str = Field(..., description="Operational and schedule impact")
    category: str = Field("Scope", description="Scope, Data, Resource, Schedule, Quality")
    owner: str = Field("Project Manager", description="Role or name accountable for mitigation")
    probability: str = Field("Medium", description="Low, Medium, High")
    impact_level: str = Field("Medium", description="Low, Medium, High")
    status: str = Field("Open", description="Open, Mitigated, Closed")
    mitigation_plan: str = Field("", description="Mitigation or preventive actions")
    contingency_plan: str = Field("", description="Contingency plan if risk event triggers")
    mitigation_date: Optional[str] = Field(None, description="Date mitigation executed or verified")
    date_identified: str = Field(default_factory=lambda: datetime.now().strftime("%Y-%m-%d"))


def append_risk_to_register(
    risk: RiskEntry,
    template_path: Optional[Union[str, Path]] = None,
    output_path: Optional[Union[str, Path]] = None,
) -> Path:
    """Safely append a new risk row to Risk Register workbook preserving formulas and styles."""
    p = template_path or get_template_path("4.5_Risk_Register_Template.xlsx")
    if not p or not Path(p).exists():
        raise FileNotFoundError("4.5_Risk_Register_Template.xlsx template not found.")

    wb = openpyxl.load_workbook(str(p))
    ws = wb["Risk Register"]

    next_row = ws.max_row + 1
    prev_row = ws.max_row if ws.max_row >= 4 else 4

    formula_num = f'=ROWS(INDIRECT("A$4:A"&ROW()))'

    values = [
        (2, formula_num),
        (3, risk.date_identified),
        (4, risk.title),
        (5, risk.description),
        (6, risk.impact_statement),
        (7, risk.category),
        (8, risk.owner),
        (9, risk.probability),
        (10, risk.impact_level),
        (11, risk.status),
        (12, risk.mitigation_plan),
        (13, risk.contingency_plan),
        (14, risk.mitigation_date or ""),
    ]

    for col_idx, val in values:
        c = ws.cell(row=next_row, column=col_idx, value=val)
        ref_cell = ws.cell(row=prev_row, column=col_idx)
        _copy_cell_style(ref_cell, c)

    out = Path(output_path) if output_path else OUTPUT_DIR / "updated_Risk_Register.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    wb.close()
    return out


# ============================================================================
# 2. Issue Log Models & Ingestion
# ============================================================================

class IssueEntry(BaseModel):
    """Payload representing an active project issue / blocker."""
    title: str = Field(..., description="Short title of the issue")
    description: str = Field(..., description="Actual roadblock or defect description")
    category: str = Field("Technical", description="Resource, Requirement, Technical, Environment")
    owner: str = Field("Lead Data Engineer", description="PIC responsible for resolution")
    reported_by: str = Field("Project Lead", description="Originator who flagged the issue")
    impacted_area: str = Field("Timeline", description="Timeline, Scope, Budget, Resource, Quality")
    severity: str = Field("Medium", description="Low, Medium, High, Critical")
    status: str = Field("Open", description="Open, In Progress, Closed")
    root_cause: str = Field("", description="Root cause explanation")
    impact_explanation: str = Field("", description="Detailed impact assessment")
    action_plan: str = Field("", description="Resolution action steps")
    date_identified: str = Field(default_factory=lambda: datetime.now().strftime("%Y-%m-%d"))


def append_issue_to_log(
    issue: IssueEntry,
    template_path: Optional[Union[str, Path]] = None,
    output_path: Optional[Union[str, Path]] = None,
) -> Path:
    """Safely append a new issue row to Issue Log workbook preserving formulas and styles."""
    p = template_path or get_template_path("4.6_Issue_Log_Template.xlsx")
    if not p or not Path(p).exists():
        raise FileNotFoundError("4.6_Issue_Log_Template.xlsx template not found.")

    wb = openpyxl.load_workbook(str(p))
    ws = wb["Issue Log"]

    next_row = ws.max_row + 1
    prev_row = ws.max_row if ws.max_row >= 4 else 4

    formula_num = f'=ROWS(INDIRECT("A$4:A"&ROW()))'

    values = [
        (2, formula_num),
        (3, issue.date_identified),
        (4, issue.title),
        (5, issue.description),
        (6, issue.category),
        (7, issue.owner),
        (8, issue.reported_by),
        (9, issue.impacted_area),
        (10, issue.severity),
        (11, issue.status),
        (12, issue.root_cause),
        (13, issue.impact_explanation),
        (14, issue.action_plan),
    ]

    for col_idx, val in values:
        c = ws.cell(row=next_row, column=col_idx, value=val)
        ref_cell = ws.cell(row=prev_row, column=col_idx)
        _copy_cell_style(ref_cell, c)

    out = Path(output_path) if output_path else OUTPUT_DIR / "updated_Issue_Log.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    wb.close()
    return out


# ============================================================================
# 3. Defect List Models & Ingestion (3.6_Defect_List_Template.xlsx)
# ============================================================================

class DefectEntry(BaseModel):
    """Payload representing a software or pipeline defect."""
    module: str = Field("A. Sales", description="Module or component where defect occurred")
    activity: str = Field("B. SIT", description="Testing phase (A. Unit Testing, B. SIT, C. UAT)")
    summary: str = Field(..., description="Short summary of the bug/defect")
    description: str = Field(..., description="Detailed description of defect condition")
    severity: str = Field("2. Major", description="1. Critical, 2. Major, 3. Medium, 4. Low")
    steps_to_reproduce: str = Field("", description="Step-by-step reproduction instructions")
    expected_result: str = Field("", description="Expected correct behavior")
    actual_result: str = Field("", description="Observed incorrect result")
    priority: str = Field("2. Medium", description="1. High, 2. Medium, 3. Low")
    status: str = Field("Open", description="Open, Assigned, In Progress, Resolved, Closed")
    pic: str = Field("Data Dev", description="Person in charge of remediation")
    test_case: str = Field("TC-01", description="Associated Test Case ID (e.g. SLS-B02)")
    resolution_date: Optional[str] = Field(None, description="Date bug was resolved")
    image_ref: Optional[str] = Field(None, description="Image screenshot or attachment reference")
    tester_comment: Optional[str] = Field(None, description="Latest tester notes or validation comments")
    entry_date: str = Field(default_factory=lambda: datetime.now().strftime("%Y-%m-%d"))


def append_defect_to_list(
    defect: DefectEntry,
    template_path: Optional[Union[str, Path]] = None,
    output_path: Optional[Union[str, Path]] = None,
) -> Path:
    """Safely append a new defect row to 3.6_Defect_List_Template.xlsx preserving formulas and styles."""
    p = template_path or get_template_path("3.6_Defect_List_Template.xlsx")
    if not p or not Path(p).exists():
        raise FileNotFoundError("3.6_Defect_List_Template.xlsx template not found.")

    wb = openpyxl.load_workbook(str(p))
    ws = wb["Defect List"]

    next_row = ws.max_row + 1
    prev_row = ws.max_row if ws.max_row >= 3 else 3

    formula_num = f'=ROWS(INDIRECT("A$3:A"&ROW()))'

    values = [
        (2, formula_num),
        (3, defect.entry_date),
        (4, defect.module),
        (5, defect.activity),
        (6, defect.summary),
        (7, defect.description),
        (8, defect.severity),
        (9, defect.steps_to_reproduce),
        (10, defect.expected_result),
        (11, defect.actual_result),
        (12, defect.priority),
        (13, defect.status),
        (14, defect.resolution_date or ""),
        (15, defect.pic),
        (16, defect.test_case),
        (17, defect.image_ref or ""),
        (18, defect.tester_comment or ""),
    ]

    for col_idx, val in values:
        c = ws.cell(row=next_row, column=col_idx, value=val)
        ref_cell = ws.cell(row=prev_row, column=col_idx)
        _copy_cell_style(ref_cell, c)

    out = Path(output_path) if output_path else OUTPUT_DIR / "updated_Defect_List.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    wb.close()
    return out


# ============================================================================
# 4. Stakeholder Register Models & Ingestion (1.3_Stakeholders_Register_Template.xlsx)
# ============================================================================

class StakeholderEntry(BaseModel):
    """Payload representing a project stakeholder."""
    company: str = Field("Client Corp", description="Organization / Employer")
    name: str = Field(..., description="Full name and title of stakeholder")
    email: str = Field(..., description="Email address")
    phone: str = Field("+62-811-0000-000", description="Contact phone number")
    department: str = Field("Finance", description="Department or business unit")
    project_role: str = Field("Project Manager", description="Role in project governance")
    is_active: str = Field("Y", description="Active flag: Y / N")
    notes: Optional[str] = Field("", description="Additional responsibilities or context")


def append_stakeholder_to_register(
    stakeholder: StakeholderEntry,
    template_path: Optional[Union[str, Path]] = None,
    output_path: Optional[Union[str, Path]] = None,
) -> Path:
    """Safely append a stakeholder row to 1.3_Stakeholders_Register_Template.xlsx preserving formulas."""
    p = template_path or get_template_path("1.3_Stakeholders_Register_Template.xlsx")
    if not p or not Path(p).exists():
        raise FileNotFoundError("1.3_Stakeholders_Register_Template.xlsx template not found.")

    wb = openpyxl.load_workbook(str(p))
    ws = wb["Stakeholders"]

    next_row = ws.max_row + 1
    prev_row = ws.max_row if ws.max_row >= 3 else 3

    formula_num = f'=ROWS(INDIRECT("A$3:A"&ROW()))'

    values = [
        (2, formula_num),
        (3, stakeholder.company),
        (4, stakeholder.name),
        (5, stakeholder.email),
        (6, stakeholder.phone),
        (7, stakeholder.department),
        (8, stakeholder.project_role),
        (9, stakeholder.is_active),
        (10, stakeholder.notes or ""),
    ]

    for col_idx, val in values:
        c = ws.cell(row=next_row, column=col_idx, value=val)
        ref_cell = ws.cell(row=prev_row, column=col_idx)
        _copy_cell_style(ref_cell, c)

    out = Path(output_path) if output_path else OUTPUT_DIR / "updated_Stakeholders_Register.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    wb.close()
    return out


# ============================================================================
# 5. Project Closeout Checklist Models (5.2_Project_Closeout_Checklist_Template.xlsx)
# ============================================================================

class CloseoutChecklistItem(BaseModel):
    """Payload representing status update for a closeout deliverable or gate."""
    item_no: int = Field(..., description="Item number in the list")
    deliverable: str = Field(..., description="Deliverable or checklist requirement")
    status: str = Field("Delivered", description="Status (Delivered, Done, Yes, Partial, TBD)")
    section: str = Field("1. Deliverables", description="1. Deliverables, 2. Ruang Lingkup, 3. Closing Checklist")


def update_closeout_checklist(
    updates: List[Dict[str, Any]],
    template_path: Optional[Union[str, Path]] = None,
    output_path: Optional[Union[str, Path]] = None,
) -> Path:
    """
    Safely update deliverable and gate statuses in 5.2_Project_Closeout_Checklist_Template.xlsx.
    updates is a list of dicts: [{"item_name": "BAST 1 sign-off", "status": "Yes"}, ...]
    or [{"row": 4, "status": "Delivered"}, ...]
    """
    p = template_path or get_template_path("5.2_Project_Closeout_Checklist_Template.xlsx")
    if not p or not Path(p).exists():
        raise FileNotFoundError("5.2_Project_Closeout_Checklist_Template.xlsx template not found.")

    wb = openpyxl.load_workbook(str(p))
    ws = wb["Closing Checklist"]

    for upd in updates:
        if "row" in upd and "status" in upd:
            ws.cell(row=int(upd["row"]), column=3, value=upd["status"])
            continue

        item_name = str(upd.get("item_name", "")).strip().lower()
        new_status = upd.get("status", "Delivered")
        if not item_name:
            continue

        for r in range(4, ws.max_row + 1):
            val = ws.cell(r, 2).value
            if val and item_name in str(val).strip().lower():
                ws.cell(r, 3, value=new_status)

    out = Path(output_path) if output_path else OUTPUT_DIR / "updated_Project_Closeout_Checklist.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    wb.close()
    return out
