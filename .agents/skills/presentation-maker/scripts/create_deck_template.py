#!/usr/bin/env python3
"""
create_deck_template.py
A reusable utility for generating modular PowerPoint presentations (.pptx)
using modern consulting grid principles and python-pptx.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

class DeckBuilder:
    def __init__(self, output_path="presentation.pptx", width_in=13.333, height_in=7.5):
        self.output_path = output_path
        self.prs = Presentation()
        self.prs.slide_width = Inches(width_in)
        self.prs.slide_height = Inches(height_in)
        
        # Default Theme Palette (Modern Consulting)
        self.bg_color = RGBColor(248, 249, 250)
        self.card_bg = RGBColor(255, 255, 255)
        self.primary_text = RGBColor(30, 41, 59)
        self.secondary_text = RGBColor(100, 116, 139)
        self.accent_color = RGBColor(15, 118, 110)      # Deep Teal
        self.highlight_color = RGBColor(217, 119, 6)    # Amber
        self.border_color = RGBColor(226, 232, 240)

    def new_slide(self):
        blank_layout = self.prs.slide_layouts[6]
        slide = self.prs.slides.add_slide(blank_layout)
        bg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), self.prs.slide_width, self.prs.slide_height
        )
        bg.fill.solid()
        bg.fill.fore_color.rgb = self.bg_color
        bg.line.fill.background()
        return slide

    def add_header(self, slide, tracker, title, subtitle=None):
        # Tracker breadcrumb
        tb_track = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.35))
        tf_tr = tb_track.text_frame
        tf_tr.word_wrap = True
        p_tr = tf_tr.paragraphs[0]
        p_tr.text = tracker.upper()
        p_tr.font.size = Pt(10)
        p_tr.font.bold = True
        p_tr.font.color.rgb = self.accent_color

        # Title
        tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.6))
        tf_t = tb_title.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = self.primary_text

        if subtitle:
            tb_sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.45), Inches(11.7), Inches(0.4))
            tf_s = tb_sub.text_frame
            tf_s.word_wrap = True
            p_s = tf_s.paragraphs[0]
            p_s.text = subtitle
            p_s.font.size = Pt(12)
            p_s.font.color.rgb = self.secondary_text

    def add_card(self, slide, left, top, width, height, title=None, items=None, accent_top=True):
        # Card shape
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = self.card_bg
        card.line.color.rgb = self.border_color
        card.line.width = Pt(1)

        if accent_top:
            stripe = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, Inches(0.1))
            stripe.fill.solid()
            stripe.fill.fore_color.rgb = self.accent_color
            stripe.line.fill.background()

        if title or items:
            tb = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.25), width - Inches(0.5), height - Inches(0.5))
            tf = tb.text_frame
            tf.word_wrap = True

            if title:
                p_t = tf.paragraphs[0]
                p_t.text = title
                p_t.font.size = Pt(14)
                p_t.font.bold = True
                p_t.font.color.rgb = self.primary_text
                p_t.space_after = Pt(8)

            if items:
                for idx, item in enumerate(items):
                    p = tf.add_paragraph() if (title or idx > 0) else tf.paragraphs[0]
                    p.text = f"• {item}"
                    p.font.size = Pt(11.5)
                    p.font.color.rgb = RGBColor(71, 85, 105)
                    p.space_after = Pt(4)
        return card

    def save(self):
        os.makedirs(os.path.dirname(os.path.abspath(self.output_path)), exist_ok=True)
        self.prs.save(self.output_path)
        print(f"Presentation saved successfully to: {self.output_path}")

if __name__ == "__main__":
    builder = DeckBuilder("sample_output.pptx")
    slide = builder.new_slide()
    builder.add_header(slide, "Introduction | Overview", "Strategic Presentation Deck Framework", "Built with modern consulting grid standards")
    
    # 3 column cards
    w = Inches(3.64)
    h = Inches(4.8)
    builder.add_card(slide, Inches(0.8), Inches(2.0), w, h, "Pillar 1: Context", ["Industry landscape", "Macro factors", "Regulatory policy"])
    builder.add_card(slide, Inches(4.84), Inches(2.0), w, h, "Pillar 2: Research Gap", ["Conflicting prior findings", "Emerging market nuance", "Mechanism deficit"])
    builder.add_card(slide, Inches(8.88), Inches(2.0), w, h, "Pillar 3: Methodology", ["Panel data regression", "Fixed effects model", "Robustness testing"])
    
    builder.save()
