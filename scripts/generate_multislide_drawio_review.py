"""
Multi-Slide Draw.io Typographic Benchmark & Icon Visualization Generator
========================================================================
Generates a 4-slide presentation in `output/proj-xyz/presentations/Project_XYZ_DrawIO_Review.pptx`:
  - Slide 1: 18pt Font Benchmark (Unconstrained Natural Aspect Ratio)
  - Slide 2: 24pt Font Benchmark (Elevated Bold Typography)
  - Slide 3: 32pt Font Benchmark (Maximum Executive Boardroom Readability)
  - Slide 4: 4th Visualization Type — 4-Tier Zero-Trust Service Mesh Architecture
             featuring integrated Lucide icons (Shield, CPU, Database, Lock).
"""

from __future__ import annotations

import sys
from pathlib import Path
from typing import Dict, List, Optional, Tuple, Any

# Ensure repo root on sys.path
_ROOT = Path(__file__).resolve().parent.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from src.ppt_engine.consulting_archetypes import (
    add_card,
    add_slide_footer,
    add_slide_header,
    add_slide_with_background,
    create_presentation,
)
from src.ppt_engine.diagram_engine import DrawIOProject
from src.ppt_engine.theme_engine import get_theme


def build_flowchart_slide(
    prs: Presentation,
    theme: Any,
    png_path: Path,
    font_label: str,
    tracker_str: str,
    current_idx: int,
    total_slides: int = 4,
):
    slide = add_slide_with_background(prs, theme)

    # Slide Header
    add_slide_header(
        slide,
        theme,
        tracker=tracker_str,
        action_title="Zero-Trust Perimeter & Continuous Auditing Enforce SOC-2 Compliance Across 450+ Stores",
        subtitle=f"Architecture topology visualization rendered at {font_label} typography with natural unconstrained aspect ratio.",
    )

    # Section Tag
    diag_tag_tb = slide.shapes.add_textbox(Inches(0.8), Inches(1.68), Inches(11.733), Inches(0.30))
    diag_tag_tf = diag_tag_tb.text_frame
    diag_tag_tf.margin_left = diag_tag_tf.margin_right = diag_tag_tf.margin_top = diag_tag_tf.margin_bottom = 0
    p_tag = diag_tag_tf.paragraphs[0]
    p_tag.text = f"TARGET ARCHITECTURE TOPOLOGY (DRAW.IO ENGINE — {font_label.upper()} TYPOGRAPHY)"
    p_tag.font.name = theme.font_family_header
    p_tag.font.size = Pt(9.5)
    p_tag.font.bold = True
    p_tag.font.color.rgb = theme.get_rgb("accent")

    # Flowchart Picture (Natural aspect ratio: ONLY width specified!)
    flowchart_y = Inches(2.05)
    flowchart_w = Inches(11.733)
    pic = slide.shapes.add_picture(
        str(png_path),
        Inches(0.8),
        flowchart_y,
        width=flowchart_w,
    )
    pic_h_inches = pic.height / 914400.0
    pic_bottom_inches = 2.05 + pic_h_inches

    # Bottom Consulting Takeaways positioned below the diagram
    bottom_y_inches = pic_bottom_inches + 0.35
    bottom_h_inches = max(6.92 - bottom_y_inches, 1.80)
    card_gap = Inches(0.28)
    card_w = (flowchart_w - (2 * card_gap)) / 3

    takeaway_cards = [
        {
            "tag": "ZERO-TRUST INGRESS",
            "title": "Mutual TLS & WAF Shield",
            "metric": "100% mTLS",
            "accent": "accent",
            "bullets": [
                "Hardware-attested mTLS for 450+ POS retail terminals",
                "AWS WAF blocks malicious DDoS & bot scrapers",
                "Okta OAuth 2.0 / JWT integration for mobile clients",
            ],
        },
        {
            "tag": "DATA PRIVACY & ISOLATION",
            "title": "Tokenized PII Vault",
            "metric": "AES-256",
            "accent": "accent_teal",
            "bullets": [
                "Real-time PII tokenization before lakehouse entry",
                "Dedicated HSM keys rotated automatically every 24h",
                "Snowflake dynamic row-level masking & RBAC",
            ],
        },
        {
            "tag": "CONTINUOUS AUDITING",
            "title": "Immutable Compliance Ledger",
            "metric": "SOC-2 Ready",
            "accent": "success",
            "bullets": [
                "Write-once S3 Glacier immutable audit logging",
                "140+ automated dbt schema contract assertions",
                "Real-time drift detection alerts Slack & PagerDuty",
            ],
        },
    ]

    for idx, cdata in enumerate(takeaway_cards):
        cx = Inches(0.8) + idx * (card_w + card_gap)
        add_card(slide, theme, cx, Inches(bottom_y_inches), card_w, Inches(bottom_h_inches), bg_color=theme.get_rgb("surface"))

        # Top Accent Stripe
        stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, Inches(bottom_y_inches), card_w, Inches(0.08))
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = theme.get_rgb(cdata["accent"])
        stripe.line.fill.background()

        # Text Content
        tb = slide.shapes.add_textbox(cx + Inches(0.2), Inches(bottom_y_inches) + Inches(0.14), card_w - Inches(0.4), Inches(bottom_h_inches) - Inches(0.25))
        tf = tb.text_frame
        tf.word_wrap = True

        p_tg = tf.paragraphs[0]
        p_tg.text = cdata["tag"]
        p_tg.font.name = theme.font_family_header
        p_tg.font.size = Pt(8.5)
        p_tg.font.bold = True
        p_tg.font.color.rgb = theme.get_rgb(cdata["accent"])

        p_t = tf.add_paragraph()
        p_t.text = cdata["title"]
        p_t.font.name = theme.font_family_header
        p_t.font.size = Pt(11.0)
        p_t.font.bold = True
        p_t.font.color.rgb = theme.get_rgb("primary")
        p_t.space_before = Pt(2)

        m_val = cdata["metric"]
        p_m = tf.add_paragraph()
        p_m.text = f"TARGET KPI: [{m_val}]"
        p_m.font.name = theme.font_family_header
        p_m.font.size = Pt(8.5)
        p_m.font.bold = True
        p_m.font.color.rgb = theme.get_rgb("secondary")
        p_m.space_before = Pt(2)
        p_m.space_after = Pt(4)

        for b in cdata["bullets"]:
            pb = tf.add_paragraph()
            pb.text = f"• {b}"
            pb.font.name = theme.font_family
            pb.font.size = Pt(8.5)
            pb.font.color.rgb = theme.get_rgb("secondary")
            pb.space_before = Pt(2)

    # Footer
    add_slide_footer(slide, theme, current_idx=current_idx, total_slides=total_slides)


def build_icon_mesh_slide(
    prs: Presentation,
    theme: Any,
    current_idx: int = 4,
    total_slides: int = 4,
):
    """Slide 4: 4-Tier Zero-Trust Service Mesh Architecture with Relevant Icons."""
    slide = add_slide_with_background(prs, theme)

    add_slide_header(
        slide,
        theme,
        tracker="ARCHITECTURE TAXONOMY & SERVICE MESH  |  TIERED ENTERPRISE VISUALIZATION",
        action_title="4-Tier Zero-Trust Architecture Decouples Ingress, Streaming, Lakehouse & Compliance",
        subtitle="End-to-end integration topology combining hardware attestation, real-time message brokering, and compliance ledger.",
    )

    # 4 Architecture Columns
    col_x = Inches(0.8)
    col_y = Inches(1.72)
    col_w = Inches(2.72)
    gap_x = Inches(0.28)
    col_h = Inches(5.15)

    tiers = [
        {
            "tier_tag": "TIER 01: EDGE INGRESS",
            "tier_title": "Hardware Attestation",
            "icon": Path("assets/icons/lucide/shield-check_10B981.png"),
            "accent": "accent",
            "badge_bg": "badge_blue_fill",
            "badge_txt": "badge_blue_text",
            "sla": "SLA: 99.99% Ingress Uptime",
            "components": [
                ("450+ POS Store Registers", "Mutual TLS with TPM hardware chips"),
                ("Mobile App Clients", "Okta OAuth 2.0 / JWT signed sessions"),
                ("AWS WAF & CloudFront", "Automated anti-DDoS rate-limiting"),
                ("Edge API Gateway", "Mutual TLS certificate validation"),
            ],
            "connector_tag": "Stream ->",
        },
        {
            "tier_tag": "TIER 02: STREAMING CORE",
            "tier_title": "Event Mesh Broker",
            "icon": Path("assets/icons/lucide/cpu_F59E0B.png"),
            "accent": "warning",
            "badge_bg": "badge_amber_fill",
            "badge_txt": "badge_amber_text",
            "sla": "P99 Latency: < 45ms",
            "components": [
                ("Amazon MSK Cluster", "3-AZ partition broker replication"),
                ("Kafka Schema Registry", "Enforces strict JSON/Avro contracts"),
                ("AES-256 In-Transit", "Full payload wire-level encryption"),
                ("Write-Ahead Log", "Zero event loss across failovers"),
            ],
            "connector_tag": "Ingest ->",
        },
        {
            "tier_tag": "TIER 03: LAKEHOUSE MESH",
            "tier_title": "Medallion & Token Vault",
            "icon": Path("assets/icons/lucide/database_0D9488.png"),
            "accent": "accent_teal",
            "badge_bg": "badge_blue_fill",
            "badge_txt": "badge_blue_text",
            "sla": "Freshness: Sub-15 Mins",
            "components": [
                ("Tokenization Vault", "Pre-ingestion PII masking & salt"),
                ("Snowflake Multi-Cluster", "Dynamic scalable virtual warehouses"),
                ("Row-Level Security", "Attribute-based RBAC access policies"),
                ("Bronze-Silver-Gold", "Conformed analytical reporting marts"),
            ],
            "connector_tag": "Audit ->",
        },
        {
            "tier_tag": "TIER 04: GOVERNANCE",
            "tier_title": "Continuous Audit Ledger",
            "icon": Path("assets/icons/lucide/lock_10B981.png"),
            "accent": "success",
            "badge_bg": "badge_green_fill",
            "badge_txt": "badge_green_text",
            "sla": "Compliance: SOC-2 / ISO",
            "components": [
                ("Immutable S3 Glacier", "WORM compliance audit log archive"),
                ("140+ dbt Assertions", "Automated data contract test suites"),
                ("Slack / PagerDuty", "Instant schema anomaly drift alerts"),
                ("Automated Evidence", "Continuous SOC-2 audit log generation"),
            ],
            "connector_tag": "Certified",
        },
    ]

    for idx, tdata in enumerate(tiers):
        cx = col_x + idx * (col_w + gap_x)
        add_card(slide, theme, cx, col_y, col_w, col_h, bg_color=theme.get_rgb("surface"))

        # Top Accent Stripe
        stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, col_y, col_w, Inches(0.08))
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = theme.get_rgb(tdata["accent"])
        stripe.line.fill.background()

        # Icon Placement inside card header
        icon_path = tdata["icon"]
        if icon_path.exists():
            slide.shapes.add_picture(
                str(icon_path),
                cx + Inches(0.20),
                col_y + Inches(0.18),
                width=Inches(0.46),
                height=Inches(0.46),
            )

        # Tier Header Box
        header_x = cx + Inches(0.74)
        tb_hdr = slide.shapes.add_textbox(header_x, col_y + Inches(0.14), col_w - Inches(0.85), Inches(0.60))
        tf_hdr = tb_hdr.text_frame
        tf_hdr.word_wrap = True
        tf_hdr.margin_left = tf_hdr.margin_right = tf_hdr.margin_top = tf_hdr.margin_bottom = 0

        p_tag = tf_hdr.paragraphs[0]
        p_tag.text = tdata["tier_tag"]
        p_tag.font.name = theme.font_family_header
        p_tag.font.size = Pt(8.5)
        p_tag.font.bold = True
        p_tag.font.color.rgb = theme.get_rgb(tdata["accent"])

        p_title = tf_hdr.add_paragraph()
        p_title.text = tdata["tier_title"]
        p_title.font.name = theme.font_family_header
        p_title.font.size = Pt(11.0)
        p_title.font.bold = True
        p_title.font.color.rgb = theme.get_rgb("primary")
        p_title.space_before = Pt(2)

        # SLA Target Banner Box
        sla_box = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            cx + Inches(0.20),
            col_y + Inches(0.82),
            col_w - Inches(0.40),
            Inches(0.32),
        )
        sla_box.fill.solid()
        sla_box.fill.fore_color.rgb = theme.get_rgb("surface_muted")
        sla_box.line.color.rgb = theme.get_rgb("border")
        sla_box.line.width = Pt(0.75)

        sla_tf = sla_box.text_frame
        sla_tf.margin_left = Inches(0.1)
        sla_tf.margin_top = Inches(0.04)
        p_sla = sla_tf.paragraphs[0]
        p_sla.text = tdata["sla"]
        p_sla.font.name = theme.font_family_header
        p_sla.font.size = Pt(8.5)
        p_sla.font.bold = True
        p_sla.font.color.rgb = theme.get_rgb("primary")

        # Components Listing
        comp_y = col_y + Inches(1.24)
        tb_comp = slide.shapes.add_textbox(cx + Inches(0.20), comp_y, col_w - Inches(0.40), col_h - Inches(1.35))
        tf_comp = tb_comp.text_frame
        tf_comp.word_wrap = True
        tf_comp.margin_left = tf_comp.margin_right = tf_comp.margin_top = tf_comp.margin_bottom = 0

        p_sec = tf_comp.paragraphs[0]
        p_sec.text = "CORE CAPABILITIES:"
        p_sec.font.name = theme.font_family_header
        p_sec.font.size = Pt(8.0)
        p_sec.font.bold = True
        p_sec.font.color.rgb = theme.get_rgb("muted")
        p_sec.space_after = Pt(4)

        for c_title, c_desc in tdata["components"]:
            p_ct = tf_comp.add_paragraph()
            p_ct.text = f"• {c_title}"
            p_ct.font.name = theme.font_family_header
            p_ct.font.size = Pt(9.0)
            p_ct.font.bold = True
            p_ct.font.color.rgb = theme.get_rgb("primary")
            p_ct.space_before = Pt(4)

            p_cd = tf_comp.add_paragraph()
            p_cd.text = f"   {c_desc}"
            p_cd.font.name = theme.font_family
            p_cd.font.size = Pt(8.0)
            p_cd.font.color.rgb = theme.get_rgb("secondary")
            p_cd.space_before = Pt(1)

    # Footer
    add_slide_footer(slide, theme, current_idx=current_idx, total_slides=total_slides)


def generate_benchmark_presentation():
    proj_dir = Path("output/proj-xyz")
    diag_dir = proj_dir / "diagrams"
    pres_dir = proj_dir / "presentations"
    pres_dir.mkdir(parents=True, exist_ok=True)

    mermaid_code = """flowchart LR
    subgraph Sources["Zero-Trust Ingress"]
        POS["Store POS Terminals"]
        Mobile["Mobile App Clients"]
        WAF["AWS WAF Shield"]
    end
    subgraph Mesh["Security & Streaming Mesh"]
        Kafka["MSK Kafka Cluster"]
        Vault["Tokenization Vault"]
        Snowflake["Snowflake Lakehouse"]
    end
    subgraph Governance["Continuous Compliance"]
        Audit["Immutable Audit Logs"]
        dbt["dbt Data Contracts"]
        SOC2["SOC-2 & ISO Gating"]
    end
    POS --> WAF
    Mobile --> WAF
    WAF --> Kafka
    Kafka --> Vault
    Vault --> Snowflake
    Kafka --> Audit
    Snowflake --> dbt
    dbt --> SOC2"""

    # 1. Export 18pt, 24pt, and 32pt PNGs
    font_samples = [
        (18.0, "18pt Font", "SAMPLE 01  |  18PT FONT BENCHMARK (STANDARD)"),
        (24.0, "24pt Font", "SAMPLE 02  |  24PT FONT BENCHMARK (ELEVATED PROMINENCE)"),
        (32.0, "32pt Font", "SAMPLE 03  |  32PT FONT BENCHMARK (MAXIMUM READABILITY)"),
    ]

    png_paths = {}
    for fs, flabel, _ in font_samples:
        proj = DrawIOProject()
        pname = f"Mesh_{int(fs)}pt"
        proj.add_mermaid_page(pname, mermaid_code, theme="corporate_navy", font_size=fs)
        png_p = diag_dir / f"mesh_{int(fs)}pt.png"
        proj.export_page(pname, png_p, format="png", scale=3.0, font_size=fs, transparent=True)
        png_paths[fs] = png_p

    # 2. Build 4-Slide Presentation
    theme = get_theme("snowblue")
    prs = create_presentation(theme)

    # Slides 1, 2, 3: Flowchart Font Benchmarks
    for idx, (fs, flabel, tracker) in enumerate(font_samples, start=1):
        build_flowchart_slide(
            prs=prs,
            theme=theme,
            png_path=png_paths[fs],
            font_label=flabel,
            tracker_str=tracker,
            current_idx=idx,
            total_slides=4,
        )

    # Slide 4: 4th Visualization Type with Relevant Icons
    build_icon_mesh_slide(
        prs=prs,
        theme=theme,
        current_idx=4,
        total_slides=4,
    )

    out_file = pres_dir / "Project_XYZ_DrawIO_Review.pptx"
    prs.save(str(out_file))
    print(f"Successfully generated 4-slide benchmark presentation: {out_file}")
    return out_file


if __name__ == "__main__":
    generate_benchmark_presentation()
