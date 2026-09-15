"""
Draw.io to PowerPoint Integration and Concurrency Hard-Test
===========================================================
Hard-tests the end-to-end workflow of:
1. Safely creating a new copy of existing PowerPoint decks while respecting
   concurrent active locks from other agents or Microsoft PowerPoint.
2. Creating a new Draw.io visualization for Project XYZ from Mermaid specification.
3. Headless compilation and export of the Draw.io diagram to high-DPI PNG and SVG.
4. Introducing the compiled visualization into the new PowerPoint copy using
   executive consulting slide archetypes (McKinsey/BCG layout system).
5. Strict verification gates: SHA256 integrity of source files, lock preservation,
   and visual asset validation in the target presentation.
"""

from __future__ import annotations

import hashlib
import os
import shutil
import sys
import tempfile
from pathlib import Path
from typing import Dict, List, Optional, Tuple, Any

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

# Ensure project root is on sys.path
_ROOT = Path(__file__).resolve().parent.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from src.ppt_engine.consulting_archetypes import (
    add_card,
    add_slide_footer,
    add_slide_header,
    add_slide_with_background,
)
from src.ppt_engine.diagram_engine import DrawIOProject
from src.ppt_engine.theme_engine import get_theme


def calculate_sha256(file_path: Path) -> str:
    """Calculate SHA256 hex digest of a file."""
    h = hashlib.sha256()
    with open(file_path, "rb") as f:
        while chunk := f.read(65536):
            h.update(chunk)
    return h.hexdigest()


def safe_copy_file(src: Path, dst: Path) -> Tuple[bool, str]:
    """
    Safely copies a file even when opened with shared locks by another process.
    Uses atomic staging to avoid partial writes.
    """
    if not src.exists():
        return False, f"Source file does not exist: {src}"

    dst.parent.mkdir(parents=True, exist_ok=True)
    temp_dst = dst.with_suffix(f".tmp_{os.getpid()}_{dst.suffix}")

    try:
        # Non-exclusive binary read to prevent blocking/failing on active handles
        with open(src, "rb") as f_in:
            with open(temp_dst, "wb") as f_out:
                shutil.copyfileobj(f_in, f_out)
        
        # Atomic rename on POSIX
        temp_dst.replace(dst)
        return True, f"Successfully created isolated copy at {dst}"
    except Exception as e:
        if temp_dst.exists():
            temp_dst.unlink()
        return False, f"Error copying file: {e}"


def run_integration_test():
    print("=" * 80)
    print("PROJECT XYZ: DRAW.IO TO POWERPOINT INTEGRATION & CONCURRENCY TEST")
    print("=" * 80)

    # Paths
    proj_dir = Path("output/proj-xyz")
    pres_dir = proj_dir / "presentations"
    diag_dir = proj_dir / "diagrams"

    src_substituted = pres_dir / "01_Project_XYZ_Executive_Strategy_Deck_Substituted.pptx"
    src_clean = pres_dir / "01_Project_XYZ_Executive_Strategy_Deck.pptx"
    src_kickoff = pres_dir / "02_Project_XYZ_KickOff_Briefing.pptx"
    lock_file = pres_dir / "~$01_Project_XYZ_Executive_Strategy_Deck_Substituted.pptx"

    # -------------------------------------------------------------------------
    # Step 1: Concurrency & Lock Analysis
    # -------------------------------------------------------------------------
    print("\n[Step 1] Concurrency & Active Lock Analysis...")
    print(f"  Target deck: {src_substituted}")
    print(f"  Lock file exists: {lock_file.exists()} ({lock_file})")

    # Baseline SHA256 hashes
    initial_hashes = {}
    for f in [src_substituted, src_clean, src_kickoff]:
        if f.exists():
            initial_hashes[f] = calculate_sha256(f)
            print(f"  Baseline SHA256 [{f.name}]: {initial_hashes[f][:16]}...")

    # -------------------------------------------------------------------------
    # Step 2: Safe Isolated Copy Creation
    # -------------------------------------------------------------------------
    print("\n[Step 2] Creating Isolated Working Copy of Existing PowerPoint Decks...")
    dst_target = pres_dir / "01_Project_XYZ_Executive_Strategy_Deck_DrawIO_Integrated.pptx"
    success, msg = safe_copy_file(src_substituted, dst_target)
    print(f"  {msg}")
    assert success, "Safe copy failed!"
    assert dst_target.exists(), "Target file does not exist!"

    # Verify copy integrity
    copy_hash = calculate_sha256(dst_target)
    assert copy_hash == initial_hashes[src_substituted], "Copied file hash mismatch!"
    print(f"  ✓ Verified copied file bit-identical to source ({copy_hash[:16]}...)")

    # Also test copying the Kick-off briefing to prove universal safety
    dst_kickoff_copy = pres_dir / "02_Project_XYZ_KickOff_Briefing_SafeCopy.pptx"
    k_success, k_msg = safe_copy_file(src_kickoff, dst_kickoff_copy)
    print(f"  ✓ {k_msg}")

    # -------------------------------------------------------------------------
    # Step 3: Author and Export Draw.io Visualization for Project XYZ
    # -------------------------------------------------------------------------
    print("\n[Step 3] Authoring and Exporting Draw.io Visualization for Project XYZ...")
    master_drawio = diag_dir / "xyz_platform_architecture.drawio"
    assert master_drawio.exists(), f"Master drawio missing: {master_drawio}"

    page_name = "Security & Governance Mesh"
    mermaid_code = """flowchart LR
    subgraph Sources["Zero-Trust Ingress"]
        POS["Store POS Terminals"]
        Mobile["Mobile App Clients"]
        WAF["AWS WAF Shield"]
    end
    subgraph Mesh["Security & Streaming Mesh"]
        Kafka["MSK Kafka Cluster"]
        Vault["Tokenization Vault"]
        Snowflake["Snowflake Lakehouse"]
    end
    subgraph Governance["Continuous Compliance"]
        Audit["Immutable Audit Logs"]
        dbt["dbt Data Contracts"]
        SOC2["SOC-2 & ISO Gating"]
    end
    POS --> WAF
    Mobile --> WAF
    WAF --> Kafka
    Kafka --> Vault
    Vault --> Snowflake
    Kafka --> Audit
    Snowflake --> dbt
    dbt --> SOC2"""

    proj = DrawIOProject.load(master_drawio)
    page_id = proj.add_mermaid_page(name=page_name, mermaid_code=mermaid_code, theme="corporate_navy")
    proj.save(master_drawio)
    print(f"  ✓ Registered page '{page_name}' (ID: {page_id}) in master Draw.io file")

    # Also save as standalone modular Draw.io file
    standalone_drawio = diag_dir / "04_security_governance_mesh.drawio"
    standalone_proj = DrawIOProject()
    standalone_proj.add_mermaid_page(name=page_name, mermaid_code=mermaid_code, theme="corporate_navy")
    standalone_proj.save(standalone_drawio)
    print(f"  ✓ Saved standalone Draw.io project: {standalone_drawio}")

    # Export to High-DPI PNG and SVG
    out_png = diag_dir / "04_security_governance_mesh.png"
    out_svg = diag_dir / "04_security_governance_mesh.svg"

    proj.export_page(
        name_or_index=page_name,
        output_path=out_png,
        format="png",
        scale=3.0,
        transparent=True,
    )
    proj.export_page(
        name_or_index=page_name,
        output_path=out_svg,
        format="svg",
        transparent=True,
    )

    assert out_png.exists(), f"Exported PNG not found: {out_png}"
    assert out_svg.exists(), f"Exported SVG not found: {out_svg}"

    with Image.open(out_png) as img:
        img_w, img_h = img.size
        print(f"  ✓ Exported PNG: {out_png} ({img_w}x{img_h}px, {out_png.stat().st_size // 1024} KB)")
        print(f"  ✓ Exported SVG: {out_svg} ({out_svg.stat().st_size // 1024} KB)")

    # -------------------------------------------------------------------------
    # Step 4: Introduce Visualization into the New PowerPoint File
    # -------------------------------------------------------------------------
    print("\n[Step 4] Introducing Draw.io Visualization into New PowerPoint Copy...")
    theme = get_theme("snowblue")
    prs = Presentation(str(dst_target))
    original_slide_count = len(prs.slides)
    print(f"  Current slide count: {original_slide_count}")

    # Add new slide with 60-30-10 consulting layout
    slide = add_slide_with_background(prs, theme)
    new_slide_idx = len(prs.slides)
    total_slides = new_slide_idx

    # Header
    add_slide_header(
        slide,
        theme,
        tracker="SECURITY & DATA GOVERNANCE MESH  |  PROJECT XYZ",
        action_title="Zero-Trust Perimeter & Continuous Auditing Enforce SOC-2 Compliance Across 450+ Stores",
        subtitle="Hardware-grade mutual TLS ingress, tokenized PII isolation vault, and automated dbt data contract assertions.",
    )

    # Top Visual Card: Draw.io Diagram Hero Container
    top_x = Inches(0.8)
    top_y = Inches(1.72)
    top_w = Inches(11.733)
    top_h = Inches(2.55)

    add_card(slide, theme, top_x, top_y, top_w, top_h, bg_color=theme.get_rgb("surface"))

    # Diagram Subtitle Badge
    badge_tb = slide.shapes.add_textbox(top_x + Inches(0.25), top_y + Inches(0.12), top_w - Inches(0.5), Inches(0.35))
    badge_tf = badge_tb.text_frame
    badge_tf.margin_left = badge_tf.margin_right = badge_tf.margin_top = badge_tf.margin_bottom = 0
    p_badge = badge_tf.paragraphs[0]
    p_badge.text = "ZERO-TRUST TOPOLOGY & DATA CONTRACT MESH (DRAW.IO ENGINE)"
    p_badge.font.name = theme.font_family_header
    p_badge.font.size = Pt(9.5)
    p_badge.font.bold = True
    p_badge.font.color.rgb = theme.get_rgb("accent")

    # Embed Draw.io Compiled PNG
    diag_img_w = top_w - Inches(0.6)
    diag_img_h = top_h - Inches(0.55)
    slide.shapes.add_picture(
        str(out_png),
        top_x + Inches(0.3),
        top_y + Inches(0.42),
        width=diag_img_w,
        height=diag_img_h,
    )

    # Bottom Row: 3 Executive Consulting Takeaway Cards
    bottom_y = Inches(4.45)
    bottom_h = Inches(2.40)
    card_gap = Inches(0.28)
    card_w = (top_w - (2 * card_gap)) / 3

    takeaway_cards = [
        {
            "tag": "ZERO-TRUST INGRESS",
            "title": "Mutual TLS & WAF Shield",
            "metric": "100% mTLS",
            "accent": "accent",
            "bullets": [
                "Hardware-attested mTLS for 450+ POS retail terminals",
                "AWS WAF blocks malicious DDoS & bot scrapers",
                "Okta OAuth 2.0 / JWT integration for mobile clients",
            ],
        },
        {
            "tag": "DATA PRIVACY & ISOLATION",
            "title": "Tokenized PII Vault",
            "metric": "AES-256",
            "accent": "accent_teal",
            "bullets": [
                "Real-time PII tokenization before lakehouse entry",
                "Dedicated HSM keys rotated automatically every 24h",
                "Snowflake dynamic row-level masking & RBAC",
            ],
        },
        {
            "tag": "CONTINUOUS AUDITING",
            "title": "Immutable Compliance Ledger",
            "metric": "SOC-2 Ready",
            "accent": "success",
            "bullets": [
                "Write-once S3 Glacier immutable audit logging",
                "140+ automated dbt schema contract assertions",
                "Real-time drift detection alerts Slack & PagerDuty",
            ],
        },
    ]

    for idx, cdata in enumerate(takeaway_cards):
        cx = top_x + idx * (card_w + card_gap)
        add_card(slide, theme, cx, bottom_y, card_w, bottom_h, bg_color=theme.get_rgb("surface"))

        # Top Accent Stripe
        stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, bottom_y, card_w, Inches(0.08))
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = theme.get_rgb(cdata["accent"])
        stripe.line.fill.background()

        # Text Content
        tb = slide.shapes.add_textbox(cx + Inches(0.2), bottom_y + Inches(0.16), card_w - Inches(0.4), bottom_h - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p_tag = tf.paragraphs[0]
        p_tag.text = cdata["tag"]
        p_tag.font.name = theme.font_family_header
        p_tag.font.size = Pt(8.5)
        p_tag.font.bold = True
        p_tag.font.color.rgb = theme.get_rgb(cdata["accent"])

        p_t = tf.add_paragraph()
        p_t.text = cdata["title"]
        p_t.font.name = theme.font_family_header
        p_t.font.size = Pt(11.5)
        p_t.font.bold = True
        p_t.font.color.rgb = theme.get_rgb("primary")
        p_t.space_before = Pt(2)

        p_m = tf.add_paragraph()
        p_m.text = f"TARGET KPI: [{cdata['metric']}]"
        p_m.font.name = theme.font_family_header
        p_m.font.size = Pt(8.5)
        p_m.font.bold = True
        p_m.font.color.rgb = theme.get_rgb("secondary")
        p_m.space_before = Pt(3)
        p_m.space_after = Pt(6)

        for b in cdata["bullets"]:
            pb = tf.add_paragraph()
            pb.text = f"• {b}"
            pb.font.name = theme.font_family
            pb.font.size = Pt(8.5)
            pb.font.color.rgb = theme.get_rgb("secondary")
            pb.space_before = Pt(2)

    # Footer
    add_slide_footer(slide, theme, current_idx=new_slide_idx, total_slides=total_slides)

    # Update slide footers for all slides in the deck to reflect new total
    for i, s in enumerate(prs.slides):
        for shape in s.shapes:
            if shape.has_text_frame and shape.text_frame.text:
                text = shape.text_frame.text
                if f"/{original_slide_count}" in text:
                    new_text = text.replace(f"/{original_slide_count}", f"/{total_slides}")
                    shape.text_frame.text = new_text

    # Atomic write to destination
    temp_target = dst_target.with_suffix(".tmp.pptx")
    prs.save(str(temp_target))
    temp_target.replace(dst_target)
    print(f"  ✓ Saved enhanced PowerPoint file: {dst_target}")

    # -------------------------------------------------------------------------
    # Step 5: Strict Verification & Concurrency Assertions
    # -------------------------------------------------------------------------
    print("\n[Step 5] Running Strict Verification Gates...")

    # Gate A: Source File Immutability Assertion
    for f, initial_hash in initial_hashes.items():
        current_hash = calculate_sha256(f)
        assert current_hash == initial_hash, f"CONCURRENCY VIOLATION! Source file {f.name} was modified!"
        print(f"  ✓ Source file untouched: {f.name} (SHA256: {current_hash[:16]}...)")

    # Gate B: Lock File Preservation Assertion
    assert lock_file.exists(), "Lock file was prematurely removed or disturbed!"
    print(f"  ✓ External lock file preserved: {lock_file.name}")

    # Gate C: Target Presentation Structural Validation
    verify_prs = Presentation(str(dst_target))
    assert len(verify_prs.slides) == original_slide_count + 1, "Slide count mismatch!"
    target_slide = verify_prs.slides[-1]
    
    # Check that image shape exists on the new slide
    picture_shapes = [sp for sp in target_slide.shapes if sp.shape_type.name == "PICTURE"]
    assert len(picture_shapes) >= 1, "Draw.io picture shape missing from target slide!"
    
    # Check slide dimensions (16:9 widescreen)
    w_in = verify_prs.slide_width / Inches(1)
    h_in = verify_prs.slide_height / Inches(1)
    assert abs(w_in - 13.333) < 0.01 and abs(h_in - 7.5) < 0.01, f"Invalid aspect ratio: {w_in}x{h_in}"
    print(f"  ✓ Target deck validated: {len(verify_prs.slides)} slides, 16:9 widescreen ({w_in:.3f}\" x {h_in:.3f}\")")
    print(f"  ✓ New slide contains Draw.io picture shape: {picture_shapes[0].name}")

    print("\n" + "=" * 80)
    print("ALL VERIFICATION GATES PASSED [100% PASS]")
    print("=" * 80)


if __name__ == "__main__":
    run_integration_test()
