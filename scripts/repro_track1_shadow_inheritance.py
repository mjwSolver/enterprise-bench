#!/usr/bin/env python3
"""
Live Probe for Track 1: Theme Drop Shadow Inheritance in lingering deck builders.
Tests:
- closing_deck.py
- weekly_progress_deck.py
- reference_slides.py
- BSIM_Migration/build_final_powerpoint.py
Checks whether shapes inherit <a:effectRef idx="2"> or lack empty <a:effectLst/>.
"""

import sys
from pathlib import Path
from pptx import Presentation

def audit_presentation_shadows(prs: Presentation, deck_name: str):
    defects = []
    total_shapes = 0
    for s_idx, slide in enumerate(prs.slides):
        for sh in slide.shapes:
            total_shapes += 1
            sp = getattr(sh, "_element", getattr(sh, "element", sh))
            style = sp.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style")
            effect_ref = style.find("{http://schemas.openxmlformats.org/drawingml/2006/main}effectRef") if style is not None else None
            idx_val = effect_ref.get("idx") if effect_ref is not None else None
            
            sp_pr = getattr(sp, "spPr", None)
            effect_lst = sp_pr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst") if sp_pr is not None else None
            is_empty_lst = effect_lst is not None and len(list(effect_lst)) == 0

            # Defect condition: either theme shadow idx is active (!= "0"), or no explicit empty effectLst
            if (idx_val is not None and idx_val != "0") or not is_empty_lst:
                defects.append({
                    "slide": s_idx + 1,
                    "shape_name": sh.name,
                    "effectRef_idx": idx_val,
                    "has_empty_effectLst": is_empty_lst,
                })
    return total_shapes, defects

def main():
    print("=== Track 1: Live-Testing Drop Shadow Probe ===")
    
    # 1. Probe BSIM_Migration/build_final_powerpoint.py
    bsim_script = Path("/Users/marcelljw/Documents/antigravity/adventurous-darwin/BSIM_Migration/build_final_powerpoint.py")
    if bsim_script.exists():
        sys.path.insert(0, str(bsim_script.parent))
        import build_final_powerpoint as bfp
        bfp.generate_comparison_presentation(output_path="/tmp/bsim_comparison_test.pptx")
        prs_bsim = Presentation("/tmp/bsim_comparison_test.pptx")
        tot, defs = audit_presentation_shadows(prs_bsim, "BSIM build_final_powerpoint")
        print(f"BSIM build_final_powerpoint: {len(defs)}/{tot} shapes have drop shadow inheritance defects!")
        if defs:
            print(f"  Evidence: Slide {defs[0]['slide']} {defs[0]['shape_name']} -> effectRef idx={defs[0]['effectRef_idx']}, empty_effectLst={defs[0]['has_empty_effectLst']}")

    # 2. Probe closing_deck.py
    try:
        from src.ppt_engine.closing_deck import ClosingDeckBuilder
        prs_closing = ClosingDeckBuilder().build_all()
        tot_c, defs_c = audit_presentation_shadows(prs_closing, "closing_deck.py")
        print(f"closing_deck.py: {len(defs_c)}/{tot_c} shapes have drop shadow inheritance defects!")
        if defs_c:
            print(f"  Evidence: Slide {defs_c[0]['slide']} {defs_c[0]['shape_name']} -> effectRef idx={defs_c[0]['effectRef_idx']}, empty_effectLst={defs_c[0]['has_empty_effectLst']}")
    except Exception as e:
        print(f"Error checking closing_deck: {e}")

    # 3. Probe weekly_progress_deck.py
    try:
        from src.ppt_engine.weekly_progress_deck import WeeklyProgressDeckBuilder
        prs_weekly = WeeklyProgressDeckBuilder().build_all()
        tot_w, defs_w = audit_presentation_shadows(prs_weekly, "weekly_progress_deck.py")
        print(f"weekly_progress_deck.py: {len(defs_w)}/{tot_w} shapes have drop shadow inheritance defects!")
        if defs_w:
            print(f"  Evidence: Slide {defs_w[0]['slide']} {defs_w[0]['shape_name']} -> effectRef idx={defs_w[0]['effectRef_idx']}, empty_effectLst={defs_w[0]['has_empty_effectLst']}")
    except Exception as e:
        print(f"Error checking weekly_progress_deck: {e}")

    # 4. Probe reference_slides.py
    try:
        from src.ppt_engine.reference_slides import ReferenceDeckBuilder
        builder_ref = ReferenceDeckBuilder()
        builder_ref.build_all()
        tot_r, defs_r = audit_presentation_shadows(builder_ref.prs, "reference_slides.py")
        print(f"reference_slides.py: {len(defs_r)}/{tot_r} shapes have drop shadow inheritance defects!")
        if defs_r:
            print(f"  Evidence: Slide {defs_r[0]['slide']} {defs_r[0]['shape_name']} -> effectRef idx={defs_r[0]['effectRef_idx']}, empty_effectLst={defs_r[0]['has_empty_effectLst']}")
    except Exception as e:
        print(f"Error checking reference_slides: {e}")

if __name__ == "__main__":
    main()
