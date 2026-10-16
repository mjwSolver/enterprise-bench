#!/usr/bin/env python3
"""
Verification script for the OpenXML drop shadow sanitizer.
Verifies that shapes compiled via python-pptx have:
1. shape.shadow.inherit == False
2. <p:style><a:effectRef idx="0">
3. Explicit empty <a:effectLst/> in spPr
4. Zero <a:outerShdw> elements
"""

import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor

# Import the permanent sanitizer
from src.ppt_engine.theme_engine import strip_shape_shadows, _flatten

def verify_shape_xml(shape, name: str) -> None:
    sp = getattr(shape, "_element", shape)
    
    # 1. Verify shadow inherit
    if hasattr(shape, "shadow"):
        assert shape.shadow.inherit is False, f"[{name}] Expected shape.shadow.inherit to be False"
        
    # 2. Verify effectRef idx="0" in style (if style exists)
    style = sp.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style")
    if style is not None:
        effect_ref = style.find("{http://schemas.openxmlformats.org/drawingml/2006/main}effectRef")
        if effect_ref is not None:
            idx_val = effect_ref.get("idx")
            assert idx_val == "0", f"[{name}] Expected style/effectRef idx='0', found idx='{idx_val}'"

    # 3. Verify empty effectLst in spPr
    sp_pr = getattr(sp, "spPr", None)
    assert sp_pr is not None, f"[{name}] Missing spPr on shape element"
    
    # Check that no outerShdw exists anywhere in spPr
    outer_shadows = sp_pr.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}outerShdw")
    assert len(outer_shadows) == 0, f"[{name}] Found {len(outer_shadows)} lingering outerShdw elements!"
    
    # Check that effectLst exists and has 0 children (explicit empty tag)
    effect_lst = sp_pr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst")
    assert effect_lst is not None, f"[{name}] Missing explicit <a:effectLst/> in spPr"
    assert len(list(effect_lst)) == 0, f"[{name}] <a:effectLst/> should be empty, but has {len(list(effect_lst))} children!"

def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # blank layout
    
    # 1. Rectangle container
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(4), Inches(2))
    rect.fill.solid()
    rect.fill.fore_color.rgb = RGBColor(248, 250, 252)
    strip_shape_shadows(rect)
    
    # 2. Rounded rectangle card
    rounded = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.5), Inches(3), Inches(1))
    rounded.fill.solid()
    rounded.fill.fore_color.rgb = RGBColor(255, 255, 255)
    _flatten(rounded)
    
    # 3. Text box
    tb = slide.shapes.add_textbox(Inches(1.6), Inches(1.6), Inches(2.8), Inches(0.8))
    tb.text_frame.text = "Encapsulated Title"
    _flatten(tb)
    
    # 4. Connector line
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(4.5), Inches(2), Inches(6.5), Inches(2))
    _flatten(conn)
    
    # Verify all shapes
    verify_shape_xml(rect, "Sharp Rectangle Container")
    verify_shape_xml(rounded, "Rounded Rectangle Card")
    verify_shape_xml(tb, "Text Box")
    verify_shape_xml(conn, "Connector Line")
    
    # Save test presentation
    out_path = Path("output/presentations/openxml_sanitizer_verification.pptx")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    
    # Re-open and verify serialized XML from saved package
    prs_reloaded = Presentation(str(out_path))
    slide_reloaded = prs_reloaded.slides[0]
    for idx, sh in enumerate(slide_reloaded.shapes):
        verify_shape_xml(sh, f"Reloaded Shape {idx} ({sh.name})")
        
    print(f"SUCCESS: All 4 shapes verified in memory and on disk ({out_path}).")
    print("Zero outer drop shadows, <a:effectRef idx='0'> verified, empty <a:effectLst/> verified.")

if __name__ == "__main__":
    main()
