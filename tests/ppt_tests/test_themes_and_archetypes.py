"""
Test Suite for Theme Engine & Consulting Archetypes Subsystem
============================================================
Verifies:
  1. Theme YAML loading, parsing, and caching.
  2. Color conversions (hex_to_rgb, rgb_to_hex, get_rgb, get_hex).
  3. Geometry and typography thresholds (corner_radius, shadows, font sizes).
  4. BCG 3-Horizon Growth framework slide generation.
  5. McKinsey MECE Hypothesis Cascade slide generation.
  6. Executive Balanced Scorecard / 4-Quadrant KPI Matrix slide generation.
  7. End-to-end multi-theme presentation generation and PPTX serialization.
"""

import os
import unittest
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

from src.ppt_engine.theme_engine import (
    Theme,
    ThemeEngine,
    hex_to_rgb,
    rgb_to_hex,
    load_theme,
    get_theme,
    list_available_themes,
)
from src.ppt_engine.consulting_archetypes import (
    ConsultingDeckBuilder,
    HorizonColumnData,
    StrategyPillarData,
    ScorecardMetric,
    ScorecardQuadrantData,
    create_presentation,
    add_slide_with_background,
    add_slide_header,
    add_card,
    add_slide_footer,
    build_bcg_3_horizon_slide,
    build_mckinsey_cascade_slide,
    build_balanced_scorecard_slide,
)

WORKSPACE_ROOT = Path(__file__).resolve().parent.parent


class TestThemeEngine(unittest.TestCase):

    def setUp(self) -> None:
        self.themes_dir = WORKSPACE_ROOT / "presets" / "themes"
        self.engine = ThemeEngine(themes_dir=self.themes_dir)

    def test_list_available_themes(self) -> None:
        """Verify list_available_themes discovers preset themes."""
        themes = self.engine.list_available_themes()
        self.assertIn("default", themes)
        self.assertIn("snowblue", themes)
        self.assertIn("brickred", themes)

    def test_load_default_theme(self) -> None:
        """Verify loading and parsing of default (Clean Corporate) theme."""
        theme = self.engine.load_theme("default")
        self.assertEqual(theme.name, "default")
        self.assertEqual(theme.get_hex("primary"), "#1E293B")
        self.assertEqual(theme.get_hex("accent"), "#2563EB")
        self.assertEqual(theme.get_hex("background"), "#F8FAFC")
        self.assertEqual(theme.get_hex("surface"), "#FFFFFF")
        self.assertEqual(theme.get_hex("border"), "#E2E8F0")

        # Geometry & Shadows
        self.assertEqual(theme.corner_radius, 0)
        self.assertFalse(theme.enable_shadows)

        # RGBColor conversion
        rgb_primary = theme.get_rgb("primary")
        self.assertIsInstance(rgb_primary, RGBColor)
        self.assertEqual(rgb_primary, RGBColor(0x1E, 0x29, 0x3B))

        rgb_accent = theme.get_rgb("accent")
        self.assertEqual(rgb_accent, RGBColor(0x25, 0x63, 0xEB))

    def test_load_snowblue_theme(self) -> None:
        """Verify loading and parsing of snowblue theme."""
        theme = self.engine.load_theme("snowblue")
        self.assertEqual(theme.name, "snowblue")
        self.assertEqual(theme.get_hex("accent"), "#0284C7")
        self.assertEqual(theme.get_hex("accent_secondary"), "#06B6D4")
        self.assertEqual(theme.get_hex("border"), "#BAE6FD")

        # Geometry & Shadows
        self.assertEqual(theme.corner_radius, 8)
        self.assertFalse(theme.enable_shadows)

        # Badges
        self.assertEqual(theme.get_hex("badge_blue_fill"), "#E0F2FE")
        self.assertEqual(theme.get_hex("badge_blue_text"), "#0369A1")

    def test_load_brickred_theme(self) -> None:
        """Verify loading and parsing of brickred theme."""
        theme = self.engine.load_theme("brickred")
        self.assertEqual(theme.name, "brickred")
        self.assertEqual(theme.get_hex("accent"), "#DC2626")
        self.assertEqual(theme.get_hex("accent_secondary"), "#991B1B")
        self.assertEqual(theme.get_hex("primary"), "#0F172A")
        self.assertEqual(theme.get_hex("border"), "#FECACA")

        # Geometry
        self.assertEqual(theme.corner_radius, 0)
        self.assertFalse(theme.enable_shadows)

    def test_color_utilities(self) -> None:
        """Verify hex_to_rgb and rgb_to_hex conversion functions."""
        # 6-char hex with #
        rgb1 = hex_to_rgb("#2563EB")
        self.assertEqual(rgb1, RGBColor(37, 99, 235))
        self.assertEqual(rgb_to_hex(rgb1), "#2563EB")

        # 3-char shorthand hex
        rgb2 = hex_to_rgb("#FFF")
        self.assertEqual(rgb2, RGBColor(255, 255, 255))
        self.assertEqual(rgb_to_hex(rgb2), "#FFFFFF")

        # Invalid hex
        with self.assertRaises(ValueError):
            hex_to_rgb("invalid_hex")

    def test_typography_and_diagram_thresholds(self) -> None:
        """Verify typography and diagram threshold dictionaries."""
        theme = self.engine.load_theme("default")
        t_thresh = theme.typography_thresholds
        self.assertIn("action_title_pt", t_thresh)
        self.assertIn("tracker_pt", t_thresh)
        self.assertGreaterEqual(t_thresh["action_title_pt"], 18.0)

        d_thresh = theme.diagram_thresholds
        self.assertIn("node_fill", d_thresh)
        self.assertIn("node_stroke", d_thresh)
        self.assertIn("canvas_bg", d_thresh)


class TestConsultingArchetypes(unittest.TestCase):

    def setUp(self) -> None:
        self.themes_dir = WORKSPACE_ROOT / "presets" / "themes"
        self.default_theme = load_theme("default", themes_dir=self.themes_dir)
        self.snowblue_theme = load_theme("snowblue", themes_dir=self.themes_dir)
        self.brickred_theme = load_theme("brickred", themes_dir=self.themes_dir)
        self.output_dir = WORKSPACE_ROOT / "project_outputs" / "test_themes"
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def test_bcg_3_horizon_slide_generation(self) -> None:
        """Test BCG 3-Horizon Growth / Modernization Framework slide builder."""
        prs = create_presentation(self.default_theme)
        slide = build_bcg_3_horizon_slide(
            prs=prs,
            theme=self.default_theme,
            tracker="Growth Strategy  |  BCG 3-Horizon Framework",
            action_title="Sequenced Portfolio Execution Accelerates Cloud Modernization",
            subtitle="Balancing near-term core optimizations with transformational platform bets.",
        )

        self.assertEqual(len(prs.slides), 1)
        # Verify shapes generated (background, header boxes, 3 columns x (card, stripe, badge, texts), footer)
        self.assertGreater(len(slide.shapes), 12)

        # Test with custom Horizon data
        custom_horizons = [
            HorizonColumnData(
                horizon_tag="HORIZON 1 (0-6M)",
                title="Immediate Quick Wins",
                strategic_focus="Automate CI/CD pipelines and eliminate downtime.",
                metric_highlight="100% CI/CD",
                metric_label="Full automation coverage",
                initiatives=["Zero-downtime blue/green deployment", "Automated container security"],
            ),
            HorizonColumnData(
                horizon_tag="HORIZON 2 (6-18M)",
                title="Microservices Mesh",
                strategic_focus="Decouple monolithic core services.",
                metric_highlight="4.2x Speed",
                metric_label="Release frequency velocity",
                initiatives=["Deploy Envoy service mesh", "Federate GraphQL data gateway"],
            ),
            HorizonColumnData(
                horizon_tag="HORIZON 3 (18-36M)",
                title="Autonomous AI Mesh",
                strategic_focus="Autonomous self-healing cluster mesh.",
                metric_highlight="Sub-10s MTTR",
                metric_label="Autonomous node failover",
                initiatives=["Pilot predictive load routing", "Multi-cloud automated failover"],
            ),
        ]
        prs2 = create_presentation(self.snowblue_theme)
        slide2 = build_bcg_3_horizon_slide(
            prs=prs2,
            theme=self.snowblue_theme,
            horizons=custom_horizons,
        )
        self.assertEqual(len(prs2.slides), 1)

    def test_mckinsey_cascade_slide_generation(self) -> None:
        """Test McKinsey MECE Hypothesis / Strategy Cascade slide builder."""
        prs = create_presentation(self.brickred_theme)
        slide = build_mckinsey_cascade_slide(
            prs=prs,
            theme=self.brickred_theme,
            tracker="Corporate Strategy  |  McKinsey MECE Model",
            action_title="Hypothesis Cascade Resolves Core Enterprise Constraints",
            subtitle="Three mutually exclusive pillars to expand operating margins.",
            hypothesis_statement="CORE HYPOTHESIS: Decoupling tier-1 monolithic services lowers infrastructure TCO by 35%.",
            synthesis_conclusion="DECISION: Fast-track microservices migration in Q2 to unlock TCO benefits.",
        )

        self.assertEqual(len(prs.slides), 1)
        # Check that shapes include hypothesis box, pillars, synthesis banner
        self.assertGreater(len(slide.shapes), 10)

    def test_balanced_scorecard_slide_generation(self) -> None:
        """Test Executive Balanced Scorecard 4-Quadrant KPI Matrix slide builder."""
        prs = create_presentation(self.default_theme)
        slide = build_balanced_scorecard_slide(
            prs=prs,
            theme=self.default_theme,
            tracker="Executive Governance  |  Balanced Scorecard",
            action_title="All 4 Strategic Dimensions Exceed Target Milestones",
            subtitle="Financial, Customer, Process, and Cyber Resilience scorecards.",
        )

        self.assertEqual(len(prs.slides), 1)
        # 4 Quadrants x (card, stripe, header text, 2 metric sub-cards with pills)
        self.assertGreater(len(slide.shapes), 15)

    def test_consulting_deck_builder_e2e(self) -> None:
        """Test ConsultingDeckBuilder end-to-end presentation assembly and file saving across themes."""
        themes_to_test = ["default", "snowblue", "brickred"]

        for theme_name in themes_to_test:
            builder = ConsultingDeckBuilder(theme=theme_name)

            # Slide 1: BCG 3-Horizon
            builder.add_bcg_3_horizon_slide()

            # Slide 2: McKinsey MECE Strategy Cascade
            builder.add_mckinsey_cascade_slide()

            # Slide 3: Executive Balanced Scorecard
            builder.add_balanced_scorecard_slide()

            self.assertEqual(len(builder.prs.slides), 3)

            output_file = self.output_dir / f"consulting_deck_{theme_name}.pptx"
            saved_path = builder.save(output_file)

            self.assertTrue(saved_path.exists())
            self.assertGreater(saved_path.stat().st_size, 5000)

            # Validate that saved PPTX can be reloaded
            reloaded_prs = Presentation(str(saved_path))
            self.assertEqual(len(reloaded_prs.slides), 3)


if __name__ == "__main__":
    unittest.main()
