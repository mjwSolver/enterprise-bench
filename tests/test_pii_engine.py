"""
Comprehensive Test Suite for Universal PII Engine
==================================================
Tests detection rules, format handlers (DOCX, PPTX, XLSX), custom format registration,
mapping generation, and end-to-end sanitization pipelines.
"""

import os
from pathlib import Path
import pytest
import docx
import pptx
import openpyxl

from src.core.pii.models import (
    PiiCategory,
    PiiConfidence,
    ReplacementMapping,
    SanitizeResult,
    TextNode,
)
from src.core.pii.detectors import (
    ContractIdDetector,
    DetectorPipeline,
    EmailDetector,
    GovernmentIdDetector,
    PhoneDetector,
)
from src.core.pii.handlers.base import BaseFormatHandler
from src.core.pii.handlers.docx_handler import DocxHandler
from src.core.pii.handlers.pptx_handler import PptxHandler
from src.core.pii.handlers.xlsx_handler import XlsxHandler
from src.core.pii.registry import HandlerRegistry, register_handler
from src.core.pii.engine import (
    PiiEngine,
    load_mapping_json,
    save_mapping_json,
)


# ============================================================================
# 1. Detector Tests
# ============================================================================

def test_detectors_regex():
    email_det = EmailDetector()
    phone_det = PhoneDetector()
    gov_det = GovernmentIdDetector()
    contract_det = ContractIdDetector()

    # Email
    node_email = TextNode(location="Test", text="Please contact yana_mulyanah@toyotatsusho.co.id regarding terms.", node_type="para")
    matches = email_det.detect(node_email)
    assert len(matches) == 1
    assert matches[0].raw_value == "yana_mulyanah@toyotatsusho.co.id"
    assert matches[0].category == PiiCategory.EMAIL

    # Indonesian Phone
    node_phone = TextNode(location="Test", text="Call me at 081310352687 or +6281807475266 today.", node_type="para")
    matches = phone_det.detect(node_phone)
    assert len(matches) == 2
    assert matches[0].category == PiiCategory.PHONE

    # Contract ID
    node_contract = TextNode(location="Test", text="Under Agreement No. 0073/ECP/IMPL/I/2026 executed in Jakarta.", node_type="para")
    matches = contract_det.detect(node_contract)
    assert len(matches) == 1
    assert matches[0].raw_value == "0073/ECP/IMPL/I/2026"
    assert matches[0].category == PiiCategory.CONTRACT_ID


def test_pipeline_entity_dictionary():
    pipeline = DetectorPipeline()
    node = TextNode(
        location="Slide 1",
        text="Presented by Agus Pramono, Project Manager at PT Mitra Integrasi Informatika for Toyota Tsusho.",
        node_type="slide_shape",
    )
    matches = pipeline.scan_node(node)
    matched_values = [m.raw_value.lower() for m in matches]
    assert any("agus pramono" in v for v in matched_values)
    assert any("mitra integrasi informatika" in v for v in matched_values)
    assert any("toyota tsusho" in v for v in matched_values)


# ============================================================================
# 2. DOCX Handler Test
# ============================================================================

def test_docx_handling(tmp_path: Path):
    doc_path = tmp_path / "test_contract.docx"
    doc = docx.Document()
    doc.core_properties.author = "Agus Pramono"
    doc.core_properties.comments = "Confidential Client Data"

    # Body
    doc.add_paragraph("Agreement between PT Toyota Tsusho Indonesia and Consultant.")
    # Table
    tbl = doc.add_table(rows=2, cols=2)
    tbl.rows[0].cells[0].paragraphs[0].text = "Contact Person"
    tbl.rows[0].cells[1].paragraphs[0].text = "Agus Pramono"
    tbl.rows[1].cells[0].paragraphs[0].text = "Email"
    tbl.rows[1].cells[1].paragraphs[0].text = "agus.Pramono@metrodata.co.id"
    # Header
    doc.sections[0].header.paragraphs[0].text = "Confidential - TTI Snowflake Implementation"

    doc.save(str(doc_path))

    # Test Extraction
    handler = DocxHandler()
    nodes = handler.extract_text_nodes(doc_path)
    assert len(nodes) >= 4
    meta = handler.extract_metadata(doc_path)
    assert meta["author"] == "Agus Pramono"

    # Test Engine Scan & Map & Sanitize
    engine = PiiEngine()
    report = engine.scan_file(doc_path)
    assert report.total_matches > 0

    mapping = engine.generate_mapping(report, strategy="jinja")
    clean_path = tmp_path / "clean_contract.docx"
    result = engine.sanitize_file(doc_path, clean_path, mapping)
    assert result.replacements_applied > 0

    # Verify sanitized output
    clean_doc = docx.Document(str(clean_path))
    clean_text = " ".join(p.text for p in clean_doc.paragraphs)
    assert "Agus Pramono" not in clean_text
    assert "Toyota Tsusho" not in clean_text
    assert clean_doc.core_properties.author != "Agus Pramono"


# ============================================================================
# 3. PPTX Handler Test
# ============================================================================

def test_pptx_handling(tmp_path: Path):
    ppt_path = tmp_path / "test_presentation.pptx"
    prs = pptx.Presentation()
    prs.core_properties.author = "Fredric Retanubun"

    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Project Kick-off for Toyota Tsusho"
    slide.placeholders[1].text = "Presented by Adam Wijaya (Adam.Wijaya@metrodata.co.id)"

    prs.save(str(ppt_path))

    # Test Extraction
    handler = PptxHandler()
    nodes = handler.extract_text_nodes(ppt_path)
    assert len(nodes) >= 2
    meta = handler.extract_metadata(ppt_path)
    assert meta["author"] == "Fredric Retanubun"

    # Test Scan & Sanitize
    engine = PiiEngine()
    report = engine.scan_file(ppt_path)
    assert report.total_matches >= 2

    clean_path = tmp_path / "clean_presentation.pptx"
    mapping = engine.generate_mapping(report, strategy="jinja")
    result = engine.sanitize_file(ppt_path, clean_path, mapping)
    assert result.replacements_applied >= 2

    # Verify
    clean_prs = pptx.Presentation(str(clean_path))
    assert clean_prs.core_properties.author != "Fredric Retanubun"
    clean_sub = clean_prs.slides[0].placeholders[1].text
    assert "Adam Wijaya" not in clean_sub
    assert "@metrodata.co.id" not in clean_sub


# ============================================================================
# 4. XLSX Handler Test
# ============================================================================

def test_xlsx_handling(tmp_path: Path):
    xlsx_path = tmp_path / "test_register.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Toyota Tsusho Contacts"
    wb.properties.creator = "Mulyanah"

    ws.append(["Name", "Email", "Phone"])
    ws.append(["Mulyanah", "yana_mulyanah@toyotatsusho.co.id", "081310352687"])
    wb.save(str(xlsx_path))
    wb.close()

    handler = XlsxHandler()
    nodes = handler.extract_text_nodes(xlsx_path)
    assert any("yana_mulyanah@toyotatsusho.co.id" in n.text for n in nodes)

    engine = PiiEngine()
    report = engine.scan_file(xlsx_path)
    assert report.total_matches >= 3  # Mulyanah, email, phone

    clean_path = tmp_path / "clean_register.xlsx"
    mapping = engine.generate_mapping(report, strategy="jinja")
    result = engine.sanitize_file(xlsx_path, clean_path, mapping)
    assert result.replacements_applied >= 3

    clean_wb = openpyxl.load_workbook(str(clean_path))
    clean_ws = clean_wb.active
    all_cells = [str(cell.value) for row in clean_ws.iter_rows() for cell in row if cell.value]
    joined_content = " ".join(all_cells)

    assert "yana_mulyanah@toyotatsusho.co.id" not in joined_content
    assert "081310352687" not in joined_content
    assert "Mulyanah" not in joined_content
    assert clean_wb.properties.creator != "Mulyanah"
    clean_wb.close()


# ============================================================================
# 5. Future Format Extensibility Test
# ============================================================================

def test_custom_format_extensibility(tmp_path: Path):
    # Register a new dummy handler for .custom files
    class CustomTextHandler(BaseFormatHandler):
        @classmethod
        def supported_extensions(cls):
            return {".custom"}

        def extract_metadata(self, path: Path):
            return {"creator": "Test Author"}

        def extract_text_nodes(self, path: Path):
            text = path.read_text(encoding="utf-8")
            return [TextNode(location="File Content", text=text, node_type="raw_text")]

        def apply_replacements(self, input_path, output_path, replacements, strip_metadata=True, metadata_overrides=None):
            text = input_path.read_text(encoding="utf-8")
            count = 0
            for target, repl in replacements.items():
                if target in text:
                    text = text.replace(target, repl)
                    count += 1
            output_path.write_text(text, encoding="utf-8")
            return SanitizeResult(
                source_path=str(input_path),
                output_path=str(output_path),
                replacements_applied=count,
                metadata_scrubbed=strip_metadata,
            )

    registry = HandlerRegistry()
    registry.register(CustomTextHandler)
    engine = PiiEngine(registry=registry)

    custom_file = tmp_path / "data.custom"
    custom_file.write_text("Hello from admin@company.org with contract 0012/ABC/DEF/X/2026", encoding="utf-8")

    report = engine.scan_file(custom_file)
    assert report.total_matches == 2  # email + contract ID

    clean_file = tmp_path / "data_clean.custom"
    mapping = engine.generate_mapping(report, strategy="jinja")
    result = engine.sanitize_file(custom_file, clean_file, mapping)

    clean_content = clean_file.read_text(encoding="utf-8")
    assert "admin@company.org" not in clean_content
    assert "{{ contract_number }}" in clean_content or "{{ contact_email" in clean_content


# ============================================================================
# 6. Mapping JSON Persistence
# ============================================================================

def test_mapping_io(tmp_path: Path):
    mapping = ReplacementMapping(
        strategy="jinja",
        replacements={"Tadahiko Onaka": "{{ client_exec }}"},
        strip_metadata=True,
    )
    json_path = tmp_path / "test_mapping.json"
    save_mapping_json(mapping, json_path)
    assert json_path.exists()

    loaded = load_mapping_json(json_path)
    assert loaded.strategy == "jinja"
    assert loaded.replacements["Tadahiko Onaka"] == "{{ client_exec }}"
