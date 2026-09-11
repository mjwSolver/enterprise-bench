"""
Unit Tests for Spreadsheet Calculator & Engine (In-Memory / Tempfile)
"""

from pathlib import Path
import openpyxl
import pytest
from typer.testing import CliRunner

from src.cli import app
from src.xlsx_engine.calculator_stamper import CalculatorStamper, calculate_spreadsheet


def test_calculator_stamper_basic(tmp_path: Path):
    # Create sample workbook template
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calculations"
    ws["A1"] = "Metric"
    ws["B1"] = "Value"
    ws["A2"] = "Initial Users"
    ws["B2"] = 100
    ws["A3"] = "Client Name: {{ client_name }}"
    ws["B3"] = 0
    ws["A4"] = "Total Sum"
    ws["B4"] = "=SUM(B2:B3)"

    template_path = tmp_path / "sample_template.xlsx"
    wb.save(str(template_path))
    wb.close()

    # Calculate with data injection
    out_path = tmp_path / "output_calculated.xlsx"
    data = {
        "Calculations!B2": 500,
        "client_name": "Acme Global",
        "append_rows": {
            "Calculations": [
                ["Row 5 Extra", 42],
            ]
        }
    }

    result = calculate_spreadsheet(template=template_path, data=data, output_path=out_path)
    assert result.exists()

    # Verify updated content
    wb_out = openpyxl.load_workbook(str(out_path), data_only=False)
    ws_out = wb_out["Calculations"]
    assert ws_out["B2"].value == 500
    assert "Acme Global" in ws_out["A3"].value
    assert ws_out["B4"].value == "=SUM(B2:B3)"  # Formula preserved
    assert ws_out["A5"].value == "Row 5 Extra"
    assert ws_out["B5"].value == 42
    wb_out.close()


def test_cli_xlsx_calculate(tmp_path: Path):
    runner = CliRunner()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Score"
    ws["A2"] = "{{ score }}"
    template_path = tmp_path / "cli_test_template.xlsx"
    wb.save(str(template_path))
    wb.close()

    data_path = tmp_path / "data.json"
    data_path.write_text('{"score": 98.5}', encoding="utf-8")

    out_path = tmp_path / "cli_out.xlsx"

    result = runner.invoke(app, [
        "xlsx", "calculate",
        "--template", str(template_path),
        "--data", str(data_path),
        "--output", str(out_path),
    ])

    assert result.exit_code == 0
    assert "Spreadsheet calculated and saved" in result.output
    assert out_path.exists()

    wb_res = openpyxl.load_workbook(str(out_path))
    assert "98.5" in str(wb_res.active["A2"].value)
    wb_res.close()
