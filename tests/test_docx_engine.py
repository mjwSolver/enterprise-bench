"""
Unit Tests for Docx Engine
"""
from pathlib import Path
import pytest
from docx import Document
from src.docx_engine.table_engine import set_cell_background, style_table
from src.docx_engine.document_linter import lint_document, LintReport
from src.docx_engine.template_stamper import TemplateStamper


def test_table_styling(tmp_path: Path):
    doc = Document()
    table = doc.add_table(rows=3, cols=3)
    for r in range(3):
        for c in range(3):
            table.cell(r, c).text = f"R{r}C{c}"

    style_table(table, header_bg="#2563EB", zebra_even_bg="#F1F5F9")
    out_file = tmp_path / "table_test.docx"
    doc.save(str(out_file))
    assert out_file.exists()


def test_document_linter(tmp_path: Path):
    # Dirty doc with unrendered jinja
    doc_dirty = Document()
    doc_dirty.add_paragraph("This is an unrendered tag: {{ client_name }}")
    dirty_path = tmp_path / "dirty.docx"
    doc_dirty.save(str(dirty_path))

    report_dirty = lint_document(dirty_path)
    assert not report_dirty.passed
    assert report_dirty.total_issues >= 1

    # Clean doc
    doc_clean = Document()
    doc_clean.add_paragraph("This is completely clean text.")
    clean_path = tmp_path / "clean.docx"
    doc_clean.save(str(clean_path))

    report_clean = lint_document(clean_path)
    assert report_clean.passed
    assert report_clean.total_issues == 0
